from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, send_from_directory, make_response, g, abort
from flask_socketio import SocketIO, emit, join_room
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from datetime import datetime, date, timedelta, timezone
try:
    from zoneinfo import ZoneInfo
except ImportError:
    ZoneInfo = None
from functools import wraps
from collections import OrderedDict
from apscheduler.schedulers.background import BackgroundScheduler
import sqlite3
import os
import threading
import time
import uuid
import json
import calendar
import re
import base64
from openpyxl import load_workbook
try:
    import xlrd
except ImportError:
    xlrd = None

try:
    from pywebpush import webpush
except ImportError:
    webpush = None

try:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import ec
except ImportError:
    serialization = None
    ec = None

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('YURT_SECRET_KEY', 'yurt-gizli-anahtar-2024')
app.config['MAX_CONTENT_LENGTH'] = 60 * 1024 * 1024  # 60 MB (video paylaşımı için)
# Yüklenen profil/gönderi dosyaları değişmeyen (immutable) adlara sahip (uuid), bu yüzden
# tarayıcının onları agresifçe önbelleklemesi güvenli — tünel/ağ üzerinden tekrar tekrar
# aynı görseli çekmeyi engelleyerek gözle görülür bir hız kazancı sağlar.
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 60 * 60 * 24 * 30  # 30 gün
# "Beni Hatırla" özelliği: kullanıcı işaretlerse oturum çerezi 30 gün boyunca
# tarayıcı kapansa bile geçerli kalır (session.permanent = True olduğunda kullanılır).
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=30)

# YURT uygulamasının tüm kullanıcıya gösterilen saatleri Türkiye saatine sabitlenir.
# SQLite CURRENT_TIMESTAMP UTC tuttuğu için görüntüleme sırasında UTC -> Europe/Istanbul
# dönüşümü yapıyoruz; böylece sunucunun saat dilimi ne olursa olsun saatler tutarlı kalır.
ISTANBUL_TZ = ZoneInfo('Europe/Istanbul') if ZoneInfo else timezone(timedelta(hours=3))

def now_tr():
    return datetime.now(ISTANBUL_TZ)

def get_app_setting(key, default=None):
    try:
        conn = get_db()
        row = conn.execute('SELECT value FROM app_settings WHERE key = ?', (key,)).fetchone()
        conn.close()
        return row['value'] if row else default
    except Exception:
        return default

def set_app_setting(key, value):
    conn = get_db()
    conn.execute('''
        INSERT INTO app_settings(key, value, updated_at) VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=CURRENT_TIMESTAMP
    ''', (key, str(value)))
    conn.commit()
    conn.close()

def utc_to_tr(value):
    if not value:
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip()
        try:
            dt = datetime.fromisoformat(text.replace('Z', '+00:00'))
        except ValueError:
            try:
                dt = datetime.strptime(text[:19], '%Y-%m-%d %H:%M:%S')
            except ValueError:
                return value
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(ISTANBUL_TZ)

def format_datetime_tr(value, with_date=True):
    dt = utc_to_tr(value)
    if not isinstance(dt, datetime):
        return str(value or '')
    return dt.strftime('%d.%m.%Y %H:%M' if with_date else '%H:%M')


# ÖNEMLİ (STABİLİTE): async_mode'u ortam değişkeniyle seçilebilir yaptık.
# Varsayılan "threading" — ekstra bağımlılık gerektirmez ve test edilmiştir.
# 150-200 kullanıcı gibi daha kalabalık/eşzamanlı kullanım ve tünel (localtunnel/ngrok)
# üzerinden erişimde daha iyi performans için YURT_ASYNC_MODE=eventlet ile
# çalıştırılabilir (bkz. run_production.py). eventlet kullanılacaksa monkey_patch()
# diğer TÜM importlardan önce çağrılmalıdır — bu nedenle app.py'nin kendisi eventlet'i
# import etmez, bunun yerine run_production.py app.py'den önce patch'i uygular.
ASYNC_MODE = os.environ.get('YURT_ASYNC_MODE', 'threading')
socketio = SocketIO(app, cors_allowed_origins="*", async_mode=ASYNC_MODE)

DATABASE = 'database.db'
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_PROFILE = os.path.join(BASE_DIR, 'static', 'uploads', 'profiles')
UPLOAD_GROUPS = os.path.join(BASE_DIR, 'static', 'uploads', 'groups')
UPLOAD_MESSAGES = os.path.join(BASE_DIR, 'static', 'uploads', 'messages')
UPLOAD_POSTS = os.path.join(BASE_DIR, 'static', 'uploads', 'posts')
UPLOAD_STORIES = os.path.join(BASE_DIR, 'static', 'uploads', 'stories')
UPLOAD_QUIZ = os.path.join(BASE_DIR, 'static', 'uploads', 'quiz')
UPLOAD_EVENTS = os.path.join(BASE_DIR, 'static', 'uploads', 'events')
UPLOAD_MENU = os.path.join(BASE_DIR, 'static', 'uploads', 'menus')
os.makedirs(UPLOAD_MENU, exist_ok=True)
os.makedirs(UPLOAD_PROFILE, exist_ok=True)
os.makedirs(UPLOAD_GROUPS, exist_ok=True)
os.makedirs(UPLOAD_MESSAGES, exist_ok=True)
os.makedirs(UPLOAD_POSTS, exist_ok=True)
os.makedirs(UPLOAD_STORIES, exist_ok=True)
os.makedirs(UPLOAD_QUIZ, exist_ok=True)
os.makedirs(UPLOAD_EVENTS, exist_ok=True)

IMAGE_EXT = {'png', 'jpg', 'jpeg', 'gif', 'webp'}
VIDEO_EXT = {'mp4', 'mov', 'webm', 'm4v'}
AUDIO_EXT = {'mp3', 'wav', 'ogg', 'm4a', 'aac', 'webm'}

# Web Push / VAPID anahtarı ilk çalıştırmada otomatik oluşturulur.
VAPID_PRIVATE_FILE = os.path.join(BASE_DIR, 'instance', 'vapid_private.pem')
VAPID_SUBJECT = os.environ.get('YURT_VAPID_SUBJECT', 'mailto:admin@yurt.local')

def _ensure_vapid_key():
    if serialization is None or ec is None:
        return
    os.makedirs(os.path.dirname(VAPID_PRIVATE_FILE), exist_ok=True)
    if os.path.exists(VAPID_PRIVATE_FILE):
        return
    key = ec.generate_private_key(ec.SECP256R1())
    pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())
    with open(VAPID_PRIVATE_FILE, 'wb') as f:
        f.write(pem)

def _vapid_public_key():
    if serialization is None or ec is None:
        return ''
    _ensure_vapid_key()
    try:
        with open(VAPID_PRIVATE_FILE, 'rb') as f:
            key = serialization.load_pem_private_key(f.read(), password=None)
        raw = key.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)
        return base64.urlsafe_b64encode(raw).rstrip(b'=').decode('ascii')
    except Exception:
        return ''

_ensure_vapid_key()

# ==================== VERİTABANI ====================

def get_db():
    # Veritabanı bağlantısını ve cursor'ı önce oluştur.
    # Çamaşırhane tabloları bu fonksiyonun içinde oluşturulduğu için
    # cursor'ın daha önce tanımlanmış olması gerekir.
    conn = sqlite3.connect(DATABASE, timeout=10)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    # SQLite connection pragmas must be applied before any transaction starts.
    # In particular, synchronous cannot be changed while a transaction is active.
    conn.execute('PRAGMA foreign_keys = ON')
    conn.execute('PRAGMA journal_mode = WAL')
    conn.execute('PRAGMA synchronous = NORMAL')
    conn.execute('PRAGMA busy_timeout = 5000')

    # ÇAMAŞIRHANE
    c.execute('''CREATE TABLE IF NOT EXISTS laundry_machines (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL UNIQUE,
        slot_minutes INTEGER DEFAULT 60,
        enabled INTEGER DEFAULT 1,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS laundry_reservations (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        machine_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        start_at TEXT NOT NULL,
        end_at TEXT NOT NULL,
        status TEXT DEFAULT 'reserved',
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(machine_id) REFERENCES laundry_machines(id) ON DELETE CASCADE,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS laundry_waitlist (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        machine_id INTEGER NOT NULL,
        user_id INTEGER NOT NULL,
        notified_at TEXT,
        created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(machine_id, user_id),
        FOREIGN KEY(machine_id) REFERENCES laundry_machines(id) ON DELETE CASCADE,
        FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
    )''')
    for machine_name in ['Çamaşır Makinesi 1','Çamaşır Makinesi 2','Çamaşır Makinesi 3','Çamaşır Makinesi 4','Çamaşır Makinesi 5','Çamaşır Makinesi 6']:
        c.execute('INSERT OR IGNORE INTO laundry_machines(name, slot_minutes, enabled) VALUES (?,60,1)', (machine_name,))

    # DDL + seed işlemlerinin transaction'ını kapat. Böylece sonraki bağlantı
    # ayarlarında "Safety level may not be changed inside a transaction" oluşmaz.
    conn.commit()
    return conn


def _ensure_column(conn, table, column, definition):
    """Var olan veritabanlarını bozmadan yeni sütun ekler (küçük migration yardımcı fonksiyonu)."""
    cols = [r['name'] for r in conn.execute(f'PRAGMA table_info({table})').fetchall()]
    if column not in cols:
        conn.execute(f'ALTER TABLE {table} ADD COLUMN {column} {definition}')


def _migrate_direct_messages_to_conversations(conn):
    """Eski birebir mesajları (direct_messages) yeni conversations/conv_messages
    şemasına bir kereye mahsus aktarır. Eski tablo silinmez, sadece okunur;
    böylece hiçbir kullanıcı ya da mesaj kaybolmaz. app_settings üzerindeki
    bayrak sayesinde bu fonksiyon her uygulama başlangıcında güvenle çağrılabilir."""
    already = conn.execute(
        "SELECT value FROM app_settings WHERE key = 'dm_migrated_v1'"
    ).fetchone()
    if already and already['value'] == '1':
        return

    rows = conn.execute('SELECT * FROM direct_messages ORDER BY id ASC').fetchall()
    pair_to_conv = {}
    last_read_tracker = {}  # (conv_id, user_id) -> last read conv_message id

    for dm in rows:
        pair_key = tuple(sorted((dm['sender_id'], dm['receiver_id'])))
        conv_id = pair_to_conv.get(pair_key)
        if conv_id is None:
            cur = conn.execute(
                "INSERT INTO conversations (type, created_by, created_at, last_message_at) VALUES ('dm', ?, ?, ?)",
                (dm['sender_id'], dm['created_at'], dm['created_at'])
            )
            conv_id = cur.lastrowid
            for uid in pair_key:
                conn.execute(
                    'INSERT OR IGNORE INTO conversation_members (conversation_id, user_id, role, joined_at) VALUES (?, ?, ?, ?)',
                    (conv_id, uid, 'member', dm['created_at'])
                )
            pair_to_conv[pair_key] = conv_id

        cur = conn.execute(
            '''INSERT INTO conv_messages (conversation_id, sender_id, content, created_at)
               VALUES (?, ?, ?, ?)''',
            (conv_id, dm['sender_id'], dm['message'], dm['created_at'])
        )
        new_msg_id = cur.lastrowid
        conn.execute('UPDATE conversations SET last_message_at = ? WHERE id = ?', (dm['created_at'], conv_id))
        if dm['is_read']:
            last_read_tracker[(conv_id, dm['receiver_id'])] = new_msg_id

    for (conv_id, user_id), last_read_id in last_read_tracker.items():
        conn.execute(
            'UPDATE conversation_members SET last_read_message_id = ? WHERE conversation_id = ? AND user_id = ?',
            (last_read_id, conv_id, user_id)
        )

    conn.execute(
        "INSERT INTO app_settings (key, value) VALUES ('dm_migrated_v1', '1') "
        "ON CONFLICT(key) DO UPDATE SET value = '1'"
    )
    conn.commit()


def _migrate_posts_allow_text(conn):
    """Eski 'posts' tablosunda media_path NOT NULL idi (sadece görsel/video).
    Metin paylaşımını desteklemek için tabloyu, veriyi koruyarak yeniden kurar."""
    info = conn.execute("PRAGMA table_info(posts)").fetchall()
    if not info:
        return  # tablo henüz yok, aşağıdaki CREATE TABLE IF NOT EXISTS zaten doğru şemayla oluşturacak
    media_col = next((r for r in info if r['name'] == 'media_path'), None)
    if media_col is None or media_col['notnull'] != 1:
        return  # zaten güncel şema

    conn.execute('ALTER TABLE posts RENAME TO posts_old_migrate')
    conn.execute('''
        CREATE TABLE posts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            caption TEXT,
            media_path TEXT,
            media_type TEXT NOT NULL CHECK(media_type IN ('image', 'video', 'text')),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')
    conn.execute('''
        INSERT INTO posts (id, user_id, caption, media_path, media_type, created_at)
        SELECT id, user_id, caption, media_path, media_type, created_at FROM posts_old_migrate
    ''')
    conn.execute('DROP TABLE posts_old_migrate')


PROFILE_THEMES = {'default': 'Standart'}

# ==================== ROL & YETKİ SİSTEMİ ====================
# Roller veritabanında (roles tablosu) tutulur. Admin panelindeki
# "Roller & Yetkiler" sekmesinden yeni rol eklenebilir, mevcut roller
# düzenlenebilir ve her role:
#   * hangi SAYFALARA erişebileceği (page_*)
#   * hangi YÖNETİM işlemlerini yapabileceği (manage_* vb.)
# tek tek atanabilir. Aşağıdaki DEFAULT_ROLES_SEED sadece ilk kurulumda
# (ve sonradan eklenen yeni varsayılan roller için) tabloya yazılır;
# admin panelinden yapılan düzenlemelerin üzerine yazılmaz.

# Sayfa erişim yetkileri: bir role bu anahtar False olarak kaydedilmedikçe
# erişim AÇIK kabul edilir (yeni sayfa eklendiğinde kimse kilitlenmesin diye).
# 'endpoints' -> Flask endpoint adları (tam eşleşme) ya da 'önek*' (başlangıç).
PAGE_PERMISSIONS = OrderedDict([
    ('page_feed', {'label': 'Sosyal Akış', 'icon': 'bi-grid-3x3', 'desc': 'Gönderi ve hikâye akışı, paylaşım yapma, beğeni/yorum.',
                   'endpoints': ['feed', 'create_post', 'like_post', 'comment_post', 'add_comment', 'delete_post',
                                 'create_story', 'get_user_stories', 'view_story', 'delete_story']}),
    ('page_events', {'label': 'Etkinlikler', 'icon': 'bi-calendar3', 'desc': 'Etkinlikleri görme, oluşturma ve katılma.',
                     'endpoints': ['events', 'create_event', 'event_detail', 'join_event', 'leave_event']}),
    ('page_menu', {'label': 'Yemekhane', 'icon': 'bi-cup-hot', 'desc': 'Günlük menü, puanlama ve yorumlar.',
                   'endpoints': ['menu_page', 'rate_menu', 'api_get_menu_date']}),
    ('page_games', {'label': 'Oyun Merkezi', 'icon': 'bi-joystick', 'desc': 'Mini oyunlar ve oyun skorları.',
                    'endpoints': ['games_hub', 'game_play', 'submit_game_score']}),
    ('page_quiz', {'label': 'Yarışmalar', 'icon': 'bi-patch-question', 'desc': 'Canlı quiz oluşturma, sunma ve katılma.',
                   'endpoints': ['quiz_*', 'api_quiz_*']}),
    ('page_leaderboard', {'label': 'Yurt Sıralaması', 'icon': 'bi-trophy', 'desc': 'XP sıralaması ve rozetler.',
                          'endpoints': ['leaderboard_page']}),
    ('page_chat', {'label': 'Genel Sohbet', 'icon': 'bi-chat-dots', 'desc': 'Tüm yurdun ortak sohbet odası.',
                   'endpoints': ['chat']}),
    ('page_messages', {'label': 'Mesajlar', 'icon': 'bi-chat-left-text', 'desc': 'Birebir ve grup mesajlaşma.',
                       'endpoints': ['messages_*', 'group_*', 'conversation_*', 'message_*', 'pin_message']}),
    ('page_announcements', {'label': 'Duyurular', 'icon': 'bi-megaphone', 'desc': 'Yönetim duyurularını okuma.',
                            'endpoints': ['announcements_page']}),
    ('page_attendance', {'label': 'Yoklama', 'icon': 'bi-clipboard-check', 'desc': 'Yoklama listesi ve detayları (zorunlu yoklama her zaman açıktır).',
                         'endpoints': ['attendance_list', 'attendance_detail']}),
    ('page_maintenance', {'label': 'Arıza & Talep', 'icon': 'bi-tools', 'desc': 'Arıza bildirme ve talepleri takip etme.',
                          'endpoints': ['maintenance_page', 'create_maintenance']}),
    ('page_laundry', {'label': 'Çamaşırhane', 'icon': 'bi-basket', 'desc': 'Makine rezervasyonu ve sıra bekleme.',
                      'endpoints': ['laundry_*']}),
    ('page_roommates', {'label': 'Oda Arkadaşı', 'icon': 'bi-people', 'desc': 'Oda arkadaşı eşleştirme.',
                        'endpoints': ['roommate_matching']}),
    ('page_market', {'label': 'Yurt Pazarı', 'icon': 'bi-bag', 'desc': 'İkinci el alım-satım ilanları.',
                     'endpoints': ['market_*']}),
])

# Yönetim yetkileri: açıkça verilmedikçe KAPALI. 'full_admin' hepsini kapsar.
ADMIN_PERMISSIONS = OrderedDict([
    ('full_admin', {'label': 'Tam Yönetici', 'icon': 'bi-shield-shaded',
                    'desc': 'Tüm sayfalara ve tüm yönetim işlemlerine sınırsız erişim (aşağıdakilerin hepsini kapsar).'}),
    ('manage_roles', {'label': 'Rol & Yetki Yönetimi', 'icon': 'bi-person-gear',
                      'desc': 'Rol oluşturma, düzenleme, silme ve rollere yetki atama.'}),
    ('manage_users', {'label': 'Kullanıcı Yönetimi', 'icon': 'bi-people',
                      'desc': 'Kullanıcılara rol atama, yurt durumu, profil teması ve hesap silme.'}),
    ('manage_attendance', {'label': 'Yoklama Yönetimi', 'icon': 'bi-clipboard-check',
                           'desc': 'Yoklama açma/kapatma/silme ve otomatik yoklama saati.'}),
    ('manage_events', {'label': 'Etkinlik Yönetimi', 'icon': 'bi-calendar-check',
                       'desc': 'Herhangi bir etkinliği yönetme ve silme.'}),
    ('manage_menu', {'label': 'Menü & Yemekhane', 'icon': 'bi-cup-hot',
                     'desc': 'Excel menü yükleme, günlük menü düzenleme ve yemek saatleri.'}),
    ('manage_announcements', {'label': 'Duyuru Yönetimi', 'icon': 'bi-megaphone',
                              'desc': 'Duyuru yayınlama ve silme.'}),
    ('manage_polls', {'label': 'Anket Yönetimi', 'icon': 'bi-bar-chart-steps',
                      'desc': 'Ana ekrandaki yurt anketlerini açma ve kapatma.'}),
    ('manage_maintenance', {'label': 'Arıza & Talep Yönetimi', 'icon': 'bi-tools',
                            'desc': 'Arıza taleplerinin durumunu güncelleme ve yanıtlama.'}),
    ('manage_laundry', {'label': 'Çamaşırhane Yönetimi', 'icon': 'bi-basket',
                        'desc': 'Makine ekleme, bakım/servis dışı bırakma.'}),
    ('manage_quizzes', {'label': 'Yarışma Yönetimi', 'icon': 'bi-patch-question',
                        'desc': 'Herkesin yarışmalarını düzenleme, sunma ve silme.'}),
    ('moderate_content', {'label': 'İçerik Moderasyonu', 'icon': 'bi-shield-exclamation',
                          'desc': 'Başkalarının gönderi, sohbet mesajı ve pazar ilanlarını kaldırma.'}),
    ('view_analytics', {'label': 'Analitik & Raporlar', 'icon': 'bi-graph-up',
                        'desc': 'Yönetim panelindeki istatistik ve raporları görme.'}),
])

ALL_PERMISSION_KEYS = list(PAGE_PERMISSIONS.keys()) + list(ADMIN_PERMISSIONS.keys())

ROLE_COLOR_PALETTE = ['#b8860b', '#ef4444', '#f97316', '#eab308', '#10b981', '#0d9488', '#06b6d4',
                      '#3b82f6', '#6366f1', '#8b5cf6', '#d946ef', '#ec4899', '#64748b', '#6b7280']
ROLE_ICON_CHOICES = ['bi-shield-shaded', 'bi-shield-check', 'bi-person-badge', 'bi-award', 'bi-star',
                     'bi-tools', 'bi-building', 'bi-person', 'bi-cup-hot', 'bi-calendar-event', 'bi-megaphone',
                     'bi-basket', 'bi-heart-pulse', 'bi-mortarboard', 'bi-controller', 'bi-camera',
                     'bi-music-note-beamed', 'bi-flag', 'bi-lightning-charge', 'bi-gem', 'bi-key', 'bi-eye']


def _pages_all(value=True):
    return {k: value for k in PAGE_PERMISSIONS}


DEFAULT_ROLES_SEED = [
    {'key': 'yurt_muduru', 'title': 'Yurt Müdürü', 'badge': 'Yurt Müdürü', 'icon': 'bi-shield-shaded',
     'color': '#b8860b', 'sort_order': 1, 'desc': 'Yurt genel yöneticisi (tam yetki)',
     'permissions': {'full_admin': True}},
    {'key': 'yurt_mudur_yardimcisi', 'title': 'Yurt Müdür Yardımcısı', 'badge': 'Yurt Müdür Yrd.', 'icon': 'bi-shield-check',
     'color': '#f97316', 'sort_order': 2, 'desc': 'Yurt müdür yardımcısı (idari yetki)',
     'permissions': {'full_admin': True}},
    {'key': 'yurt_memuru', 'title': 'Yurt Memuru', 'badge': 'Yurt Memuru', 'icon': 'bi-person-badge',
     'color': '#3b82f6', 'sort_order': 3, 'desc': 'İdari işler ve kayıt sorumlusu',
     'permissions': {'full_admin': True}},
    {'key': 'yurt_temsilcisi', 'title': 'Yurt Temsilcisi', 'badge': 'Yurt Temsilcisi', 'icon': 'bi-award',
     'color': '#8b5cf6', 'sort_order': 4, 'desc': 'Öğrenci temsilcisi ve etkinlik koordinatörü',
     'permissions': {'full_admin': True}},
    {'key': 'yurt_gorevlisi', 'title': 'Yurt Görevlisi', 'badge': 'Yurt Görevlisi', 'icon': 'bi-tools',
     'color': '#eab308', 'sort_order': 5, 'desc': 'Arıza ve teknik taleplerden sorumlu',
     'permissions': {'manage_maintenance': True, 'manage_laundry': True}},
    {'key': 'kat_1_sorumlusu', 'title': '1. Kat Sorumlusu', 'badge': '1. Kat Sorumlusu', 'icon': 'bi-building',
     'color': '#10b981', 'sort_order': 6, 'desc': '1. kat düzeni ve yoklama sorumlusu',
     'permissions': {'manage_attendance': True}},
    {'key': 'kat_2_sorumlusu', 'title': '2. Kat Sorumlusu', 'badge': '2. Kat Sorumlusu', 'icon': 'bi-building',
     'color': '#10b981', 'sort_order': 7, 'desc': '2. kat düzeni ve yoklama sorumlusu',
     'permissions': {'manage_attendance': True}},
    {'key': 'kat_3_sorumlusu', 'title': '3. Kat Sorumlusu', 'badge': '3. Kat Sorumlusu', 'icon': 'bi-building',
     'color': '#10b981', 'sort_order': 8, 'desc': '3. kat düzeni ve yoklama sorumlusu',
     'permissions': {'manage_attendance': True}},
    {'key': 'kat_4_sorumlusu', 'title': '4. Kat Sorumlusu', 'badge': '4. Kat Sorumlusu', 'icon': 'bi-building',
     'color': '#10b981', 'sort_order': 9, 'desc': '4. kat düzeni ve yoklama sorumlusu',
     'permissions': {'manage_attendance': True}},
    # Uygulamanın sayfalarına göre hazır gelen yeni roller
    {'key': 'yemekhane_sorumlusu', 'title': 'Yemekhane Sorumlusu', 'badge': 'Yemekhane', 'icon': 'bi-cup-hot',
     'color': '#ef4444', 'sort_order': 20, 'desc': 'Menüleri yükler, günlük menüyü ve yemek saatlerini günceller',
     'permissions': {'manage_menu': True, 'manage_polls': True}},
    {'key': 'etkinlik_koordinatoru', 'title': 'Etkinlik Koordinatörü', 'badge': 'Etkinlik Koord.', 'icon': 'bi-calendar-event',
     'color': '#ec4899', 'sort_order': 21, 'desc': 'Etkinlik, duyuru, anket ve yarışmaları yönetir',
     'permissions': {'manage_events': True, 'manage_announcements': True, 'manage_polls': True, 'manage_quizzes': True}},
    {'key': 'camasirhane_sorumlusu', 'title': 'Çamaşırhane Sorumlusu', 'badge': 'Çamaşırhane', 'icon': 'bi-basket',
     'color': '#06b6d4', 'sort_order': 22, 'desc': 'Makineleri ve rezervasyon düzenini yönetir',
     'permissions': {'manage_laundry': True}},
    {'key': 'icerik_moderatoru', 'title': 'İçerik Moderatörü', 'badge': 'Moderatör', 'icon': 'bi-eye',
     'color': '#6366f1', 'sort_order': 23, 'desc': 'Akış, sohbet ve pazar içeriklerini denetler',
     'permissions': {'moderate_content': True, 'view_analytics': True}},
    {'key': 'ogrenci', 'title': 'Öğrenci', 'badge': 'Öğrenci', 'icon': 'bi-person',
     'color': '#6b7280', 'sort_order': 99, 'desc': 'Yurt sakini öğrenci (varsayılan rol)',
     'permissions': {}},
]
LOCKED_ROLE_KEYS = {'yurt_muduru', 'ogrenci'}  # silinemez; yurt_muduru'nun yetkileri kısıtlanamaz


def _role_row_to_dict(row):
    d = dict(row)
    try:
        d['permissions'] = json.loads(d.get('permissions') or '{}')
    except (ValueError, TypeError):
        d['permissions'] = {}
    d['class'] = 'role-badge-r-' + d['key']
    d['is_admin'] = role_has_admin_access(d)
    d['can_maintenance'] = role_perm(d, 'manage_maintenance')
    d['can_attendance'] = role_perm(d, 'manage_attendance')
    return d


def role_perm(role, perm_key):
    """Bir rol sözlüğünün belirli bir yetkiye sahip olup olmadığını döndürür."""
    if not role:
        return False
    perms = role.get('permissions') or {}
    if perms.get('full_admin'):
        return True
    if perm_key in PAGE_PERMISSIONS:
        return perms.get(perm_key, True) is not False
    return bool(perms.get(perm_key))


def role_has_admin_access(role):
    perms = (role or {}).get('permissions') or {}
    return any(bool(perms.get(k)) for k in ADMIN_PERMISSIONS)


USER_ROLES = OrderedDict()


def _seed_role_values(seed):
    return (seed['key'], seed['title'], seed['badge'], seed['icon'], seed['color'], seed['desc'],
            json.dumps(seed['permissions']), 1, seed['sort_order'])


def _init_roles_table(conn):
    conn.execute('''
        CREATE TABLE IF NOT EXISTS roles (
            key TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            badge TEXT NOT NULL,
            icon TEXT NOT NULL DEFAULT 'bi-person',
            color TEXT NOT NULL DEFAULT '#6b7280',
            desc TEXT,
            permissions TEXT NOT NULL DEFAULT '{}',
            is_system INTEGER DEFAULT 0,
            sort_order INTEGER DEFAULT 100,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    # Eksik varsayılan rolleri ekle; admin tarafından düzenlenmiş satırlara dokunma.
    for seed in DEFAULT_ROLES_SEED:
        conn.execute('''INSERT OR IGNORE INTO roles (key, title, badge, icon, color, desc, permissions, is_system, sort_order)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''', _seed_role_values(seed))
    # Yurt müdürü her zaman tam yetkili kalsın (kilitlenmeyi önler)
    conn.execute("UPDATE roles SET permissions = ? WHERE key = 'yurt_muduru'", (json.dumps({'full_admin': True}),))
    # Tanımsız role sahip kullanıcıları öğrenciye çek
    conn.execute("UPDATE users SET role = 'ogrenci' WHERE role IS NULL OR role = '' OR role NOT IN (SELECT key FROM roles)")
    conn.execute("UPDATE users SET role = 'yurt_muduru' WHERE id = 1")
    # users.is_admin sütununu rol yetkileriyle senkronize et
    for row in conn.execute('SELECT key, permissions FROM roles').fetchall():
        try:
            perms = json.loads(row['permissions'] or '{}')
        except (ValueError, TypeError):
            perms = {}
        flag = 1 if any(bool(perms.get(k)) for k in ADMIN_PERMISSIONS) else 0
        conn.execute('UPDATE users SET is_admin = ? WHERE role = ?', (flag, row['key']))


def load_user_roles_from_db():
    conn = get_db()
    try:
        rows = conn.execute('SELECT * FROM roles ORDER BY sort_order ASC, title ASC').fetchall()
    except sqlite3.OperationalError:
        rows = []
    conn.close()
    roles = OrderedDict()
    for row in rows:
        d = _role_row_to_dict(row)
        roles[d['key']] = d
    if not roles:
        for seed in DEFAULT_ROLES_SEED:
            d = dict(seed)
            d['permissions'] = dict(seed['permissions'])
            d['is_system'] = 1
            d['class'] = 'role-badge-r-' + d['key']
            d['is_admin'] = role_has_admin_access(d)
            roles[d['key']] = d
    return roles


def reload_user_roles():
    global USER_ROLES
    fresh = load_user_roles_from_db()
    USER_ROLES.clear()
    USER_ROLES.update(fresh)
    return USER_ROLES


def slugify_role_key(text):
    txt = normalize_tr_text(text or '').lower()
    txt = re.sub(r'[^a-z0-9]+', '_', txt).strip('_')
    return (txt or 'rol')[:40]


def resolve_role_key(user_dict):
    if not user_dict:
        return 'ogrenci'
    key = user_dict.get('role') or ('yurt_muduru' if user_dict.get('is_admin') else 'ogrenci')
    return key if key in USER_ROLES else 'ogrenci'


def user_has_perm(user_dict, perm_key):
    """Kullanıcının (rolü üzerinden) bir yetkiye sahip olup olmadığı."""
    if not user_dict:
        return False
    if user_dict.get('id') == 1:
        return True  # kurucu admin hesabı asla kilitlenmez
    return role_perm(USER_ROLES.get(resolve_role_key(user_dict)), perm_key)


def user_has_admin_access(user_dict):
    if not user_dict:
        return False
    if user_dict.get('id') == 1:
        return True
    return role_has_admin_access(USER_ROLES.get(resolve_role_key(user_dict)))


def is_admin_role(role_key):
    return role_has_admin_access(USER_ROLES.get(role_key))


def can_user_manage_maintenance(user_dict):
    return user_has_perm(user_dict, 'manage_maintenance')


def can_user_manage_attendance(user_dict):
    return user_has_perm(user_dict, 'manage_attendance')


def current_user_dict():
    """İstek başına bir kez veritabanından okunan oturum kullanıcısı (g içinde önbellekli)."""
    if 'user_id' not in session:
        return None
    cached = getattr(g, '_current_user_dict', None)
    if cached is not None:
        return cached
    conn = get_db()
    row = conn.execute('SELECT id, username, full_name, profile_photo, is_admin, role, room_number FROM users WHERE id = ?',
                       (session['user_id'],)).fetchone()
    conn.close()
    g._current_user_dict = dict(row) if row else {}
    return g._current_user_dict


def has_perm(perm_key):
    return user_has_perm(current_user_dict(), perm_key)


def page_permission_for_endpoint(endpoint):
    if not endpoint:
        return None
    for key, meta in PAGE_PERMISSIONS.items():
        for pat in meta['endpoints']:
            if pat.endswith('*'):
                if endpoint.startswith(pat[:-1]):
                    return key
            elif endpoint == pat:
                return key
    return None


def normalize_tr_text(text):
    if not text:
        return ""
    tr_map = {
        'İ': 'i', 'I': 'ı', 'ı': 'i', 'Ş': 's', 'ş': 's',
        'Ğ': 'g', 'ğ': 'g', 'Ü': 'u', 'ü': 'u', 'Ö': 'o',
        'ö': 'o', 'Ç': 'c', 'ç': 'c'
    }
    s = str(text)
    for k, v in tr_map.items():
        s = s.replace(k, v)
    return s.lower().strip()

POSTS_PER_PAGE = 12

# Oyunlar sekmesinde yer alan mini oyunlar
GAMES = {
    'snake': {
        'name': 'Yılan',
        'icon': 'bi-dice-5-fill',
        'desc': 'Klasik yılan oyunu — elmaları topla, kendine ya da duvara çarpma!',
        'sort': 'desc',
        'unit': 'puan',
    },
    'reflex': {
        'name': 'Refleks Testi',
        'icon': 'bi-lightning-charge-fill',
        'desc': 'Ekran yeşile dönünce en hızlı sen tıkla, reaksiyon sürenle yarış!',
        'sort': 'asc',
        'unit': 'ms',
    },
    'whackamole': {
        'name': 'Köstebek Vurma',
        'icon': 'bi-hammer',
        'desc': '30 saniyede çıkan köstebekleri yakala, en yüksek skoru sen yap!',
        'sort': 'desc',
        'unit': 'puan',
    },
    'penalty': {'name':'Penaltı Shootout','icon':'bi-dribbble','desc':'10 penaltıda kaleciyi geç ve yurttaki en yüksek skoru yap!','sort':'desc','unit':'gol'},
    'memory': {'name':'Hafıza Kartları','icon':'bi-grid-3x3-gap','desc':'Kartları eşleştir, süreye karşı hafızanı test et!','sort':'asc','unit':'sn'},
    '2048': {'name':'2048','icon':'bi-grid-3x3','desc':'Sayıları birleştir, 2048 hedefine ulaş!','sort':'desc','unit':'puan'},
}



def init_db():
    conn = get_db()
    c = conn.cursor()

    # Kullanıcılar
    c.execute('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            full_name TEXT NOT NULL,
            room_number TEXT,
            is_admin INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # Etkinlikler
    c.execute('''
        CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            event_type TEXT NOT NULL,
            location TEXT NOT NULL,
            event_date TEXT NOT NULL,
            max_participants INTEGER,
            image_path TEXT,
            created_by INTEGER,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS event_participants (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_id INTEGER,
            user_id INTEGER,
            joined_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (event_id) REFERENCES events (id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            UNIQUE(event_id, user_id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            message TEXT NOT NULL,
            room TEXT DEFAULT 'genel',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # YOKLAMA TABLOLARI
    c.execute('''
        CREATE TABLE IF NOT EXISTS attendances (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            attendance_date TEXT NOT NULL,
            is_active INTEGER DEFAULT 1,
            is_mandatory INTEGER DEFAULT 0,
            created_by INTEGER,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS attendance_records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            attendance_id INTEGER,
            user_id INTEGER,
            status TEXT NOT NULL CHECK(status IN ('present', 'absent')),
            marked_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (attendance_id) REFERENCES attendances (id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            UNIQUE(attendance_id, user_id)
        )
    ''')

    # YÖNETİCİ AYARLARI — otomatik yoklama saati vb.
    c.execute('''
        CREATE TABLE IF NOT EXISTS app_settings (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('attendance_auto_enabled', '1')")
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('attendance_auto_time', '22:00')")
    # YEMEKHANE SAATLERİ — Türkiye'deki yurt varsayılanları. Yönetici panelinden değiştirilebilir.
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('meal_breakfast_weekday_start', '07:00')")
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('meal_breakfast_weekday_end', '09:00')")
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('meal_breakfast_weekend_start', '09:00')")
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('meal_breakfast_weekend_end', '11:00')")
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('meal_dinner_start', '18:00')")
    c.execute("INSERT OR IGNORE INTO app_settings(key, value) VALUES ('meal_dinner_end', '20:00')")

    # BİLDİRİM TABLOSU
    c.execute('''
        CREATE TABLE IF NOT EXISTS notifications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            title TEXT NOT NULL,
            message TEXT NOT NULL,
            link TEXT,
            is_read INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')

    # SOSYAL PAYLAŞIM (Instagram tarzı akış) — görsel, video ya da sadece metin
    _migrate_posts_allow_text(conn)
    c.execute('''
        CREATE TABLE IF NOT EXISTS posts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            caption TEXT,
            media_path TEXT,
            media_type TEXT NOT NULL CHECK(media_type IN ('image', 'video', 'text')),
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS post_likes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            post_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (post_id) REFERENCES posts (id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            UNIQUE(post_id, user_id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS post_comments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            post_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            comment TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (post_id) REFERENCES posts (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')

    _ensure_column(conn, 'events', 'image_path', 'TEXT')

    # WEB PUSH ABONELİKLERİ
    c.execute('''
        CREATE TABLE IF NOT EXISTS push_subscriptions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            endpoint TEXT NOT NULL UNIQUE,
            p256dh TEXT NOT NULL,
            auth TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        )
    ''')

    # ÖZEL MESAJLAR (DM)
    c.execute('''
        CREATE TABLE IF NOT EXISTS direct_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sender_id INTEGER NOT NULL,
            receiver_id INTEGER NOT NULL,
            message TEXT NOT NULL,
            is_read INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (sender_id) REFERENCES users (id),
            FOREIGN KEY (receiver_id) REFERENCES users (id)
        )
    ''')

    # ==================== MODERN MESAJLAŞMA SİSTEMİ (DM + GRUP) ====================
    # Eski 'direct_messages' tablosu korunur (geçmiş veri/uyumluluk için); yeni
    # sistem 'conversations' tabanlı çalışır ve hem birebir hem grup sohbetlerini
    # tek bir yapıda yönetir. Var olan veriler _migrate_direct_messages_to_conversations
    # ile kaybolmadan yeni şemaya aktarılır.
    c.execute('''
        CREATE TABLE IF NOT EXISTS conversations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            type TEXT NOT NULL CHECK(type IN ('dm', 'group')),
            name TEXT,
            description TEXT,
            photo TEXT,
            created_by INTEGER,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            last_message_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users (id)
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS conversation_members (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            role TEXT NOT NULL DEFAULT 'member' CHECK(role IN ('member', 'admin')),
            joined_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            last_read_message_id INTEGER NOT NULL DEFAULT 0,
            is_pinned INTEGER DEFAULT 0,
            is_muted INTEGER DEFAULT 0,
            is_archived INTEGER DEFAULT 0,
            hidden_at TIMESTAMP,
            left_at TIMESTAMP,
            UNIQUE(conversation_id, user_id),
            FOREIGN KEY (conversation_id) REFERENCES conversations (id) ON DELETE CASCADE,
            FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE CASCADE
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS conv_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            sender_id INTEGER NOT NULL,
            content TEXT,
            reply_to_id INTEGER,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            edited_at TIMESTAMP,
            deleted_for_everyone INTEGER DEFAULT 0,
            deleted_at TIMESTAMP,
            is_system INTEGER DEFAULT 0,
            FOREIGN KEY (conversation_id) REFERENCES conversations (id) ON DELETE CASCADE,
            FOREIGN KEY (sender_id) REFERENCES users (id),
            FOREIGN KEY (reply_to_id) REFERENCES conv_messages (id)
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS message_hides (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            message_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(message_id, user_id),
            FOREIGN KEY (message_id) REFERENCES conv_messages (id) ON DELETE CASCADE,
            FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE CASCADE
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS message_pins (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id INTEGER NOT NULL,
            message_id INTEGER NOT NULL,
            pinned_by INTEGER,
            pinned_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(conversation_id, message_id),
            FOREIGN KEY (conversation_id) REFERENCES conversations (id) ON DELETE CASCADE,
            FOREIGN KEY (message_id) REFERENCES conv_messages (id) ON DELETE CASCADE
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS blocked_users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            blocker_id INTEGER NOT NULL,
            blocked_id INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(blocker_id, blocked_id),
            FOREIGN KEY (blocker_id) REFERENCES users (id) ON DELETE CASCADE,
            FOREIGN KEY (blocked_id) REFERENCES users (id) ON DELETE CASCADE
        )
    ''')
    c.execute('''
        CREATE TABLE IF NOT EXISTS message_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            message_id INTEGER,
            conversation_id INTEGER,
            reported_by INTEGER NOT NULL,
            reason TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (reported_by) REFERENCES users (id)
        )
    ''')
    c.execute('CREATE INDEX IF NOT EXISTS idx_convmembers_user ON conversation_members(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_convmembers_conv ON conversation_members(conversation_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_convmsg_conv ON conv_messages(conversation_id, id)')
    _ensure_column(conn, 'conv_messages', 'image', 'TEXT')
    c.execute('CREATE INDEX IF NOT EXISTS idx_blocked_pair ON blocked_users(blocker_id, blocked_id)')
    conn.commit()
    _migrate_direct_messages_to_conversations(conn)

    # TAKİP SİSTEMİ
    c.execute('''
        CREATE TABLE IF NOT EXISTS follows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            follower_id INTEGER NOT NULL,
            followed_id INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (follower_id) REFERENCES users (id),
            FOREIGN KEY (followed_id) REFERENCES users (id),
            UNIQUE(follower_id, followed_id)
        )
    ''')

    # HİKAYELER (24 saat sonra kaybolan paylaşımlar)
    c.execute('''
        CREATE TABLE IF NOT EXISTS stories (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            media_path TEXT NOT NULL,
            media_type TEXT NOT NULL CHECK(media_type IN ('image', 'video')),
            caption TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            expires_at TIMESTAMP NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS story_views (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            story_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            viewed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (story_id) REFERENCES stories (id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            UNIQUE(story_id, user_id)
        )
    ''')

    # ==================== KAHOOT TARZI YARIŞMA (QUIZ) SİSTEMİ ====================
    c.execute('''
        CREATE TABLE IF NOT EXISTS quizzes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            created_by INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS quiz_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            quiz_id INTEGER NOT NULL,
            question_text TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT,
            option_d TEXT,
            correct_option TEXT NOT NULL CHECK(correct_option IN ('a','b','c','d')),
            time_limit INTEGER DEFAULT 20,
            order_index INTEGER DEFAULT 0,
            FOREIGN KEY (quiz_id) REFERENCES quizzes (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS quiz_sessions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            quiz_id INTEGER NOT NULL,
            host_id INTEGER NOT NULL,
            code TEXT UNIQUE NOT NULL,
            status TEXT NOT NULL DEFAULT 'waiting',
            current_index INTEGER DEFAULT -1,
            question_started_at TIMESTAMP,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (quiz_id) REFERENCES quizzes (id),
            FOREIGN KEY (host_id) REFERENCES users (id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS quiz_participants (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            nickname TEXT NOT NULL,
            score INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (session_id) REFERENCES quiz_sessions (id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            UNIQUE(session_id, user_id)
        )
    ''')

    c.execute('''
        CREATE TABLE IF NOT EXISTS quiz_answers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id INTEGER NOT NULL,
            question_id INTEGER NOT NULL,
            participant_id INTEGER NOT NULL,
            selected_option TEXT,
            is_correct INTEGER DEFAULT 0,
            answer_ms INTEGER,
            points INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (session_id) REFERENCES quiz_sessions (id),
            FOREIGN KEY (question_id) REFERENCES quiz_questions (id),
            FOREIGN KEY (participant_id) REFERENCES quiz_participants (id),
            UNIQUE(question_id, participant_id)
        )
    ''')

    # ==================== OYUNLAR (Mini oyunlar & skor tablosu) ====================
    c.execute('''
        CREATE TABLE IF NOT EXISTS game_scores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            game_key TEXT NOT NULL,
            user_id INTEGER NOT NULL,
            score INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
    ''')

    # ==================== YENİ MODÜLLER ====================
    c.execute("""CREATE TABLE IF NOT EXISTS meal_menus (id INTEGER PRIMARY KEY AUTOINCREMENT, menu_date TEXT NOT NULL UNIQUE, breakfast TEXT DEFAULT '', lunch TEXT DEFAULT '', dinner TEXT DEFAULT '', snack TEXT DEFAULT '', source_file TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS menu_ratings (id INTEGER PRIMARY KEY AUTOINCREMENT, menu_date TEXT NOT NULL, user_id INTEGER NOT NULL, breakfast_rating INTEGER, dinner_rating INTEGER, overall_rating INTEGER NOT NULL, comment TEXT DEFAULT '', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, UNIQUE(menu_date,user_id), FOREIGN KEY(user_id) REFERENCES users(id))""")
    c.execute("""CREATE TABLE IF NOT EXISTS announcements (id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, content TEXT NOT NULL, priority TEXT DEFAULT 'normal', expires_at TEXT, created_by INTEGER, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(created_by) REFERENCES users(id))""")
    c.execute("""CREATE TABLE IF NOT EXISTS maintenance_requests (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, category TEXT DEFAULT 'Diğer', room_number TEXT, description TEXT NOT NULL, photo_path TEXT, priority TEXT DEFAULT 'normal', status TEXT DEFAULT 'pending', admin_note TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES users(id))""")
    c.execute("""CREATE TABLE IF NOT EXISTS marketplace_items (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, title TEXT NOT NULL, description TEXT, price REAL, category TEXT DEFAULT 'Diğer', status TEXT DEFAULT 'active', created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES users(id))""")
    c.execute("""CREATE TABLE IF NOT EXISTS xp_transactions (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, amount INTEGER NOT NULL, reason TEXT, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, FOREIGN KEY(user_id) REFERENCES users(id))""")
    c.execute("""CREATE TABLE IF NOT EXISTS badges (id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT UNIQUE NOT NULL, name TEXT NOT NULL, description TEXT, icon TEXT DEFAULT '🏆')""")
    c.execute("""CREATE TABLE IF NOT EXISTS user_badges (id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL, badge_id INTEGER NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, UNIQUE(user_id,badge_id), FOREIGN KEY(user_id) REFERENCES users(id), FOREIGN KEY(badge_id) REFERENCES badges(id))""")
    c.execute("""CREATE TABLE IF NOT EXISTS daily_challenges (id INTEGER PRIMARY KEY AUTOINCREMENT, challenge_date TEXT UNIQUE NOT NULL, title TEXT NOT NULL, description TEXT, target INTEGER DEFAULT 1, xp_reward INTEGER DEFAULT 50, action TEXT DEFAULT 'game')""")
    c.execute("""CREATE TABLE IF NOT EXISTS challenge_completions (id INTEGER PRIMARY KEY AUTOINCREMENT, challenge_id INTEGER NOT NULL, user_id INTEGER NOT NULL, progress INTEGER DEFAULT 0, completed INTEGER DEFAULT 0, completed_at TIMESTAMP, UNIQUE(challenge_id,user_id), FOREIGN KEY(challenge_id) REFERENCES daily_challenges(id), FOREIGN KEY(user_id) REFERENCES users(id))""")
    c.execute("INSERT OR IGNORE INTO badges(code,name,description,icon) VALUES ('legend','Yurt Efsanesi','500 XP ve üzeri','👑')")

    # ANKETLER ("Bugün Yurtta" ekranı için)
    c.execute("""CREATE TABLE IF NOT EXISTS polls (id INTEGER PRIMARY KEY AUTOINCREMENT, question TEXT NOT NULL, is_active INTEGER DEFAULT 1, expires_at TEXT, created_by INTEGER, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)""")
    c.execute("""CREATE TABLE IF NOT EXISTS poll_options (id INTEGER PRIMARY KEY AUTOINCREMENT, poll_id INTEGER NOT NULL, label TEXT NOT NULL, FOREIGN KEY(poll_id) REFERENCES polls(id) ON DELETE CASCADE)""")
    c.execute("""CREATE TABLE IF NOT EXISTS poll_votes (id INTEGER PRIMARY KEY AUTOINCREMENT, poll_id INTEGER NOT NULL, option_id INTEGER NOT NULL, user_id INTEGER NOT NULL, created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP, UNIQUE(poll_id,user_id), FOREIGN KEY(poll_id) REFERENCES polls(id) ON DELETE CASCADE, FOREIGN KEY(option_id) REFERENCES poll_options(id) ON DELETE CASCADE)""")
    c.execute('CREATE INDEX IF NOT EXISTS idx_poll_votes_poll ON poll_votes(poll_id)')

    # Eski veritabanlarına yeni sütunları güvenle ekle
    _ensure_column(conn, 'users', 'profile_photo', "TEXT")
    _ensure_column(conn, 'users', 'last_seen', "TEXT")
    _ensure_column(conn, 'users', 'bio', "TEXT")
    _ensure_column(conn, 'users', 'in_dorm', "INTEGER DEFAULT 1")
    _ensure_column(conn, 'users', 'profile_theme', "TEXT DEFAULT 'default'")
    _ensure_column(conn, 'attendances', 'is_mandatory', "INTEGER DEFAULT 0")
    _ensure_column(conn, 'post_comments', 'parent_id', "INTEGER")
    _ensure_column(conn, 'stories', 'text_overlays', "TEXT")
    _ensure_column(conn, 'quiz_questions', 'image_path', "TEXT")
    _ensure_column(conn, 'quiz_questions', 'audio_path', "TEXT")
    _ensure_column(conn, 'quiz_questions', 'option_a_image', "TEXT")
    _ensure_column(conn, 'quiz_questions', 'option_b_image', "TEXT")
    _ensure_column(conn, 'quiz_questions', 'option_c_image', "TEXT")
    _ensure_column(conn, 'quiz_questions', 'option_d_image', "TEXT")

    # ------------------------------------------------------------------
    # YENİ MODÜLLER İÇİN GERİYE DÖNÜK UYUMLU MIGRATION
    # ------------------------------------------------------------------
    # Kullanıcı eski database.db ile yeni sürümü çalıştırdığında
    # CREATE TABLE IF NOT EXISTS mevcut tabloyu değiştirmez. Bu nedenle
    # yeni sürümün kullandığı kolonları ayrıca garanti ediyoruz.
    # Böylece "no such column" hataları oluşmaz ve mevcut veriler korunur.
    new_schema = {
        'meal_menus': {
            'menu_date': "TEXT", 'breakfast': "TEXT DEFAULT ''",
            'lunch': "TEXT DEFAULT ''", 'dinner': "TEXT DEFAULT ''",
            'snack': "TEXT DEFAULT ''", 'source_file': "TEXT",
            'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'menu_ratings': {
            'menu_date': "TEXT", 'user_id': "INTEGER", 'breakfast_rating': "INTEGER",
            'dinner_rating': "INTEGER", 'overall_rating': "INTEGER DEFAULT 0", 'comment': "TEXT DEFAULT ''",
            'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP", 'updated_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'announcements': {
            'title': "TEXT", 'content': "TEXT",
            'priority': "TEXT DEFAULT 'normal'", 'expires_at': "TEXT",
            'created_by': "INTEGER", 'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'maintenance_requests': {
            'user_id': "INTEGER", 'category': "TEXT DEFAULT 'Diğer'",
            'room_number': "TEXT", 'description': "TEXT", 'photo_path': "TEXT",
            'priority': "TEXT DEFAULT 'normal'", 'status': "TEXT DEFAULT 'pending'",
            'admin_note': "TEXT", 'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP",
            'updated_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'marketplace_items': {
            'user_id': "INTEGER", 'title': "TEXT", 'description': "TEXT",
            'price': "REAL", 'category': "TEXT DEFAULT 'Diğer'",
            'status': "TEXT DEFAULT 'active'", 'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'xp_transactions': {
            'user_id': "INTEGER", 'amount': "INTEGER", 'reason': "TEXT",
            'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'badges': {
            'code': "TEXT", 'name': "TEXT", 'description': "TEXT",
            'icon': "TEXT DEFAULT '🏆'"
        },
        'user_badges': {
            'user_id': "INTEGER", 'badge_id': "INTEGER",
            'created_at': "TIMESTAMP DEFAULT CURRENT_TIMESTAMP"
        },
        'daily_challenges': {
            'challenge_date': "TEXT", 'title': "TEXT", 'description': "TEXT",
            'target': "INTEGER DEFAULT 1", 'xp_reward': "INTEGER DEFAULT 50",
            'action': "TEXT DEFAULT 'game'"
        },
        'challenge_completions': {
            'challenge_id': "INTEGER", 'user_id': "INTEGER",
            'progress': "INTEGER DEFAULT 0", 'completed': "INTEGER DEFAULT 0",
            'completed_at': "TIMESTAMP"
        }
    }
    for table, columns in new_schema.items():
        for column, definition in columns.items():
            _ensure_column(conn, table, column, definition)

    # Bazı eski veritabanlarında bu tablolar hiç yoksa, yukarıdaki CREATE
    # TABLE IF NOT EXISTS ile zaten oluşturulmuş olurlar. Benzersizlik/arama
    # için gerekli indeksleri de mevcut DB üzerinde güvenle oluştur.
    c.execute('CREATE INDEX IF NOT EXISTS idx_meal_menus_date ON meal_menus(menu_date)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_menu_ratings_date ON menu_ratings(menu_date)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_menu_ratings_user ON menu_ratings(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_announcements_created ON announcements(created_at DESC)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_announcements_expiry ON announcements(expires_at)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_maintenance_user ON maintenance_requests(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_maintenance_status ON maintenance_requests(status)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_marketplace_status ON marketplace_items(status)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_marketplace_user ON marketplace_items(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_xp_user ON xp_transactions(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_xp_created ON xp_transactions(created_at DESC)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_user_badges_user ON user_badges(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_challenges_date ON daily_challenges(challenge_date)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_challenge_completions_user ON challenge_completions(user_id)')

    # PERFORMANS: 150-200+ kullanıcı ölçeğinde sık sorgulanan sütunlara indeks.
    # Bunlar olmadan her sorgu tabloyu baştan sona tarar; kullanıcı sayısı ve
    # gönderi/mesaj/bildirim hacmi arttıkça bu ciddi yavaşlamaya yol açar.
    c.execute('CREATE INDEX IF NOT EXISTS idx_attendance_records_attendance ON attendance_records(attendance_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_attendance_records_user ON attendance_records(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_event_participants_event ON event_participants(event_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_event_participants_user ON event_participants(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_posts_user ON posts(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_posts_created ON posts(created_at DESC)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_post_likes_post ON post_likes(post_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_post_likes_post_user ON post_likes(post_id, user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_post_comments_post ON post_comments(post_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_notifications_user_read ON notifications(user_id, is_read)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_dm_pair ON direct_messages(sender_id, receiver_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_dm_receiver_read ON direct_messages(receiver_id, is_read)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_push_subscriptions_user ON push_subscriptions(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_users_in_dorm ON users(in_dorm)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_follows_follower ON follows(follower_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_follows_followed ON follows(followed_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_stories_user ON stories(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_stories_expires ON stories(expires_at)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_story_views_story ON story_views(story_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_post_comments_parent ON post_comments(parent_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_quiz_questions_quiz ON quiz_questions(quiz_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_quiz_sessions_code ON quiz_sessions(code)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_quiz_participants_session ON quiz_participants(session_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_quiz_answers_session ON quiz_answers(session_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_game_scores_game ON game_scores(game_key)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_game_scores_user ON game_scores(user_id)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_laundry_res_machine_time ON laundry_reservations(machine_id, start_at, end_at, status)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_laundry_res_user ON laundry_reservations(user_id, start_at)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_laundry_waitlist_machine ON laundry_waitlist(machine_id, notified_at)')

    # ODA ARKADAŞI EŞLEŞTİRME
    c.execute('''
        CREATE TABLE IF NOT EXISTS roommate_profiles (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL UNIQUE,
            preferences_json TEXT NOT NULL DEFAULT '{}',
            completed INTEGER DEFAULT 0,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(user_id) REFERENCES users(id) ON DELETE CASCADE
        )
    ''')
    c.execute('CREATE INDEX IF NOT EXISTS idx_roommate_profiles_completed ON roommate_profiles(completed)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_roommate_profiles_user ON roommate_profiles(user_id)')

    # Maintenance requests tablosundaki kısıtlayıcı eski CHECK constraint temizliği
    try:
        schema_row = conn.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='maintenance_requests'").fetchone()
        if schema_row and ('CHECK(status IN' in schema_row['sql'] or 'CHECK(category IN' in schema_row['sql'] or 'urgency' in schema_row['sql']):
            conn.execute('ALTER TABLE maintenance_requests RENAME TO maintenance_requests_old')
            conn.execute('''
                CREATE TABLE maintenance_requests (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL,
                    category TEXT DEFAULT 'Diğer',
                    room_number TEXT,
                    description TEXT NOT NULL,
                    photo_path TEXT,
                    priority TEXT DEFAULT 'normal',
                    status TEXT DEFAULT 'pending',
                    admin_note TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(user_id) REFERENCES users(id)
                )
            ''')
            old_rows = conn.execute('SELECT * FROM maintenance_requests_old').fetchall()
            for row in old_rows:
                old_status = row['status']
                if old_status in ['Bekliyor', 'bekliyor']:
                    new_status = 'pending'
                elif old_status in ['İnceleniyor', 'inceleniyor', 'islemde', 'in_progress']:
                    new_status = 'in_progress'
                elif old_status in ['Çözüldü', 'cozuldu', 'tamamlandi', 'resolved']:
                    new_status = 'resolved'
                elif old_status in ['İptal', 'iptal', 'cancelled']:
                    new_status = 'cancelled'
                else:
                    new_status = old_status or 'pending'

                row_keys = row.keys()
                old_priority = row['urgency'] if 'urgency' in row_keys else (row['priority'] if 'priority' in row_keys else 'normal')
                if old_priority in ['Yüksek', 'Acil', 'urgent']:
                    new_prio = 'urgent'
                elif old_priority in ['Orta', 'high']:
                    new_prio = 'high'
                else:
                    new_prio = 'normal'

                r_room = row['room_number'] if 'room_number' in row_keys else ''
                r_cat = row['category'] if 'category' in row_keys else 'Diğer'
                r_desc = row['description'] if 'description' in row_keys else ''
                r_photo = row['photo_path'] if 'photo_path' in row_keys else None
                r_note = row['admin_note'] if 'admin_note' in row_keys else ''
                r_created = row['created_at'] if 'created_at' in row_keys else ''
                r_updated = row['updated_at'] if 'updated_at' in row_keys else ''

                conn.execute('''
                    INSERT INTO maintenance_requests (id, user_id, category, room_number, description, photo_path, priority, status, admin_note, created_at, updated_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (row['id'], row['user_id'], r_cat, r_room, r_desc, r_photo, new_prio, new_status, r_note, r_created, r_updated))
            conn.execute('DROP TABLE maintenance_requests_old')
    except Exception as e:
        print("Maintenance migration note:", e)

    # Kullanıcı rolü sütunu ve varsayılan yönetici
    try:
        c.execute("ALTER TABLE users ADD COLUMN role TEXT DEFAULT 'ogrenci'")
    except Exception:
        pass
    c.execute('CREATE INDEX IF NOT EXISTS idx_users_role ON users(role)')

    # Varsayılan admin
    admin_hash = generate_password_hash('admin123')
    c.execute('''
        INSERT OR IGNORE INTO users (id, username, password, full_name, is_admin, in_dorm, role)
        VALUES (1, 'admin', ?, 'Yurt Müdürü', 1, 0, 'yurt_muduru')
    ''', (admin_hash,))
    # Admin rolünü güncelle
    c.execute("UPDATE users SET role = 'yurt_muduru', is_admin = 1 WHERE id = 1 AND (role IS NULL OR role = 'ogrenci' OR role = 'default')")

    _laundry_cleanup_overlaps(conn)
    _repair_misclassified_menus(conn)
    _init_roles_table(conn)
    _init_quiz_extras(conn)

    conn.commit()
    conn.close()
    reload_user_roles()


# ==================== YARDIMCI FONKSİYONLAR ====================

def normalize_existing_usernames():
    """Eski hesapların kullanıcı adlarını da küçük harfe getirir."""
    conn = get_db()
    rows = conn.execute('SELECT id, username FROM users').fetchall()
    seen = set()
    for row in rows:
        username = (row['username'] or '').strip().lower()
        if username and username not in seen and username != row['username']:
            conn.execute('UPDATE users SET username = ? WHERE id = ?', (username, row['id']))
        seen.add(username)
    conn.commit()
    conn.close()

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash('Lütfen önce giriş yapın.', 'warning')
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def _wants_json():
    return (request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
            or request.path.startswith('/api/'))


def _deny(message):
    if _wants_json():
        return jsonify({'ok': False, 'error': message}), 403
    flash(message, 'danger')
    return redirect(url_for('dashboard'))


def admin_required(f):
    """Yönetim paneline giriş: kullanıcının rolünde EN AZ BİR yönetim yetkisi olmalı.
    Panel içindeki sekme ve işlemler ayrıca kendi yetkileriyle korunur."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if not user_has_admin_access(current_user_dict()):
            return _deny('Bu sayfaya erişim yetkiniz yok!')
        return f(*args, **kwargs)
    return decorated_function


def permission_required(*perm_keys):
    """Verilen yetkilerden en az birine sahip olmayı zorunlu kılar."""
    def wrapper(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'user_id' not in session:
                return redirect(url_for('login'))
            me = current_user_dict()
            if not any(user_has_perm(me, p) for p in perm_keys):
                labels = ', '.join(ADMIN_PERMISSIONS.get(p, PAGE_PERMISSIONS.get(p, {})).get('label', p) for p in perm_keys)
                return _deny(f'Bu işlem için "{labels}" yetkisi gerekiyor.')
            return f(*args, **kwargs)
        return decorated_function
    return wrapper


def maintenance_manager_required(f):
    return permission_required('manage_maintenance')(f)


def attendance_manager_required(f):
    return permission_required('manage_attendance')(f)


@app.before_request
def enforce_page_permissions():
    """Rolün erişimi kapatılmış sayfalarına girişi engeller ve oturumdaki
    is_admin bayrağını rolün güncel yetkileriyle senkron tutar."""
    if 'user_id' not in session or request.endpoint in (None, 'static'):
        return
    me = current_user_dict()
    if not me:
        session.clear()
        return redirect(url_for('login'))
    session['is_admin'] = user_has_admin_access(me)
    page_key = page_permission_for_endpoint(request.endpoint)
    if page_key and not user_has_perm(me, page_key):
        label = PAGE_PERMISSIONS[page_key]['label']
        msg = f'Rolünüzün "{label}" sayfasına erişim yetkisi yok.'
        if _wants_json():
            return jsonify({'ok': False, 'error': msg}), 403
        flash(msg, 'warning')
        return redirect(url_for('dashboard'))


def _push_to_subscription(subscription, title, message, link=None):
    if webpush is None or not os.path.exists(VAPID_PRIVATE_FILE):
        return False
    payload = json.dumps({'title': title, 'body': message, 'url': link or '/', 'icon': '/static/img/icons/icon-192.png'}, ensure_ascii=False)
    info = {'endpoint': subscription['endpoint'], 'keys': {'p256dh': subscription['p256dh'], 'auth': subscription['auth']}}
    try:
        webpush(subscription_info=info, data=payload, vapid_private_key=VAPID_PRIVATE_FILE, vapid_claims={'sub': VAPID_SUBJECT})
        return True
    except Exception as exc:
        status = getattr(getattr(exc, 'response', None), 'status_code', None)
        if status in (404, 410):
            try:
                conn = get_db(); conn.execute('DELETE FROM push_subscriptions WHERE endpoint = ?', (subscription['endpoint'],)); conn.commit(); conn.close()
            except Exception:
                pass
        else:
            print(f'[WEB PUSH] gönderilemedi: {exc}')
        return False

def push_notification_to_user(user_id, title, message, link=None):
    if webpush is None:
        return
    conn = get_db()
    rows = conn.execute('SELECT id, endpoint, p256dh, auth FROM push_subscriptions WHERE user_id = ?', (user_id,)).fetchall()
    conn.close()
    for row in rows:
        _push_to_subscription(row, title, message, link)

def push_notification_to_all(user_ids, title, message, link=None):
    if webpush is None or not user_ids:
        return
    conn = get_db()
    marks = ','.join('?' for _ in user_ids)
    rows = conn.execute(f'SELECT id, endpoint, p256dh, auth FROM push_subscriptions WHERE user_id IN ({marks})', tuple(user_ids)).fetchall()
    conn.close()
    for row in rows:
        _push_to_subscription(row, title, message, link)

def create_notification(user_id, title, message, link=None, push=True):
    """Uygulama içi bildirimi oluşturur; push=False ise telefon bildirim paneline göndermez."""
    conn = get_db()
    conn.execute('INSERT INTO notifications (user_id, title, message, link) VALUES (?, ?, ?, ?)', (user_id, title, message, link))
    conn.commit(); conn.close()
    socketio.emit('new_notification', {'title': title, 'message': message, 'link': link or '/', 'time': now_tr().strftime('%H:%M'), 'browser_push': bool(push)}, room=f'user_{user_id}')
    if push:
        socketio.start_background_task(push_notification_to_user, user_id, title, message, link)

def create_notification_for_all(title, message, link=None, only_in_dorm=False):
    conn = get_db()
    users = conn.execute('SELECT id FROM users WHERE in_dorm = 1').fetchall() if only_in_dorm else conn.execute('SELECT id FROM users').fetchall()
    ids = [u['id'] for u in users]
    conn.executemany('INSERT INTO notifications (user_id, title, message, link) VALUES (?, ?, ?, ?)', [(uid, title, message, link) for uid in ids])
    conn.commit(); conn.close()
    for uid in ids:
        socketio.emit('new_notification', {'title': title, 'message': message, 'link': link or '/', 'time': now_tr().strftime('%H:%M')}, room=f'user_{uid}')
    socketio.start_background_task(push_notification_to_all, ids, title, message, link)


def get_unread_count(user_id):
    conn = get_db()
    count = conn.execute('''
        SELECT COUNT(*) as c FROM notifications
        WHERE user_id = ? AND is_read = 0
    ''', (user_id,)).fetchone()['c']
    conn.close()
    return count


def get_unread_dm_count(user_id):
    """Yeni conversations şemasına göre toplam okunmamış mesaj sayısı
    (DM + grup, sessize alınmış/arşivlenmiş sohbetler hariç değil - rozet
    her zaman gerçek toplamı yansıtır)."""
    conn = get_db()
    count = conn.execute('''
        SELECT COUNT(*) as c
        FROM conv_messages m
        JOIN conversation_members cm ON cm.conversation_id = m.conversation_id
        WHERE cm.user_id = ?
          AND cm.hidden_at IS NULL
          AND cm.left_at IS NULL
          AND m.sender_id != ?
          AND m.id > cm.last_read_message_id
          AND NOT (m.deleted_for_everyone = 1)
          AND m.id NOT IN (SELECT message_id FROM message_hides WHERE user_id = ?)
    ''', (user_id, user_id, user_id)).fetchone()['c']
    conn.close()
    return count


def is_blocked_between(conn, user_a, user_b):
    """user_a ile user_b arasında herhangi bir yönde engelleme var mı?"""
    row = conn.execute('''
        SELECT 1 FROM blocked_users
        WHERE (blocker_id = ? AND blocked_id = ?) OR (blocker_id = ? AND blocked_id = ?)
        LIMIT 1
    ''', (user_a, user_b, user_b, user_a)).fetchone()
    return row is not None


def get_or_create_dm_conversation(conn, user_a, user_b):
    """İki kullanıcı arasındaki DM sohbetini bulur, yoksa oluşturur.
    Bir kullanıcı sohbeti kendi tarafından 'silmiş' olsa bile (hidden_at)
    sohbet tekrar bulunur; yalnızca kendi hidden_at bayrağı sıfırlanır."""
    row = conn.execute('''
        SELECT c.id FROM conversations c
        JOIN conversation_members m1 ON m1.conversation_id = c.id AND m1.user_id = ?
        JOIN conversation_members m2 ON m2.conversation_id = c.id AND m2.user_id = ?
        WHERE c.type = 'dm'
        LIMIT 1
    ''', (user_a, user_b)).fetchone()
    if row:
        return row['id']
    cur = conn.execute("INSERT INTO conversations (type, created_by) VALUES ('dm', ?)", (user_a,))
    conv_id = cur.lastrowid
    conn.execute('INSERT INTO conversation_members (conversation_id, user_id) VALUES (?, ?)', (conv_id, user_a))
    conn.execute('INSERT INTO conversation_members (conversation_id, user_id) VALUES (?, ?)', (conv_id, user_b))
    conn.commit()
    return conv_id


def get_conversation_membership(conn, conversation_id, user_id):
    """Kullanıcının bu sohbetteki üyelik satırını döndürür; üye değilse
    veya gruptan ayrılmışsa None döner. Yetkilendirme kontrollerinin temeli
    budur - tüm mesaj/grup işlemleri bu fonksiyondan geçmeli."""
    return conn.execute('''
        SELECT * FROM conversation_members
        WHERE conversation_id = ? AND user_id = ? AND left_at IS NULL
    ''', (conversation_id, user_id)).fetchone()


def conv_message_to_dict(conn, msg, viewer_id, users_cache=None):
    """Bir conv_messages satırını frontend'e gönderilecek sözlüğe çevirir.
    Silinmiş / benden-silinmiş mesajlar için içerik gizlenir."""
    users_cache = users_cache if users_cache is not None else {}
    sender_id = msg['sender_id']
    if sender_id not in users_cache:
        users_cache[sender_id] = conn.execute(
            'SELECT id, full_name, username, profile_photo FROM users WHERE id = ?', (sender_id,)
        ).fetchone()
    sender = users_cache[sender_id]

    hidden_for_me = conn.execute(
        'SELECT 1 FROM message_hides WHERE message_id = ? AND user_id = ?', (msg['id'], viewer_id)
    ).fetchone() is not None

    deleted = bool(msg['deleted_for_everyone'])
    content = msg['content']
    if content:
        # Eski (normalize edilmeden kaydedilmiş) mesajlarda da balonun aşırı
        # dikey şişmesini önlemek için görüntülemeden hemen önce de uygula.
        content = _normalize_message_text(content)
    image = msg['image'] if 'image' in msg.keys() else None
    if deleted:
        content = None
        image = None

    reply_preview = None
    if msg['reply_to_id']:
        rmsg = conn.execute('SELECT * FROM conv_messages WHERE id = ?', (msg['reply_to_id'],)).fetchone()
        if rmsg:
            rsender = conn.execute('SELECT full_name FROM users WHERE id = ?', (rmsg['sender_id'],)).fetchone()
            reply_preview = {
                'id': rmsg['id'],
                'sender_name': rsender['full_name'] if rsender else 'Kullanıcı',
                'content': ('Bu mesaj silindi.' if rmsg['deleted_for_everyone'] else (rmsg['content'] or '')[:120])
            }

    return {
        'id': msg['id'],
        'conversation_id': msg['conversation_id'],
        'sender_id': sender_id,
        'sender_name': sender['full_name'] if sender else 'Kullanıcı',
        'sender_username': sender['username'] if sender else '',
        'sender_photo': (sender['profile_photo'] if sender and sender['profile_photo'] else None),
        'content': content,
        'image': image,
        'deleted': deleted,
        'hidden_for_me': hidden_for_me,
        'edited': bool(msg['edited_at']),
        'created_at': utc_to_tr(msg['created_at']).strftime('%H:%M'),
        'created_at_full': utc_to_tr(msg['created_at']).strftime('%d.%m.%Y %H:%M'),
        'is_mine': sender_id == viewer_id,
        'reply_to': reply_preview,
        'is_system': bool(msg['is_system']),
    }


def get_pending_mandatory_attendance(user_id, in_dorm):
    """Kullanıcının doldurmadığı, aktif ve zorunlu bir yoklama var mı?"""
    if not in_dorm:
        return None
    conn = get_db()
    pending = conn.execute('''
        SELECT a.* FROM attendances a
        WHERE a.is_active = 1 AND a.is_mandatory = 1
          AND NOT EXISTS (
              SELECT 1 FROM attendance_records ar
              WHERE ar.attendance_id = a.id AND ar.user_id = ?
          )
        ORDER BY a.created_at DESC LIMIT 1
    ''', (user_id,)).fetchone()
    conn.close()
    return pending


def nest_comments(comments):
    """Düz yorum listesini üst yorum + yanıtlar (tek seviye) şeklinde gruplar."""
    top = []
    by_id = {}
    for c in comments:
        row = dict(c)
        row['replies'] = []
        by_id[row['id']] = row
    for c in comments:
        row = by_id[c['id']]
        if row.get('parent_id') and row['parent_id'] in by_id:
            by_id[row['parent_id']]['replies'].append(row)
        else:
            top.append(row)
    return top


def is_following(follower_id, followed_id):
    conn = get_db()
    row = conn.execute('SELECT id FROM follows WHERE follower_id = ? AND followed_id = ?',
                        (follower_id, followed_id)).fetchone()
    conn.close()
    return row is not None


def get_follow_counts(user_id):
    conn = get_db()
    followers = conn.execute('SELECT COUNT(*) as c FROM follows WHERE followed_id = ?',
                              (user_id,)).fetchone()['c']
    following = conn.execute('SELECT COUNT(*) as c FROM follows WHERE follower_id = ?',
                              (user_id,)).fetchone()['c']
    conn.close()
    return followers, following


def generate_quiz_code():
    import random, string
    conn = get_db()
    while True:
        code = ''.join(random.choices(string.digits, k=6))
        existing = conn.execute('SELECT id FROM quiz_sessions WHERE code = ?', (code,)).fetchone()
        if not existing:
            conn.close()
            return code


def quiz_room_name(session_id):
    return f"quiz_{session_id}"


def allowed_file(filename, exts):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in exts


def save_media(file_storage, folder, allowed):
    ext = file_storage.filename.rsplit('.', 1)[1].lower()
    filename = f"{uuid.uuid4().hex}.{ext}"
    path = os.path.join(folder, filename)
    file_storage.save(path)
    return filename


def award_xp(user_id, amount, reason):
    if amount <= 0: return
    conn=get_db(); conn.execute('INSERT INTO xp_transactions(user_id,amount,reason) VALUES (?,?,?)',(user_id,amount,reason))
    total=conn.execute('SELECT COALESCE(SUM(amount),0) xp FROM xp_transactions WHERE user_id=?',(user_id,)).fetchone()['xp']
    if total >= 500:
        b=conn.execute("SELECT id FROM badges WHERE code='legend'").fetchone()
        if b: conn.execute('INSERT OR IGNORE INTO user_badges(user_id,badge_id) VALUES (?,?)',(user_id,b['id']))
    conn.commit(); conn.close()

def get_user_xp(user_id):
    conn=get_db(); row=conn.execute('SELECT COALESCE(SUM(amount),0) xp FROM xp_transactions WHERE user_id=?',(user_id,)).fetchone(); conn.close(); return row['xp']

# ==================== TÜRKÇE TARİH VE MENÜ YARDIMCILARI ====================
MONTHS_TR = {
    1: 'Ocak', 2: 'Şubat', 3: 'Mart', 4: 'Nisan', 5: 'Mayıs', 6: 'Haziran',
    7: 'Temmuz', 8: 'Ağustos', 9: 'Eylül', 10: 'Ekim', 11: 'Kasım', 12: 'Aralık'
}
DAYS_TR = {
    0: 'Pazartesi', 1: 'Salı', 2: 'Çarşamba', 3: 'Perşembe', 4: 'Cuma', 5: 'Cumartesi', 6: 'Pazar'
}

def format_date_tr(date_val):
    if not date_val:
        return ''
    try:
        if isinstance(date_val, str):
            dt = datetime.strptime(date_val[:10], '%Y-%m-%d').date()
        elif isinstance(date_val, datetime):
            dt = date_val.date()
        elif isinstance(date_val, date):
            dt = date_val
        else:
            return str(date_val)
        return f"{dt.day} {MONTHS_TR.get(dt.month, '')} {dt.year}"
    except Exception:
        return str(date_val)

def day_name_tr(date_val):
    if not date_val:
        return ''
    try:
        if isinstance(date_val, str):
            dt = datetime.strptime(date_val[:10], '%Y-%m-%d').date()
        elif isinstance(date_val, datetime):
            dt = date_val.date()
        elif isinstance(date_val, date):
            dt = date_val
        else:
            return ''
        return DAYS_TR.get(dt.weekday(), '')
    except Exception:
        return ''

def split_menu_courses(text):
    if not text:
        return []
    raw_items = [i.strip() for i in re.split(r'[•\n\r;]+', str(text)) if i.strip()]
    return [i for i in raw_items if i and not re.fullmatch(r'[-–—*.,\s]+', i)]

# ---------------------------------------------------------------------------
# Öğün sınıflandırıcı (kahvaltı mı, akşam yemeği mi?)
# ---------------------------------------------------------------------------
# ESKİ HATA: Excel "otomatik" modda içe aktarılırken içerikte 'bal', 'zeytin',
# 'peynir' gibi ALT METİNLER aranıyordu. Bu yüzden "Balık", "Zeytinyağlı
# Fasulye", "Peynirli Börek" gibi akşam yemekleri kahvaltı sütununa yazılıyor,
# akşam sütunu boş kalıyordu -> ana ekranda "menü henüz girilmedi" görünüyordu.
# Artık kelime bazlı ve iki yönlü puanlama yapılıyor.
_BREAKFAST_WORDS = {
    'kahvaltı', 'kahvalti', 'bal', 'reçel', 'recel', 'tereyağı', 'tereyag', 'tereyağ', 'kaymak',
    'zeytin', 'zeytini', 'peynir', 'peyniri', 'kaşar', 'kasar', 'lor', 'labne', 'yumurta',
    'omlet', 'menemen', 'simit', 'poğaça', 'pogaca', 'açma', 'acma', 'gevrek', 'gevreği',
    'çay', 'cay', 'süt', 'sut', 'kakao', 'helva', 'tahin', 'pekmez', 'sucuk', 'sucuklu',
    'salam', 'sosis', 'nutella', 'krep', 'pankek', 'yulaf', 'granola', 'haşhaşlı', 'krem',
    'kahvaltılık', 'salatalık', 'salatalik', 'domates', 'biber', 'mısır', 'fındık', 'çikolata',
}
_DINNER_WORDS = {
    'çorba', 'çorbası', 'corba', 'corbasi', 'pilav', 'pilavı', 'pilavi', 'makarna', 'makarnası',
    'köfte', 'köftesi', 'kofte', 'tavuk', 'tavuklu', 'et', 'etli', 'kıyma', 'kıymalı', 'kebap',
    'kebabı', 'kebab', 'güveç', 'güveci', 'sote', 'sotesi', 'yahni', 'fasulye', 'fasulyesi',
    'nohut', 'mercimek', 'bulgur', 'zeytinyağlı', 'zeytinyagli', 'salata', 'salatası', 'cacık',
    'ayran', 'tatlı', 'tatlısı', 'sütlaç', 'kemalpaşa', 'baklava', 'revani', 'puding', 'meyve',
    'yoğurt', 'yoğurtlu', 'musakka', 'dolma', 'dolması', 'sarma', 'sarması', 'mantı', 'lahmacun',
    'pide', 'balık', 'hamsi', 'fırın', 'fırında', 'kızartma', 'kızartması', 'patates', 'türlü',
    'ıspanak', 'pırasa', 'kabak', 'patlıcan', 'karnıyarık', 'hünkar', 'hünkâr', 'iskender',
    'döner', 'şinitzel', 'schnitzel', 'haşlama', 'hoşaf', 'komposto', 'turşu', 'erişte',
    'şehriye', 'kavurma', 'bamya', 'barbunya', 'kuru', 'rosto', 'biftek', 'pirzola', 'burger',
    'hamburger', 'lazanya', 'spagetti', 'bezelye', 'enginar', 'kereviz', 'lahana', 'karnabahar',
    'brokoli', 'mücver', 'ezogelin', 'tarhana', 'yayla', 'pane', 'nugget', 'fajita', 'kaşarlı',
}


def _menu_tokens(text):
    txt = str(text or '').replace('İ', 'i').replace('I', 'ı').lower()
    return re.findall(r"[a-zçğıöşüâîû]+", txt)


def classify_meal_text(text):
    """Bir öğün metninin kahvaltı mı akşam yemeği mi olduğunu tahmin eder.
    Dönüş: 'breakfast' | 'dinner' | None (karar verilemedi)."""
    tokens = _menu_tokens(text)
    if not tokens:
        return None
    b = sum(1 for t in tokens if t in _BREAKFAST_WORDS)
    d = sum(1 for t in tokens if t in _DINNER_WORDS)
    if b == 0 and d == 0:
        return None
    return 'breakfast' if b > d else 'dinner'


def _repair_misclassified_menus(conn):
    """Eski sınıflandırıcı yüzünden akşam yemeği 'kahvaltı' sütununa yazılmış
    Excel kayıtlarını tek seferlik düzeltir (elle girilenlere dokunmaz)."""
    try:
        done = conn.execute("SELECT value FROM app_settings WHERE key='menu_classifier_v2_repaired'").fetchone()
        if done:
            return 0
        rows = conn.execute("""SELECT id, breakfast FROM meal_menus
                               WHERE COALESCE(TRIM(dinner),'') = '' AND COALESCE(TRIM(breakfast),'') <> ''
                                 AND COALESCE(source_file,'') <> 'manual_edit'""").fetchall()
        fixed = 0
        for r in rows:
            if classify_meal_text(r['breakfast']) == 'dinner':
                conn.execute("UPDATE meal_menus SET dinner = breakfast, breakfast = '' WHERE id = ?", (r['id'],))
                fixed += 1
        conn.execute("""INSERT INTO app_settings(key, value, updated_at) VALUES ('menu_classifier_v2_repaired', ?, CURRENT_TIMESTAMP)
                        ON CONFLICT(key) DO UPDATE SET value=excluded.value""", (str(fixed),))
        if fixed:
            print(f'[MENÜ] {fixed} günün akşam yemeği yanlışlıkla kahvaltıya yazılmıştı, düzeltildi.')
        return fixed
    except Exception as exc:
        print(f'[MENÜ] onarım atlandı: {exc}')
        return 0


app.jinja_env.filters['format_date_tr'] = format_date_tr
app.jinja_env.filters['day_name_tr'] = day_name_tr
app.jinja_env.filters['split_courses'] = split_menu_courses
app.jinja_env.filters['datetime_tr'] = format_datetime_tr
app.jinja_env.filters['time_tr'] = lambda value: format_datetime_tr(value, with_date=False)

def parse_menu_excel(path, meal_type='auto'):
    """Aylık yemekhane Excel menülerini akıllıca okur.
    
    Yatay TÜGVA formatlarını ve standart dikey satır formatlarını destekler.
    Gramaj satırlarını ('60 GR KIYMA', 'PİŞMİŞ 85 GR' vb.) otomatik ayıklar.
    """
    ext = os.path.splitext(path)[1].lower()
    sheets = []

    if ext == '.xls':
        if xlrd is None:
            raise RuntimeError(".xls dosyalarını okuyabilmek için xlrd==2.0.1 gereklidir. requirements.txt ile kurun.")
        book = xlrd.open_workbook(path, formatting_info=False)
        for sh in book.sheets():
            sheets.append((sh.name, [sh.row_values(r) for r in range(sh.nrows)]))
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            sheets.append((ws.title, list(ws.iter_rows(values_only=True))))

    def norm(v):
        return ' '.join(str(v or '').strip().lower().replace('\n', ' ').split())

    months = {'ocak':1,'şubat':2,'mart':3,'nisan':4,'mayıs':5,'haziran':6,
              'temmuz':7,'ağustos':8,'eylül':9,'ekim':10,'kasım':11,'aralık':12}

    def parse_date(raw):
        if raw is None or raw == '':
            return None
        if isinstance(raw, (datetime, date)):
            return raw.strftime('%Y-%m-%d')
        if isinstance(raw, (int, float)) and ext == '.xls' and xlrd is not None:
            try:
                y, m, d, *_ = xlrd.xldate_as_tuple(raw, 0)
                if y > 1900 and 1 <= m <= 12 and 1 <= d <= 31:
                    return date(y, m, d).isoformat()
            except Exception:
                pass
        txt = str(raw).strip()
        for fmt in ('%d.%m.%Y','%d/%m/%Y','%d-%m-%Y','%Y-%m-%d','%d.%m.%y'):
            try:
                return datetime.strptime(txt, fmt).strftime('%Y-%m-%d')
            except ValueError:
                pass
        m = re.search(r'(\d{1,2})\s+([a-zçğıöşü]+)(?:\s+(\d{4}))?', txt.lower())
        if m and m.group(2) in months:
            return date(int(m.group(3) or now_tr().year), months[m.group(2)], int(m.group(1))).isoformat()
        return None

    def column_index(headers, keywords):
        for i, h in enumerate(headers):
            if any(k in h for k in keywords):
                return i
        return None

    def clean_item(v):
        txt = str(v or '').strip()
        if not txt:
            return ''
        norm_t = txt.lower()
        # Gramaj, porsiyon, kalori satırlarını temizle
        if re.search(r'\b\d+\s*(?:gr|gram|kg|kcal|kalori)\b', norm_t) or re.search(r'\b(?:gr|gram|kg|kcal|kalori)\s*\d+\b', norm_t):
            return ''
        if 'pişmiş' in norm_t and re.search(r'\d+', norm_t):
            return ''
        if norm_t in ('gramaj', 'kalori', 'protein', 'porsiyon', 'toplam kalori'):
            return ''
        if re.fullmatch(r'[\d\s.,/+-]+', txt):
            return ''
        return txt

    result = {}

    for sheet_name, rows in sheets:
        if not rows:
            continue

        # --- Format A: Tarihler yatay sütunlarda (TÜGVA aylık tablosu) ---
        horizontal_date_rows = []
        for ri, row in enumerate(rows):
            dates = [(ci, parse_date(v)) for ci, v in enumerate(row)]
            valid = [(ci, ds) for ci, ds in dates if ds]
            if len(valid) >= 2:
                horizontal_date_rows.append((ri, valid))

        if horizontal_date_rows:
            for idx, (ri, valid_dates) in enumerate(horizontal_date_rows):
                next_date_row = horizontal_date_rows[idx + 1][0] if idx + 1 < len(horizontal_date_rows) else len(rows)
                for ci, ds in valid_dates:
                    items = []
                    for rr in range(ri + 1, min(next_date_row, ri + 8)):
                        if ci >= len(rows[rr]):
                            continue
                        item = clean_item(rows[rr][ci])
                        if item and item not in items:
                            items.append(item)
                    if not items:
                        continue
                    
                    meal_str = ' • '.join(items)
                    old = result.get(ds, ['', '', '', ''])
                    
                    if meal_type == 'breakfast':
                        result[ds] = [meal_str, old[1], old[2], old[3]]
                    elif meal_type == 'dinner':
                        result[ds] = [old[0], old[1], meal_str, old[3]]
                    else:
                        # Otomatik tespit: önce sayfa adı, sonra kelime bazlı içerik puanı
                        norm_s = norm(sheet_name)
                        if any(k in norm_s for k in ['kahvalt', 'breakfast', 'sabah']):
                            target = 'breakfast'
                        elif any(k in norm_s for k in ['akşam', 'aksam', 'dinner', 'öğle', 'ogle']):
                            target = 'dinner'
                        else:
                            target = classify_meal_text(' • '.join(items)) or 'dinner'
                        result[ds] = [meal_str if target == 'breakfast' else old[0],
                                      old[1],
                                      meal_str if target == 'dinner' else old[2],
                                      old[3]]
            continue

        # --- Format B: Her satırda tek bir tarih (Dikey Format) ---
        header_row = 0
        best_score = -1
        for ri, row in enumerate(rows[:20]):
            hs = [norm(v) for v in row]
            score = 0
            if column_index(hs, ['tarih', 'date', 'gün', 'gun']) is not None: score += 3
            if column_index(hs, ['kahvalt', 'breakfast', 'sabah']) is not None: score += 2
            if column_index(hs, ['akşam', 'aksam', 'dinner']) is not None: score += 2
            if score > best_score:
                best_score, header_row = score, ri

        headers = [norm(v) for v in rows[header_row]]
        date_i = column_index(headers, ['tarih', 'date', 'gün', 'gun'])
        if date_i is None:
            date_i = 0
        b_i = column_index(headers, ['kahvalt', 'breakfast', 'sabah'])
        d_i = column_index(headers, ['akşam', 'aksam', 'dinner'])
        l_i = column_index(headers, ['öğle', 'ogle', 'lunch'])
        s_i = column_index(headers, ['ikram', 'snack', 'ara öğün', 'ara ogun'])
        fallback = next((i for i in range(len(headers)) if i != date_i and headers[i]), None)
        
        if meal_type == 'dinner' and d_i is None:
            d_i = fallback
        elif meal_type == 'breakfast' and b_i is None:
            b_i = fallback
        else:
            if b_i is None and any(k in norm(sheet_name) for k in ['kahvalt','breakfast','sabah']): b_i = fallback
            if d_i is None and any(k in norm(sheet_name) for k in ['akşam','aksam','dinner']): d_i = fallback
            if b_i is None and d_i is None:
                d_i = fallback

        for row in rows[header_row + 1:]:
            ds = parse_date(row[date_i] if date_i < len(row) else None)
            if not ds:
                continue
            def val(i):
                return clean_item(row[i]) if i is not None and i < len(row) else ''
            old = result.get(ds, ['', '', '', ''])
            result[ds] = [val(b_i) or old[0], val(l_i) or old[1], val(d_i) or old[2], val(s_i) or old[3]]

    return [(ds, vals[0], vals[1], vals[2], vals[3]) for ds, vals in sorted(result.items())]

@app.context_processor
def inject_globals():
    base = {'user_roles': USER_ROLES, 'page_permissions': PAGE_PERMISSIONS,
            'admin_permissions': ADMIN_PERMISSIONS}
    if 'user_id' in session:
        me_dict = current_user_dict() or {}
        role_info = USER_ROLES.get(resolve_role_key(me_dict)) or next(iter(USER_ROLES.values()), {})

        def can(perm_key):
            return user_has_perm(me_dict, perm_key)

        base.update({
            'unread_count': get_unread_count(session['user_id']),
            'unread_dm_count': get_unread_dm_count(session['user_id']),
            'current_user': me_dict or None,
            'current_role_info': role_info,
            'can': can,
            'has_admin_access': user_has_admin_access(me_dict),
            'can_manage_maintenance': can_user_manage_maintenance(me_dict),
            'can_manage_attendance': can_user_manage_attendance(me_dict),
            'current_xp': get_user_xp(session['user_id']),
        })
        return base
    base['can'] = lambda perm_key: False
    return base


# ==================== ROTALAR ====================

@app.route('/manifest.json')
def pwa_manifest():
    return send_from_directory('static', 'manifest.json', mimetype='application/manifest+json')


@app.route('/sw.js')
@app.route('/service-worker.js')
def pwa_service_worker():
    response = make_response(send_from_directory('static', 'sw.js', mimetype='application/javascript'))
    response.headers['Service-Worker-Allowed'] = '/'
    response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    return response


@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username', '').strip().lower()
        password = request.form.get('password', '')
        remember = bool(request.form.get('remember'))

        conn = get_db()
        user = conn.execute('SELECT * FROM users WHERE username = ?',
                             (username,)).fetchone()
        conn.close()

        if user and check_password_hash(user['password'], password):
            # "Beni Hatırla" işaretliyse oturumu kalıcı yap (PERMANENT_SESSION_LIFETIME
            # kadar, tarayıcı kapansa bile) — işaretli değilse tarayıcı kapanınca
            # oturum sona erer (klasik davranış).
            session.permanent = remember
            session['user_id'] = user['id']
            session['username'] = user['username']
            session['is_admin'] = bool(user['is_admin'])
            if not (user['full_name'] or '').strip():
                flash('Devam etmek için Ad Soyad bilginizi doldurmanız gerekiyor.', 'warning')
                return redirect(url_for('edit_profile', required=1))
            flash(f'Hoş geldin, {user["full_name"]}!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Kullanıcı adı veya şifre hatalı!', 'danger')

    return render_template('login.html')


@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form.get('username', '').strip().lower()
        password = request.form.get('password', '')
        full_name = request.form.get('full_name', '').strip()
        room_number = request.form.get('room_number', '').strip()
        if not full_name:
            flash('Ad Soyad alanı zorunludur.', 'danger')
            conn.close()
            return redirect(url_for('edit_profile', required=1))
        in_dorm = 1 if request.form.get('in_dorm') else 0

        if not username or not password or not full_name:
            flash('Lütfen tüm zorunlu alanları doldurun.', 'danger')
            return redirect(url_for('register'))

        if not re.fullmatch(r'[a-z0-9çğıöşü._-]+', username):
            flash('Kullanıcı adı sadece küçük harf, rakam ve . _ - karakterlerini içerebilir.', 'danger')
            return redirect(url_for('register'))

        if len(password) < 4:
            flash('Şifre en az 4 karakter olmalıdır.', 'danger')
            return redirect(url_for('register'))

        conn = get_db()
        existing = conn.execute('SELECT id FROM users WHERE username = ?',
                                 (username,)).fetchone()
        if existing:
            conn.close()
            flash('Bu kullanıcı adı zaten kullanılıyor!', 'danger')
            return redirect(url_for('register'))

        hashed = generate_password_hash(password)
        conn.execute('''
            INSERT INTO users (username, password, full_name, room_number, in_dorm)
            VALUES (?, ?, ?, ?, ?)
        ''', (username, hashed, full_name, room_number, in_dorm))
        conn.commit()
        conn.close()

        flash('Kayıt başarılı! Şimdi giriş yapabilirsiniz.', 'success')
        return redirect(url_for('login'))

    return render_template('login.html', register=True)


@app.route('/logout')
def logout():
    session.clear()
    flash('Çıkış yapıldı.', 'info')
    return redirect(url_for('login'))


ACTIVE_WINDOW_SECONDS = 300  # Son 5 dakikada işlem yapan kullanıcı "aktif" sayılır.


@app.before_request
def touch_last_seen():
    """Giriş yapmış kullanıcının 'son görülme' zamanını (en fazla 60 sn'de bir) günceller."""
    if 'user_id' not in session or request.endpoint in (None, 'static'):
        return
    now_ts = int(now_tr().timestamp())
    if now_ts - int(session.get('_seen_ts', 0)) < 60:
        return
    session['_seen_ts'] = now_ts
    try:
        conn = get_db()
        conn.execute('UPDATE users SET last_seen = ? WHERE id = ?',
                     (now_tr().strftime('%Y-%m-%d %H:%M:%S'), session['user_id']))
        conn.commit()
        conn.close()
    except Exception as exc:
        print(f'[LAST SEEN] güncellenemedi: {exc}')


def greeting_for(now):
    h = now.hour
    if h < 5:
        return 'İyi geceler'
    if h < 12:
        return 'Günaydın'
    if h < 18:
        return 'İyi günler'
    if h < 22:
        return 'İyi akşamlar'
    return 'İyi geceler'


def get_meal_schedule(now=None):
    """Yemekhane saatlerini admin ayarlarından okur (menü sayfasıyla aynı kaynak)."""
    now = now or now_tr()
    weekend = now.weekday() >= 5
    return {
        'is_weekend': weekend,
        'breakfast_start': get_app_setting('meal_breakfast_weekend_start' if weekend else 'meal_breakfast_weekday_start', '09:00' if weekend else '07:00'),
        'breakfast_end': get_app_setting('meal_breakfast_weekend_end' if weekend else 'meal_breakfast_weekday_end', '11:00' if weekend else '09:00'),
        'dinner_start': get_app_setting('meal_dinner_start', '18:00'),
        'dinner_end': get_app_setting('meal_dinner_end', '20:00'),
    }


def _hm_to_minutes(value):
    try:
        h, m = [int(x) for x in str(value).split(':', 1)]
        return h * 60 + m
    except (ValueError, TypeError):
        return 0


def meal_status(schedule, now):
    """Şu an hangi öğün açık ya da sıradaki öğün hangisi?"""
    cur = now.hour * 60 + now.minute
    b0, b1 = _hm_to_minutes(schedule['breakfast_start']), _hm_to_minutes(schedule['breakfast_end'])
    d0, d1 = _hm_to_minutes(schedule['dinner_start']), _hm_to_minutes(schedule['dinner_end'])
    if b0 <= cur < b1:
        return {'key': 'breakfast', 'state': 'open', 'text': f"Kahvaltı şu an açık, saat {schedule['breakfast_end']}'a kadar"}
    if d0 <= cur < d1:
        return {'key': 'dinner', 'state': 'open', 'text': f"Akşam yemeği şu an açık, saat {schedule['dinner_end']}'a kadar"}
    if cur < b0:
        return {'key': 'breakfast', 'state': 'next', 'text': f"Sıradaki öğün kahvaltı, saat {schedule['breakfast_start']}"}
    if cur < d0:
        return {'key': 'dinner', 'state': 'next', 'text': f"Sıradaki öğün akşam yemeği, saat {schedule['dinner_start']}"}
    return {'key': 'dinner', 'state': 'done', 'text': 'Bugünün yemekleri bitti'}


def _parse_event_dt(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace(' ', 'T')).replace(tzinfo=ISTANBUL_TZ)
    except Exception:
        return None


def _meal_focus(menu, schedule, status, tomorrow_menu=None, tomorrow_schedule=None):
    """Şu an en ilgili öğünü (açık olan ya da sıradaki) kısa bir satır olarak hazırlar.

    Öncelik: şu an açık/sıradaki öğün -> o öğün boşsa bugünün diğer öğünü ->
    akşam yemeği bittiyse yarının kahvaltısı. Böylece menü girilmişken ekranda
    yanlışlıkla "henüz girilmedi" yazmaz."""
    def meal_text(row, key):
        try:
            return (row[key] or '').strip() if row else ''
        except (IndexError, KeyError):
            return ''

    def build(row, key, sched, prefix=''):
        label = 'Kahvaltı' if key == 'breakfast' else 'Akşam'
        time_txt = (f"{sched['breakfast_start']} – {sched['breakfast_end']}" if key == 'breakfast'
                    else f"{sched['dinner_start']} – {sched['dinner_end']}")
        courses = split_menu_courses(meal_text(row, key))
        return {'label': prefix + label, 'time': time_txt, 'key': key,
                'line': ' + '.join(courses[:4]), 'more': max(len(courses) - 4, 0), 'empty_text': ''}

    key = status['key']
    other = 'dinner' if key == 'breakfast' else 'breakfast'

    # Akşam yemeği bitti -> yarının kahvaltısını göster (varsa)
    if status.get('state') == 'done' and meal_text(tomorrow_menu, 'breakfast'):
        return build(tomorrow_menu, 'breakfast', tomorrow_schedule or schedule, 'Yarın · ')

    if meal_text(menu, key):
        return build(menu, key, schedule)
    if meal_text(menu, other):
        return build(menu, other, schedule, 'Bugün · ')

    res = build(menu, key, schedule)
    res['empty_text'] = 'Bugünün menüsü henüz girilmedi' if not menu else 'Bu öğünün menüsü henüz girilmedi'
    return res


def get_today_stats(user_id):
    """'Bugün Yurtta' ekranının tüm canlı sayıları. Hem sayfa hem /api/today bunu kullanır."""
    now = now_tr()
    today = now.date().isoformat()
    now_str = now.strftime('%Y-%m-%d %H:%M:%S')
    conn = get_db()

    active_cut = (now - timedelta(seconds=ACTIVE_WINDOW_SECONDS)).strftime('%Y-%m-%d %H:%M:%S')
    active_count = conn.execute('SELECT COUNT(*) c FROM users WHERE last_seen >= ?', (active_cut,)).fetchone()['c']
    active_count = max(active_count, 1)  # sayfaya bakan kişi de aktif

    # Sıradaki etkinlik: bugün ve henüz bitmemiş olanlar öncelikli, yoksa en yakın gelecek etkinlik.
    ev_rows = conn.execute('''
        SELECT e.id, e.title, e.location, e.event_date,
               (SELECT COUNT(*) FROM event_participants WHERE event_id = e.id) participant_count
        FROM events e WHERE substr(e.event_date,1,10) >= ? ORDER BY e.event_date ASC LIMIT 20
    ''', (today,)).fetchall()
    next_event = None
    for e in ev_rows:
        dt = _parse_event_dt(e['event_date'])
        if dt is None or dt >= now - timedelta(hours=2):
            next_event = dict(e)
            next_event['is_today'] = e['event_date'][:10] == today
            next_event['time'] = e['event_date'][11:16] if len(e['event_date']) >= 16 else ''
            next_event['date'] = e['event_date'][:10]
            break

    # Çamaşırhane: şu an rezervasyonu süren makineler dolu sayılır.
    laundry_total = conn.execute('SELECT COUNT(*) c FROM laundry_machines WHERE enabled=1').fetchone()['c']
    busy = conn.execute('''SELECT COUNT(DISTINCT r.machine_id) c FROM laundry_reservations r
                           JOIN laundry_machines m ON m.id = r.machine_id AND m.enabled = 1
                           WHERE r.status='reserved' AND r.start_at <= ? AND r.end_at > ?''', (now_str, now_str)).fetchone()['c']
    laundry_free = max(laundry_total - busy, 0)

    # Menü + bugünkü görüşler
    menu = conn.execute('SELECT * FROM meal_menus WHERE menu_date = ?', (today,)).fetchone()
    tomorrow_dt = now + timedelta(days=1)
    tomorrow_menu = conn.execute('SELECT * FROM meal_menus WHERE menu_date = ?',
                                 (tomorrow_dt.date().isoformat(),)).fetchone()
    rating =conn.execute('SELECT ROUND(AVG(overall_rating),1) avg_r, COUNT(id) cnt FROM menu_ratings WHERE menu_date = ?', (today,)).fetchone()
    comments = conn.execute('''
        SELECT r.comment, r.overall_rating, r.updated_at, u.full_name, u.username
        FROM menu_ratings r JOIN users u ON u.id = r.user_id
        WHERE r.menu_date = ? AND TRIM(COALESCE(r.comment,'')) <> ''
        ORDER BY r.updated_at DESC LIMIT 4
    ''', (today,)).fetchall()
    comment_total = conn.execute("SELECT COUNT(*) c FROM menu_ratings WHERE menu_date = ? AND TRIM(COALESCE(comment,'')) <> ''", (today,)).fetchone()['c']
    my_rating = conn.execute('SELECT overall_rating FROM menu_ratings WHERE menu_date = ? AND user_id = ?', (today, user_id)).fetchone()

    # Gündem: son 24 saatteki paylaşım / yorum (SQLite CURRENT_TIMESTAMP UTC tutar)
    since = (datetime.utcnow() - timedelta(hours=24)).strftime('%Y-%m-%d %H:%M:%S')
    new_posts = conn.execute('SELECT COUNT(*) c FROM posts WHERE created_at >= ?', (since,)).fetchone()['c']
    new_comments = conn.execute('SELECT COUNT(*) c FROM post_comments WHERE created_at >= ?', (since,)).fetchone()['c']

    # Anket
    poll = conn.execute('''SELECT * FROM polls WHERE is_active = 1 AND (expires_at IS NULL OR expires_at = '' OR expires_at >= ?)
                           ORDER BY created_at DESC LIMIT 1''', (now_str,)).fetchone()
    poll_data = None
    if poll:
        opts = conn.execute('''SELECT o.id, o.label, (SELECT COUNT(*) FROM poll_votes v WHERE v.option_id = o.id) votes
                               FROM poll_options o WHERE o.poll_id = ? ORDER BY o.id''', (poll['id'],)).fetchall()
        mine = conn.execute('SELECT option_id FROM poll_votes WHERE poll_id = ? AND user_id = ?', (poll['id'], user_id)).fetchone()
        total = sum(o['votes'] for o in opts)
        poll_data = {
            'id': poll['id'], 'question': poll['question'], 'total': total,
            'my_option': mine['option_id'] if mine else None,
            'options': [{'id': o['id'], 'label': o['label'], 'votes': o['votes'],
                         'pct': round(o['votes'] * 100 / total) if total else 0} for o in opts],
        }
    conn.close()

    schedule = get_meal_schedule(now)
    return {
        'greeting': greeting_for(now),
        'active_count': active_count,
        'event': next_event,
        'schedule': schedule,
        'meal_status': meal_status(schedule, now),
        'laundry_free': laundry_free,
        'laundry_total': laundry_total,
        'rating_avg': rating['avg_r'] if rating and rating['avg_r'] is not None else None,
        'rating_count': rating['cnt'] if rating else 0,
        'my_rating': my_rating['overall_rating'] if my_rating else None,
        'comments': [{'text': c['comment'], 'stars': c['overall_rating'],
                      'name': (c['full_name'] or c['username']).split(' ')[0]} for c in comments],
        'comment_total': comment_total,
        'new_posts': new_posts,
        'new_comments': new_comments,
        'poll': poll_data,
        'meal_focus': _meal_focus(menu, schedule, meal_status(schedule, now),
                                  tomorrow_menu, get_meal_schedule(tomorrow_dt)),
        'has_menu_today': bool(menu),
    }


@app.route('/dashboard')
@login_required
def dashboard():
    conn = get_db()

    recent_posts = conn.execute('''
        SELECT p.*, u.full_name, u.username, u.profile_photo
        FROM posts p
        JOIN users u ON p.user_id = u.id
        ORDER BY p.created_at DESC
        LIMIT 4
    ''').fetchall()

    active_attendances = conn.execute('''
        SELECT a.*,
               (SELECT status FROM attendance_records
                WHERE attendance_id = a.id AND user_id = ?) as my_status
        FROM attendances a
        WHERE a.is_active = 1
        ORDER BY a.created_at DESC
    ''', (session['user_id'],)).fetchall()

    me = conn.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    announcements = conn.execute("SELECT * FROM announcements WHERE (expires_at IS NULL OR expires_at >= ?) ORDER BY CASE priority WHEN 'urgent' THEN 0 WHEN 'high' THEN 1 ELSE 2 END, created_at DESC LIMIT 3", (now_tr().date().isoformat(),)).fetchall()
    conn.close()

    pending_attendance = get_pending_mandatory_attendance(session['user_id'], me['in_dorm'])
    today = get_today_stats(session['user_id'])

    return render_template('dashboard.html', recent_posts=recent_posts,
                            active_attendances=active_attendances, me=me,
                            pending_attendance=pending_attendance,
                            announcements=announcements, today=today)


@app.route('/api/today')
@login_required
def api_today():
    resp = jsonify(get_today_stats(session['user_id']))
    resp.headers['Cache-Control'] = 'no-store'
    return resp


@app.route('/poll/<int:poll_id>/vote', methods=['POST'])
@login_required
def poll_vote(poll_id):
    data = request.get_json(silent=True) or {}
    try:
        option_id = int(data.get('option_id') or request.form.get('option_id', 0))
    except (ValueError, TypeError):
        return jsonify({'ok': False, 'error': 'Geçersiz seçenek.'}), 400
    conn = get_db()
    poll = conn.execute('SELECT id FROM polls WHERE id = ? AND is_active = 1', (poll_id,)).fetchone()
    opt = conn.execute('SELECT id FROM poll_options WHERE id = ? AND poll_id = ?', (option_id, poll_id)).fetchone()
    if not poll or not opt:
        conn.close()
        return jsonify({'ok': False, 'error': 'Anket kapalı ya da seçenek bulunamadı.'}), 404
    conn.execute('''INSERT INTO poll_votes(poll_id, option_id, user_id) VALUES(?,?,?)
                    ON CONFLICT(poll_id,user_id) DO UPDATE SET option_id=excluded.option_id, created_at=CURRENT_TIMESTAMP''',
                 (poll_id, option_id, session['user_id']))
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'poll': get_today_stats(session['user_id'])['poll']})


@app.route('/admin/poll/create', methods=['POST'])
@permission_required('manage_polls')
def create_poll():
    question = request.form.get('question', '').strip()[:200]
    options = [o.strip()[:80] for o in request.form.getlist('option') if o.strip()]
    if not question or len(options) < 2:
        flash('Anket için bir soru ve en az iki seçenek gerekli.', 'danger')
        return redirect(url_for('admin_panel'))
    conn = get_db()
    # Aynı anda tek aktif anket: yenisi açılınca eskisi kapanır.
    conn.execute('UPDATE polls SET is_active = 0 WHERE is_active = 1')
    cur = conn.execute('INSERT INTO polls(question, created_by) VALUES(?,?)', (question, session['user_id']))
    pid = cur.lastrowid
    conn.executemany('INSERT INTO poll_options(poll_id, label) VALUES(?,?)', [(pid, o) for o in options[:6]])
    conn.commit()
    conn.close()
    flash('Anket yayınlandı. Ana sayfada görünecek.', 'success')
    return redirect(url_for('admin_panel'))


@app.route('/admin/poll/<int:poll_id>/close', methods=['POST'])
@permission_required('manage_polls')
def close_poll(poll_id):
    conn = get_db()
    conn.execute('UPDATE polls SET is_active = 0 WHERE id = ?', (poll_id,))
    conn.commit()
    conn.close()
    flash('Anket kapatıldı.', 'success')
    return redirect(url_for('admin_panel'))


# ==================== BİLDİRİMLER ====================


@app.route('/api/push/vapid-public-key')
@login_required
def vapid_public_key():
    return jsonify({'publicKey': _vapid_public_key()})

@app.route('/api/push/subscribe', methods=['POST'])
@login_required
def push_subscribe():
    data = request.get_json(silent=True) or {}
    sub = data.get('subscription') or data
    endpoint = (sub.get('endpoint') or '').strip()
    keys = sub.get('keys') or {}
    p256dh = (keys.get('p256dh') or '').strip()
    auth = (keys.get('auth') or '').strip()
    if not endpoint or not p256dh or not auth:
        return jsonify({'ok': False, 'error': 'Geçersiz push aboneliği.'}), 400
    conn = get_db()
    conn.execute('''
        INSERT INTO push_subscriptions (user_id, endpoint, p256dh, auth, updated_at)
        VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(endpoint) DO UPDATE SET user_id=excluded.user_id, p256dh=excluded.p256dh, auth=excluded.auth, updated_at=CURRENT_TIMESTAMP
    ''', (session['user_id'], endpoint, p256dh, auth))
    conn.commit(); conn.close()
    return jsonify({'ok': True})

@app.route('/api/push/unsubscribe', methods=['POST'])
@login_required
def push_unsubscribe():
    data = request.get_json(silent=True) or {}
    endpoint = (data.get('endpoint') or '').strip()
    conn = get_db()
    if endpoint:
        conn.execute('DELETE FROM push_subscriptions WHERE endpoint = ? AND user_id = ?', (endpoint, session['user_id']))
    else:
        conn.execute('DELETE FROM push_subscriptions WHERE user_id = ?', (session['user_id'],))
    conn.commit(); conn.close()
    return jsonify({'ok': True})

@app.route('/api/push/test', methods=['POST'])
@login_required
def push_test():
    socketio.start_background_task(push_notification_to_user, session['user_id'], 'YURT. Bildirimler Açık 🔔', 'Telefon bildirimleri başarıyla etkinleştirildi.', url_for('notifications'))
    return jsonify({'ok': True})

@app.route('/notifications')
@login_required
def notifications():
    conn = get_db()
    notifs = conn.execute('''
        SELECT * FROM notifications
        WHERE user_id = ?
        ORDER BY created_at DESC
    ''', (session['user_id'],)).fetchall()
    conn.close()
    return render_template('notifications.html', notifications=notifs)


@app.route('/api/notifications')
@login_required
def api_notifications():
    conn = get_db()
    notifs = conn.execute('''
        SELECT id, title, message, link, is_read, created_at
        FROM notifications
        WHERE user_id = ?
        ORDER BY created_at DESC
        LIMIT 10
    ''', (session['user_id'],)).fetchall()

    unread = get_unread_count(session['user_id'])
    conn.close()

    result = []
    for n in notifs:
        item = {key: n[key] for key in n.keys()}
        item['time'] = format_datetime_tr(n['created_at'], with_date=False)
        item['created_at'] = format_datetime_tr(n['created_at'], with_date=True)
        result.append(item)

    return jsonify({'notifications': result, 'unread_count': unread})


@app.route('/notification/<int:notif_id>/read')
@login_required
def mark_notification_read(notif_id):
    conn = get_db()
    conn.execute('UPDATE notifications SET is_read = 1 WHERE id = ? AND user_id = ?',
                 (notif_id, session['user_id']))
    conn.commit()
    notif = conn.execute('SELECT link FROM notifications WHERE id = ?', (notif_id,)).fetchone()
    conn.close()

    if notif and notif['link']:
        return redirect(notif['link'])
    return redirect(url_for('notifications'))


@app.route('/notifications/read-all')
@login_required
def mark_all_read():
    conn = get_db()
    conn.execute('UPDATE notifications SET is_read = 1 WHERE user_id = ?',
                 (session['user_id'],))
    conn.commit()
    conn.close()
    flash('Tüm bildirimler okundu olarak işaretlendi.', 'success')
    return redirect(url_for('notifications'))


# ==================== GLOBAL SEARCH API ====================

@app.route('/api/search')
@login_required
def api_global_search():
    q = request.args.get('q', '').strip()
    if not q or len(q) < 1:
        return jsonify({'users': [], 'announcements': [], 'events': [], 'menus': []})
    
    norm_q = normalize_tr_text(q)
    conn = get_db()
    
    # 1. Kullanıcılar (Öğrenciler & Yetkililer)
    all_users = conn.execute('''
        SELECT id, username, full_name, profile_photo, room_number, role, is_admin
        FROM users
        ORDER BY is_admin DESC, full_name ASC
    ''').fetchall()
    
    # Kullanıcı adı ve ad-soyad her ikisi de aranır; en alakalı sonuçlar üstte çıksın
    # diye basit bir skorlama uygulanır (tam eşleşme > başında geçme > içinde geçme).
    scored_users = []
    for u in all_users:
        u_dict = dict(u)
        role_key = u_dict.get('role') or ('yurt_muduru' if u_dict.get('is_admin') else 'ogrenci')
        role_info = USER_ROLES.get(role_key, USER_ROLES['ogrenci'])

        norm_username = normalize_tr_text(u['username'])
        norm_fullname = normalize_tr_text(u['full_name'])
        norm_room = normalize_tr_text(u['room_number'] or '')
        norm_role = normalize_tr_text(f"{role_info['title']} {role_info['badge']}")
        norm_text = f"{norm_username} {norm_fullname} {norm_room} {norm_role}"

        if norm_q not in norm_text:
            continue

        score = 4
        if norm_username == norm_q or norm_fullname == norm_q:
            score = 0
        elif norm_username.startswith(norm_q) or any(part.startswith(norm_q) for part in norm_fullname.split()):
            score = 1
        elif norm_q in norm_username or norm_q in norm_fullname:
            score = 2

        scored_users.append((score, u_dict, role_key, role_info))

    scored_users.sort(key=lambda item: item[0])

    matched_users = []
    for score, u, role_key, role_info in scored_users[:8]:
        matched_users.append({
            'id': u['id'],
            'username': u['username'],
            'full_name': u['full_name'],
            'profile_photo': u['profile_photo'],
            'room_number': u['room_number'] or '',
            'role': role_key,
            'role_badge': role_info['badge'],
            'role_class': role_info['class'],
            'role_icon': role_info['icon'],
            'url': url_for('profile', username=u['username'])
        })

    # 2. Duyurular
    all_announcements = conn.execute('''
        SELECT id, title, content, priority, created_at
        FROM announcements
        ORDER BY created_at DESC
        LIMIT 40
    ''').fetchall()
    matched_announcements = []
    for a in all_announcements:
        searchable_text = f"{a['title']} {a['content']}"
        if norm_q in normalize_tr_text(searchable_text) or q.lower() in searchable_text.lower():
            matched_announcements.append({
                'id': a['id'],
                'title': a['title'],
                'content': a['content'][:90] + ('...' if len(a['content']) > 90 else ''),
                'priority': a['priority'],
                'url': url_for('announcements_page')
            })
            if len(matched_announcements) >= 4:
                break

    # 3. Etkinlikler
    all_events = conn.execute('''
        SELECT id, title, description, event_date, location
        FROM events
        ORDER BY event_date DESC
        LIMIT 40
    ''').fetchall()
    matched_events = []
    for e in all_events:
        searchable_text = f"{e['title']} {e['description'] or ''} {e['location'] or ''}"
        if norm_q in normalize_tr_text(searchable_text) or q.lower() in searchable_text.lower():
            matched_events.append({
                'id': e['id'],
                'title': e['title'],
                'event_date': e['event_date'],
                'location': e['location'] or '',
                'url': url_for('event_detail', event_id=e['id'])
            })
            if len(matched_events) >= 4:
                break

    # 4. Yemek Menüleri
    all_menus = conn.execute('''
        SELECT id, menu_date, breakfast, dinner, snack
        FROM meal_menus
        ORDER BY menu_date DESC
        LIMIT 60
    ''').fetchall()
    matched_menus = []
    for m in all_menus:
        searchable_text = f"{m['menu_date']} {m['breakfast'] or ''} {m['dinner'] or ''} {m['snack'] or ''}"
        if norm_q in normalize_tr_text(searchable_text) or q.lower() in searchable_text.lower():
            matched_menus.append({
                'id': m['id'],
                'menu_date': m['menu_date'],
                'menu_date_tr': format_date_tr(m['menu_date']),
                'breakfast': m['breakfast'] or '',
                'dinner': m['dinner'] or '',
                'url': url_for('menu_page')
            })
            if len(matched_menus) >= 4:
                break

    conn.close()

    return jsonify({
        'users': matched_users,
        'announcements': matched_announcements,
        'events': matched_events,
        'menus': matched_menus
    })


@app.route('/api/menu/<menu_date>')
@login_required
def api_get_menu_date(menu_date):
    conn = get_db()
    menu = conn.execute('SELECT * FROM meal_menus WHERE menu_date = ?', (menu_date,)).fetchone()
    conn.close()
    if not menu:
        return jsonify({'ok': False, 'error': 'Menü bulunamadı'}), 404
    return jsonify({
        'ok': True,
        'menu': {
            'id': menu['id'],
            'menu_date': menu['menu_date'],
            'menu_date_tr': format_date_tr(menu['menu_date']),
            'day_name': day_name_tr(menu['menu_date']),
            'breakfast': menu['breakfast'] or '',
            'dinner': menu['dinner'] or '',
            'snack': menu['snack'] or ''
        }
    })


# ==================== SOHBET (GENEL) ====================

@app.route('/chat')
@login_required
def chat():
    conn = get_db()
    messages = conn.execute('''
        SELECT * FROM messages WHERE room = 'genel'
        ORDER BY created_at ASC LIMIT 100
    ''').fetchall()
    conn.close()
    return render_template('chat.html', messages=messages)


@socketio.on('connect')
def handle_connect():
    if session.get('user_id'):
        join_room(f"user_{session['user_id']}")


@socketio.on('send_message')
def handle_message(data):
    username = session.get('username', 'Misafir')
    message = data.get('message', '').strip()

    if message:
        conn = get_db()
        conn.execute('''
            INSERT INTO messages (username, message, room)
            VALUES (?, ?, 'genel')
        ''', (username, message))
        conn.commit()
        conn.close()

        emit('new_message', {
            'username': username,
            'message': message,
            'time': now_tr().strftime('%H:%M')
        }, broadcast=True)


# ==================== MODERN MESAJLAŞMA (DM + GRUP) ====================
# Eski dm_room_name / join_dm / send_dm kaldırıldı. Artık her kullanıcı zaten
# bağlantı anında kendi 'user_{id}' odasına katılıyor (bkz. @socketio.on('connect')),
# bu yüzden ayrıca sohbet odalarına katılmaya gerek yok: sunucu her mesaj/olayı
# ilgili sohbetin TÜM üyelerinin kişisel odasına yollar, istemci ise sadece o an
# açık olan sohbetle eşleşen olayları ekrana basar. Bu hem daha az socket
# karmaşıklığı hem de çoklu cihaz senkronizasyonu (aynı kullanıcı 2 cihazda) sağlar.

MAX_MESSAGE_LEN = 4000


def _normalize_message_text(text):
    """Art arda çok sayıda boş satırı (kullanıcı defalarca Enter'a basınca
    oluşan) en fazla tek bir boş satıra indirger; balonun gereksiz yere
    dikey olarak şişmesini engeller. Tek/çift satır atlamalarına dokunmaz."""
    if not text:
        return text
    return re.sub(r'\n{3,}', '\n\n', text)
_typing_state = {}  # (conversation_id) -> {user_id: last_ping_ts}  (bellek içi, kalıcı değil)


def _conversation_member_ids(conn, conversation_id):
    rows = conn.execute(
        'SELECT user_id FROM conversation_members WHERE conversation_id = ? AND left_at IS NULL',
        (conversation_id,)
    ).fetchall()
    return [r['user_id'] for r in rows]


def _emit_to_members(event, payload, member_ids, exclude_user_id=None):
    for uid in member_ids:
        if exclude_user_id is not None and uid == exclude_user_id:
            continue
        socketio.emit(event, payload, room=f'user_{uid}')


def _conversation_summary(conn, conversation_id, viewer_id):
    """Sohbet listesi kartı için tek bir sohbetin özet verisini üretir."""
    conv = conn.execute('SELECT * FROM conversations WHERE id = ?', (conversation_id,)).fetchone()
    if not conv:
        return None
    membership = get_conversation_membership(conn, conversation_id, viewer_id)
    if not membership:
        return None

    last_msg = conn.execute('''
        SELECT * FROM conv_messages
        WHERE conversation_id = ? AND id NOT IN (SELECT message_id FROM message_hides WHERE user_id = ?)
        ORDER BY id DESC LIMIT 1
    ''', (conversation_id, viewer_id)).fetchone()

    unread = conn.execute('''
        SELECT COUNT(*) c FROM conv_messages
        WHERE conversation_id = ? AND sender_id != ? AND id > ?
          AND deleted_for_everyone = 0
          AND id NOT IN (SELECT message_id FROM message_hides WHERE user_id = ?)
    ''', (conversation_id, viewer_id, membership['last_read_message_id'], viewer_id)).fetchone()['c']

    if conv['type'] == 'dm':
        other_row = conn.execute('''
            SELECT u.* FROM conversation_members cm JOIN users u ON u.id = cm.user_id
            WHERE cm.conversation_id = ? AND cm.user_id != ? LIMIT 1
        ''', (conversation_id, viewer_id)).fetchone()
        title = other_row['full_name'] if other_row else 'Silinmiş kullanıcı'
        photo = other_row['profile_photo'] if other_row else None
        subtitle_id = other_row['id'] if other_row else None
        is_blocked = other_row and is_blocked_between(conn, viewer_id, other_row['id'])
    else:
        title = conv['name'] or 'Grup'
        photo = conv['photo']
        subtitle_id = None
        is_blocked = False

    last_preview = None
    if last_msg:
        if last_msg['deleted_for_everyone']:
            last_preview = 'Bu mesaj silindi.'
        else:
            prefix = 'Sen: ' if last_msg['sender_id'] == viewer_id else ''
            if conv['type'] == 'group' and last_msg['sender_id'] != viewer_id:
                sname = conn.execute('SELECT full_name FROM users WHERE id=?', (last_msg['sender_id'],)).fetchone()
                prefix = f"{(sname['full_name'].split(' ')[0] if sname else '')}: "
            body = last_msg['content'] if last_msg['content'] else ('📷 Fotoğraf' if last_msg['image'] else '')
            last_preview = (prefix + body).strip()

    last_dt = utc_to_tr(last_msg['created_at']) if last_msg else utc_to_tr(conv['created_at'])
    last_time_display = ''
    if last_dt:
        today = now_tr().date()
        if last_dt.date() == today:
            last_time_display = last_dt.strftime('%H:%M')
        elif (today - last_dt.date()).days == 1:
            last_time_display = 'Dün'
        elif (today - last_dt.date()).days < 7:
            gunler = ['Pzt', 'Sal', 'Çar', 'Per', 'Cum', 'Cmt', 'Paz']
            last_time_display = gunler[last_dt.weekday()]
        else:
            last_time_display = last_dt.strftime('%d.%m.%Y')

    return {
        'conversation_id': conversation_id,
        'type': conv['type'],
        'title': title,
        'photo': photo,
        'other_user_id': subtitle_id,
        'last_message': last_preview,
        'last_time_display': last_time_display,
        'last_sort_time': last_msg['created_at'] if last_msg else conv['created_at'],
        'unread': unread,
        'is_pinned': bool(membership['is_pinned']),
        'is_muted': bool(membership['is_muted']),
        'is_archived': bool(membership['is_archived']),
        'is_blocked': bool(is_blocked),
    }


@app.route('/messages')
@login_required
def messages_list():
    uid = session['user_id']
    conn = get_db()
    conv_ids = [r['conversation_id'] for r in conn.execute('''
        SELECT conversation_id FROM conversation_members
        WHERE user_id = ? AND left_at IS NULL AND hidden_at IS NULL
    ''', (uid,)).fetchall()]

    summaries = [s for s in (_conversation_summary(conn, cid, uid) for cid in conv_ids) if s]
    summaries.sort(key=lambda s: (not s['is_pinned'], s['last_sort_time'] or ''), reverse=True)
    active_list = [s for s in summaries if not s['is_archived']]
    archived_list = [s for s in summaries if s['is_archived']]

    all_users = conn.execute('''
        SELECT id, username, full_name, profile_photo FROM users
        WHERE id != ? ORDER BY full_name ASC
    ''', (uid,)).fetchall()
    conn.close()

    return render_template('messages.html', conversations=active_list,
                            archived_conversations=archived_list, all_users=all_users)


@app.route('/messages/<int:other_id>')
@login_required
def messages_thread(other_id):
    """Geriye dönük uyumluluk: profil / oda arkadaşı sayfalarındaki
    'Mesaj gönder' linkleri hâlâ kullanıcı id'si ile buraya yönlendirir.
    Burada ilgili DM sohbeti bulunur/oluşturulur ve grup şablonuyla aynı
    thread arayüzü render edilir."""
    uid = session['user_id']
    conn = get_db()
    other = conn.execute('SELECT * FROM users WHERE id = ?', (other_id,)).fetchone()
    if not other:
        conn.close()
        flash('Kullanıcı bulunamadı.', 'danger')
        return redirect(url_for('messages_list'))
    if other_id == uid:
        conn.close()
        flash('Kendinize mesaj gönderemezsiniz.', 'danger')
        return redirect(url_for('messages_list'))

    conv_id = get_or_create_dm_conversation(conn, uid, other_id)
    # Kullanıcı bu sohbeti daha önce kendi tarafından silmişse (hidden_at),
    # yeniden mesajlaşmaya başlarken görünür hale getir.
    conn.execute('UPDATE conversation_members SET hidden_at = NULL WHERE conversation_id = ? AND user_id = ?',
                 (conv_id, uid))
    conn.commit()
    return _render_thread(conn, conv_id, uid)


@app.route('/messages/group/<int:conversation_id>')
@login_required
def group_thread(conversation_id):
    uid = session['user_id']
    conn = get_db()
    conv = conn.execute('SELECT * FROM conversations WHERE id = ? AND type = "group"', (conversation_id,)).fetchone()
    if not conv:
        conn.close()
        flash('Grup bulunamadı.', 'danger')
        return redirect(url_for('messages_list'))
    membership = get_conversation_membership(conn, conversation_id, uid)
    if not membership:
        conn.close()
        flash('Bu grubun üyesi değilsiniz.', 'danger')
        return redirect(url_for('messages_list'))
    conn.execute('UPDATE conversation_members SET hidden_at = NULL WHERE conversation_id = ? AND user_id = ?',
                 (conversation_id, uid))
    conn.commit()
    return _render_thread(conn, conversation_id, uid)


def _render_thread(conn, conversation_id, uid):
    conv = conn.execute('SELECT * FROM conversations WHERE id = ?', (conversation_id,)).fetchone()
    membership = get_conversation_membership(conn, conversation_id, uid)
    if not conv or not membership:
        conn.close()
        flash('Bu sohbete erişiminiz yok.', 'danger')
        return redirect(url_for('messages_list'))

    msg_rows = conn.execute('''
        SELECT * FROM conv_messages
        WHERE conversation_id = ? AND id NOT IN (SELECT message_id FROM message_hides WHERE user_id = ?)
        ORDER BY id DESC LIMIT 50
    ''', (conversation_id, uid)).fetchall()
    users_cache = {}
    messages = [conv_message_to_dict(conn, m, uid, users_cache) for m in reversed(msg_rows)]
    has_more = len(msg_rows) == 50

    other = None
    is_blocked_by_me = False
    is_blocking_me = False
    if conv['type'] == 'dm':
        other = conn.execute('''
            SELECT u.* FROM conversation_members cm JOIN users u ON u.id = cm.user_id
            WHERE cm.conversation_id = ? AND cm.user_id != ? LIMIT 1
        ''', (conversation_id, uid)).fetchone()
        if other:
            is_blocked_by_me = conn.execute('SELECT 1 FROM blocked_users WHERE blocker_id=? AND blocked_id=?',
                                             (uid, other['id'])).fetchone() is not None
            is_blocking_me = conn.execute('SELECT 1 FROM blocked_users WHERE blocker_id=? AND blocked_id=?',
                                           (other['id'], uid)).fetchone() is not None

    members = []
    if conv['type'] == 'group':
        members = conn.execute('''
            SELECT u.id, u.username, u.full_name, u.profile_photo, cm.role, cm.last_read_message_id
            FROM conversation_members cm JOIN users u ON u.id = cm.user_id
            WHERE cm.conversation_id = ? AND cm.left_at IS NULL
            ORDER BY (cm.role = 'admin') DESC, u.full_name ASC
        ''', (conversation_id,)).fetchall()

    pinned_row = conn.execute('''
        SELECT mp.message_id, cmsg.content, u.full_name as sender_name
        FROM message_pins mp
        JOIN conv_messages cmsg ON cmsg.id = mp.message_id
        JOIN users u ON u.id = cmsg.sender_id
        WHERE mp.conversation_id = ? ORDER BY mp.pinned_at DESC LIMIT 1
    ''', (conversation_id,)).fetchone()

    # Sohbeti açtım = okundum say + karşı tarafa çift tik bilgisini yolla.
    last_id = msg_rows[0]['id'] if msg_rows else membership['last_read_message_id']
    conn.execute('UPDATE conversation_members SET last_read_message_id = ? WHERE conversation_id = ? AND user_id = ?',
                 (last_id, conversation_id, uid))
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()

    _emit_to_members('messages_read', {'conversation_id': conversation_id, 'reader_id': uid, 'up_to': last_id},
                      member_ids, exclude_user_id=uid)

    return render_template('messages_thread.html',
                            conversation=conv, other=other, messages=messages, has_more=has_more,
                            members=members, my_role=membership['role'], pinned=pinned_row,
                            is_blocked_by_me=is_blocked_by_me, is_blocking_me=is_blocking_me,
                            is_muted=bool(membership['is_muted']))


@app.route('/messages/<int:conversation_id>/older')
@login_required
def messages_older(conversation_id):
    uid = session['user_id']
    conn = get_db()
    if not get_conversation_membership(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'forbidden'}), 403
    before_id = request.args.get('before_id', type=int) or 10**12
    msg_rows = conn.execute('''
        SELECT * FROM conv_messages
        WHERE conversation_id = ? AND id < ? AND id NOT IN (SELECT message_id FROM message_hides WHERE user_id = ?)
        ORDER BY id DESC LIMIT 50
    ''', (conversation_id, before_id, uid)).fetchall()
    users_cache = {}
    messages = [conv_message_to_dict(conn, m, uid, users_cache) for m in reversed(msg_rows)]
    has_more = len(msg_rows) == 50
    conn.close()
    return jsonify({'messages': messages, 'has_more': has_more})


@app.route('/messages/new-group', methods=['GET', 'POST'])
@login_required
def group_new():
    uid = session['user_id']
    conn = get_db()
    if request.method == 'GET':
        all_users = conn.execute('SELECT id, username, full_name, profile_photo FROM users WHERE id != ? ORDER BY full_name ASC',
                                  (uid,)).fetchall()
        conn.close()
        return render_template('group_new.html', all_users=all_users)

    name = (request.form.get('name') or '').strip()
    description = (request.form.get('description') or '').strip()
    member_ids = request.form.getlist('member_ids', type=int)
    member_ids = [m for m in set(member_ids) if m != uid]

    if not name:
        flash('Grup adı boş olamaz.', 'danger')
        conn.close()
        return redirect(url_for('group_new'))
    if len(member_ids) < 1:
        flash('Grup için en az bir üye seçmelisiniz.', 'danger')
        conn.close()
        return redirect(url_for('group_new'))

    valid_users = conn.execute(
        f"SELECT id FROM users WHERE id IN ({','.join('?' * len(member_ids))})", member_ids
    ).fetchall()
    member_ids = [r['id'] for r in valid_users]

    photo_filename = None
    photo_file = request.files.get('photo')
    if photo_file and photo_file.filename and allowed_file(photo_file.filename, IMAGE_EXT):
        photo_filename = save_media(photo_file, UPLOAD_GROUPS, IMAGE_EXT)

    cur = conn.execute(
        'INSERT INTO conversations (type, name, description, photo, created_by) VALUES (?, ?, ?, ?, ?)',
        ('group', name[:80], description[:300], photo_filename, uid)
    )
    conv_id = cur.lastrowid
    conn.execute('INSERT INTO conversation_members (conversation_id, user_id, role) VALUES (?, ?, ?)', (conv_id, uid, 'admin'))
    for mid in member_ids:
        conn.execute('INSERT OR IGNORE INTO conversation_members (conversation_id, user_id, role) VALUES (?, ?, ?)', (conv_id, mid, 'member'))
    sys_text = f"{name} grubu oluşturuldu."
    conn.execute('INSERT INTO conv_messages (conversation_id, sender_id, content, is_system) VALUES (?, ?, ?, 1)', (conv_id, uid, sys_text))
    conn.execute('UPDATE conversations SET last_message_at = CURRENT_TIMESTAMP WHERE id = ?', (conv_id,))
    conn.commit()

    creator = conn.execute('SELECT full_name FROM users WHERE id=?', (uid,)).fetchone()
    conn.close()
    for mid in member_ids:
        socketio.start_background_task(
            push_notification_to_user, mid, f"{creator['full_name']} seni bir gruba ekledi 👥",
            f'"{name}" grubuna eklendin.', url_for('group_thread', conversation_id=conv_id)
        )
    flash('Grup oluşturuldu.', 'success')
    return redirect(url_for('group_thread', conversation_id=conv_id))


def _require_group_admin(conn, conversation_id, uid):
    m = get_conversation_membership(conn, conversation_id, uid)
    return m is not None and m['role'] == 'admin'


@app.route('/messages/group/<int:conversation_id>/update', methods=['POST'])
@login_required
def group_update(conversation_id):
    uid = session['user_id']
    conn = get_db()
    if not _require_group_admin(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'Bu işlem için yönetici olmalısınız.'}), 403
    name = (request.form.get('name') or '').strip()
    description = (request.form.get('description') or '').strip()
    updates = {}
    if name:
        updates['name'] = name[:80]
    if 'description' in request.form:
        updates['description'] = description[:300]
    photo_file = request.files.get('photo')
    if photo_file and photo_file.filename and allowed_file(photo_file.filename, IMAGE_EXT):
        updates['photo'] = save_media(photo_file, UPLOAD_GROUPS, IMAGE_EXT)
    if updates:
        set_clause = ', '.join(f'{k} = ?' for k in updates)
        conn.execute(f'UPDATE conversations SET {set_clause} WHERE id = ?', (*updates.values(), conversation_id))
        conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    _emit_to_members('group_updated', {'conversation_id': conversation_id}, member_ids)
    return jsonify({'ok': True})


@app.route('/messages/group/<int:conversation_id>/add-members', methods=['POST'])
@login_required
def group_add_members(conversation_id):
    uid = session['user_id']
    conn = get_db()
    if not _require_group_admin(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'Bu işlem için yönetici olmalısınız.'}), 403
    new_ids = request.form.getlist('member_ids', type=int)
    conv = conn.execute('SELECT name FROM conversations WHERE id=?', (conversation_id,)).fetchone()
    added = []
    for mid in set(new_ids):
        existing = conn.execute('SELECT * FROM conversation_members WHERE conversation_id=? AND user_id=?', (conversation_id, mid)).fetchone()
        if existing and existing['left_at'] is None:
            continue
        if existing:
            conn.execute('UPDATE conversation_members SET left_at = NULL, role="member" WHERE id = ?', (existing['id'],))
        else:
            user_exists = conn.execute('SELECT 1 FROM users WHERE id=?', (mid,)).fetchone()
            if not user_exists:
                continue
            conn.execute('INSERT INTO conversation_members (conversation_id, user_id, role) VALUES (?, ?, "member")', (conversation_id, mid))
        added.append(mid)
    for mid in added:
        uname = conn.execute('SELECT full_name FROM users WHERE id=?', (mid,)).fetchone()
        conn.execute('INSERT INTO conv_messages (conversation_id, sender_id, content, is_system) VALUES (?, ?, ?, 1)',
                     (conversation_id, uid, f"{uname['full_name']} gruba eklendi."))
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    if added:
        _emit_to_members('group_updated', {'conversation_id': conversation_id}, member_ids)
        for mid in added:
            socketio.start_background_task(
                push_notification_to_user, mid, "Bir gruba eklendin 👥",
                f'"{conv["name"]}" grubuna eklendin.', url_for('group_thread', conversation_id=conversation_id)
            )
    return jsonify({'ok': True, 'added': added})


@app.route('/messages/group/<int:conversation_id>/remove-member/<int:member_id>', methods=['POST'])
@login_required
def group_remove_member(conversation_id, member_id):
    uid = session['user_id']
    conn = get_db()
    if not _require_group_admin(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'Bu işlem için yönetici olmalısınız.'}), 403
    if member_id == uid:
        conn.close()
        return jsonify({'error': "Kendinizi bu şekilde çıkaramazsınız, 'gruptan ayrıl' kullanın."}), 400
    target = get_conversation_membership(conn, conversation_id, member_id)
    if not target:
        conn.close()
        return jsonify({'error': 'Kullanıcı bu grubun üyesi değil.'}), 404
    conn.execute('UPDATE conversation_members SET left_at = CURRENT_TIMESTAMP WHERE id = ?', (target['id'],))
    uname = conn.execute('SELECT full_name FROM users WHERE id=?', (member_id,)).fetchone()
    conn.execute('INSERT INTO conv_messages (conversation_id, sender_id, content, is_system) VALUES (?, ?, ?, 1)',
                 (conversation_id, uid, f"{uname['full_name']} gruptan çıkarıldı."))
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    _emit_to_members('group_updated', {'conversation_id': conversation_id}, member_ids)
    socketio.emit('removed_from_group', {'conversation_id': conversation_id}, room=f'user_{member_id}')
    return jsonify({'ok': True})


@app.route('/messages/group/<int:conversation_id>/set-role/<int:member_id>', methods=['POST'])
@login_required
def group_set_role(conversation_id, member_id):
    uid = session['user_id']
    new_role = request.form.get('role')
    if new_role not in ('admin', 'member'):
        return jsonify({'error': 'Geçersiz rol.'}), 400
    conn = get_db()
    if not _require_group_admin(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'Bu işlem için yönetici olmalısınız.'}), 403
    target = get_conversation_membership(conn, conversation_id, member_id)
    if not target:
        conn.close()
        return jsonify({'error': 'Kullanıcı bu grubun üyesi değil.'}), 404
    if new_role == 'member' and target['role'] == 'admin':
        admin_count = conn.execute(
            "SELECT COUNT(*) c FROM conversation_members WHERE conversation_id=? AND role='admin' AND left_at IS NULL",
            (conversation_id,)
        ).fetchone()['c']
        if admin_count <= 1:
            conn.close()
            return jsonify({'error': 'Grubun en az bir yöneticisi olmalı.'}), 400
    conn.execute('UPDATE conversation_members SET role = ? WHERE id = ?', (new_role, target['id']))
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    _emit_to_members('group_updated', {'conversation_id': conversation_id}, member_ids)
    return jsonify({'ok': True})


@app.route('/messages/group/<int:conversation_id>/leave', methods=['POST'])
@login_required
def group_leave(conversation_id):
    uid = session['user_id']
    conn = get_db()
    membership = get_conversation_membership(conn, conversation_id, uid)
    if not membership:
        conn.close()
        return jsonify({'error': 'Bu grubun üyesi değilsiniz.'}), 404
    conn.execute('UPDATE conversation_members SET left_at = CURRENT_TIMESTAMP WHERE id = ?', (membership['id'],))
    uname = conn.execute('SELECT full_name FROM users WHERE id=?', (uid,)).fetchone()

    # Son yönetici ayrılıyorsa, kalan en kıdemli üyeyi otomatik yönetici yap
    # (edge case: "gruptaki son yönetici ayrılırsa" grup yönetimsiz kalmasın).
    if membership['role'] == 'admin':
        remaining_admins = conn.execute(
            "SELECT COUNT(*) c FROM conversation_members WHERE conversation_id=? AND role='admin' AND left_at IS NULL AND user_id != ?",
            (conversation_id, uid)
        ).fetchone()['c']
        if remaining_admins == 0:
            next_member = conn.execute(
                "SELECT id FROM conversation_members WHERE conversation_id=? AND left_at IS NULL AND user_id != ? ORDER BY joined_at ASC LIMIT 1",
                (conversation_id, uid)
            ).fetchone()
            if next_member:
                conn.execute("UPDATE conversation_members SET role='admin' WHERE id=?", (next_member['id'],))

    conn.execute('INSERT INTO conv_messages (conversation_id, sender_id, content, is_system) VALUES (?, ?, ?, 1)',
                 (conversation_id, uid, f"{uname['full_name']} gruptan ayrıldı."))
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    _emit_to_members('group_updated', {'conversation_id': conversation_id}, member_ids)
    return jsonify({'ok': True, 'redirect': url_for('messages_list')})


# -------- Sohbet (DM veya grup) düzeyi işlemler --------

@app.route('/messages/conversation/<int:conversation_id>/<action>', methods=['POST'])
@login_required
def conversation_action(conversation_id, action):
    uid = session['user_id']
    conn = get_db()
    membership = get_conversation_membership(conn, conversation_id, uid)
    if not membership:
        conn.close()
        return jsonify({'error': 'Bu sohbete erişiminiz yok.'}), 403

    if action == 'pin':
        conn.execute('UPDATE conversation_members SET is_pinned = NOT is_pinned WHERE id = ?', (membership['id'],))
    elif action == 'mute':
        conn.execute('UPDATE conversation_members SET is_muted = NOT is_muted WHERE id = ?', (membership['id'],))
    elif action == 'archive':
        conn.execute('UPDATE conversation_members SET is_archived = NOT is_archived WHERE id = ?', (membership['id'],))
    elif action == 'read':
        last = conn.execute('SELECT MAX(id) m FROM conv_messages WHERE conversation_id=?', (conversation_id,)).fetchone()['m'] or 0
        conn.execute('UPDATE conversation_members SET last_read_message_id = ? WHERE id = ?', (last, membership['id']))
    elif action == 'delete':
        # "Sohbeti sil": yalnızca benim görünümümden kaldırır, karşı tarafı etkilemez.
        conn.execute('UPDATE conversation_members SET hidden_at = CURRENT_TIMESTAMP WHERE id = ?', (membership['id'],))
    else:
        conn.close()
        return jsonify({'error': 'Geçersiz işlem.'}), 400

    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/messages/conversation/<int:conversation_id>/block', methods=['POST'])
@login_required
def conversation_block(conversation_id):
    uid = session['user_id']
    conn = get_db()
    conv = conn.execute('SELECT * FROM conversations WHERE id=? AND type="dm"', (conversation_id,)).fetchone()
    membership = get_conversation_membership(conn, conversation_id, uid)
    if not conv or not membership:
        conn.close()
        return jsonify({'error': 'Bu işlem yalnızca birebir sohbetlerde yapılabilir.'}), 400
    other = conn.execute('SELECT user_id FROM conversation_members WHERE conversation_id=? AND user_id != ?',
                          (conversation_id, uid)).fetchone()
    if not other:
        conn.close()
        return jsonify({'error': 'Kullanıcı bulunamadı.'}), 404
    already = conn.execute('SELECT 1 FROM blocked_users WHERE blocker_id=? AND blocked_id=?', (uid, other['user_id'])).fetchone()
    if already:
        conn.execute('DELETE FROM blocked_users WHERE blocker_id=? AND blocked_id=?', (uid, other['user_id']))
        blocked_now = False
    else:
        conn.execute('INSERT INTO blocked_users (blocker_id, blocked_id) VALUES (?, ?)', (uid, other['user_id']))
        blocked_now = True
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'blocked': blocked_now})


@app.route('/messages/conversation/<int:conversation_id>/report', methods=['POST'])
@login_required
def conversation_report(conversation_id):
    uid = session['user_id']
    reason = (request.form.get('reason') or '').strip()[:500]
    message_id = request.form.get('message_id', type=int)
    conn = get_db()
    if not get_conversation_membership(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'Bu sohbete erişiminiz yok.'}), 403
    conn.execute('INSERT INTO message_reports (message_id, conversation_id, reported_by, reason) VALUES (?, ?, ?, ?)',
                 (message_id, conversation_id, uid, reason))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/messages/conversation/<int:conversation_id>/pin-message/<int:message_id>', methods=['POST'])
@login_required
def pin_message(conversation_id, message_id):
    uid = session['user_id']
    conn = get_db()
    if not get_conversation_membership(conn, conversation_id, uid):
        conn.close()
        return jsonify({'error': 'Bu sohbete erişiminiz yok.'}), 403
    msg = conn.execute('SELECT * FROM conv_messages WHERE id=? AND conversation_id=?', (message_id, conversation_id)).fetchone()
    if not msg or msg['deleted_for_everyone']:
        conn.close()
        return jsonify({'error': 'Mesaj bulunamadı.'}), 404
    already = conn.execute('SELECT 1 FROM message_pins WHERE conversation_id=? AND message_id=?', (conversation_id, message_id)).fetchone()
    if already:
        conn.execute('DELETE FROM message_pins WHERE conversation_id=? AND message_id=?', (conversation_id, message_id))
        pinned_now = False
    else:
        conn.execute('DELETE FROM message_pins WHERE conversation_id=?', (conversation_id,))  # tek sabit mesaj
        conn.execute('INSERT INTO message_pins (conversation_id, message_id, pinned_by) VALUES (?, ?, ?)', (conversation_id, message_id, uid))
        pinned_now = True
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    _emit_to_members('group_updated', {'conversation_id': conversation_id}, member_ids)
    return jsonify({'ok': True, 'pinned': pinned_now})


# -------- Mesaj düzeyi işlemler --------

@app.route('/messages/message/<int:message_id>/edit', methods=['POST'])
@login_required
def message_edit(message_id):
    uid = session['user_id']
    new_content = _normalize_message_text((request.form.get('content') or '').strip())
    conn = get_db()
    msg = conn.execute('SELECT * FROM conv_messages WHERE id = ?', (message_id,)).fetchone()
    if not msg or msg['deleted_for_everyone']:
        conn.close()
        return jsonify({'error': 'Mesaj bulunamadı.'}), 404
    if msg['sender_id'] != uid:
        conn.close()
        return jsonify({'error': 'Yalnızca kendi mesajınızı düzenleyebilirsiniz.'}), 403
    if not get_conversation_membership(conn, msg['conversation_id'], uid):
        conn.close()
        return jsonify({'error': 'Bu sohbete erişiminiz yok.'}), 403
    if not new_content:
        conn.close()
        return jsonify({'error': 'Mesaj boş olamaz.'}), 400
    if len(new_content) > MAX_MESSAGE_LEN:
        conn.close()
        return jsonify({'error': 'Mesaj çok uzun.'}), 400

    conn.execute('UPDATE conv_messages SET content = ?, edited_at = CURRENT_TIMESTAMP WHERE id = ?', (new_content, message_id))
    conn.commit()
    member_ids = _conversation_member_ids(conn, msg['conversation_id'])
    conn.close()
    _emit_to_members('message_edited', {
        'conversation_id': msg['conversation_id'], 'message_id': message_id, 'content': new_content
    }, member_ids)
    return jsonify({'ok': True})


@app.route('/messages/message/<int:message_id>/delete-for-everyone', methods=['POST'])
@login_required
def message_delete_everyone(message_id):
    uid = session['user_id']
    conn = get_db()
    msg = conn.execute('SELECT * FROM conv_messages WHERE id = ?', (message_id,)).fetchone()
    if not msg:
        conn.close()
        return jsonify({'error': 'Mesaj bulunamadı.'}), 404
    membership = get_conversation_membership(conn, msg['conversation_id'], uid)
    if not membership:
        conn.close()
        return jsonify({'error': 'Bu sohbete erişiminiz yok.'}), 403
    # Yetki: mesaj sahibi HER ZAMAN silebilir; grup yöneticisi başkasının
    # mesajını da kaldırabilir. Bu kontrol backend'de yapılır (IDOR koruması).
    if msg['sender_id'] != uid and membership['role'] != 'admin':
        conn.close()
        return jsonify({'error': 'Bu mesajı herkesten silme yetkiniz yok.'}), 403

    conn.execute('UPDATE conv_messages SET deleted_for_everyone = 1, deleted_at = CURRENT_TIMESTAMP, content = NULL WHERE id = ?', (message_id,))
    conn.commit()
    member_ids = _conversation_member_ids(conn, msg['conversation_id'])
    conn.close()
    _emit_to_members('message_deleted', {
        'conversation_id': msg['conversation_id'], 'message_id': message_id, 'scope': 'everyone'
    }, member_ids)
    return jsonify({'ok': True})


@app.route('/messages/message/<int:message_id>/delete-for-me', methods=['POST'])
@login_required
def message_delete_me(message_id):
    uid = session['user_id']
    conn = get_db()
    msg = conn.execute('SELECT * FROM conv_messages WHERE id = ?', (message_id,)).fetchone()
    if not msg:
        conn.close()
        return jsonify({'error': 'Mesaj bulunamadı.'}), 404
    if not get_conversation_membership(conn, msg['conversation_id'], uid):
        conn.close()
        return jsonify({'error': 'Bu sohbete erişiminiz yok.'}), 403
    conn.execute('INSERT OR IGNORE INTO message_hides (message_id, user_id) VALUES (?, ?)', (message_id, uid))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


def _validate_send_target(conn, uid, conversation_id):
    """Mesaj/görsel göndermeden önce ortak kontroller: üyelik ve engelleme.
    Hem metin hem görsel gönderme yolları bu fonksiyonu paylaşır."""
    membership = get_conversation_membership(conn, conversation_id, uid)
    if not membership:
        return None, {'error': 'Bu sohbete erişiminiz yok.'}, 403
    conv = conn.execute('SELECT * FROM conversations WHERE id = ?', (conversation_id,)).fetchone()
    if not conv:
        return None, {'error': 'Sohbet bulunamadı.'}, 404
    if conv['type'] == 'dm':
        other = conn.execute('SELECT user_id FROM conversation_members WHERE conversation_id=? AND user_id != ?',
                              (conversation_id, uid)).fetchone()
        if other and is_blocked_between(conn, uid, other['user_id']):
            return None, {'error': 'Bu kullanıcıyla mesajlaşamazsınız.'}, 403
    return conv, None, None


def _broadcast_new_message(conn, uid, conversation_id, conv, new_id):
    """Yeni eklenen bir conv_messages satırını tüm üyelere Socket.IO ile
    yollar ve sessize alınmamış üyelere push bildirimi gönderir. Hem metin
    hem görsel mesaj gönderme yolları bu fonksiyonu paylaşır."""
    conn.execute('UPDATE conversations SET last_message_at = CURRENT_TIMESTAMP WHERE id = ?', (conversation_id,))
    conn.execute('UPDATE conversation_members SET last_read_message_id = ? WHERE conversation_id = ? AND user_id = ?',
                 (new_id, conversation_id, uid))
    conn.commit()

    msg_row = conn.execute('SELECT * FROM conv_messages WHERE id = ?', (new_id,)).fetchone()
    payload = conv_message_to_dict(conn, msg_row, uid)
    member_ids = _conversation_member_ids(conn, conversation_id)
    sender = conn.execute('SELECT full_name FROM users WHERE id=?', (uid,)).fetchone()

    for mid in member_ids:
        p = dict(payload)
        p['is_mine'] = (mid == uid)
        socketio.emit('new_message', p, room=f'user_{mid}')

    other_members = [m for m in member_ids if m != uid]
    unmuted_recipients = []
    for mid in other_members:
        member_row = conn.execute('SELECT is_muted FROM conversation_members WHERE conversation_id=? AND user_id=?',
                                   (conversation_id, mid)).fetchone()
        if not member_row or not member_row['is_muted']:
            unmuted_recipients.append(mid)

    notif_body = msg_row['content'] if msg_row['content'] else '📷 Fotoğraf gönderdi'
    notif_title = f"{sender['full_name']} ({conv['name']}) 💬" if conv['type'] == 'group' else f"{sender['full_name']} sana mesaj gönderdi 💬"
    link = url_for('group_thread', conversation_id=conversation_id) if conv['type'] == 'group' else url_for('messages_thread', other_id=uid)
    for mid in unmuted_recipients:
        socketio.start_background_task(push_notification_to_user, mid, notif_title, notif_body[:120], link)

    return payload


@app.route('/messages/conversation/<int:conversation_id>/send', methods=['POST'])
@login_required
def message_send(conversation_id):
    """Mesaj göndermenin REST karşılığı: socket bağlantısı koparsa (örn.
    sayfa yenilendi / internet kısa kesildi) istemci bu uca düşerek mesajın
    kaybolmamasını sağlar. Socket üzerinden gelen 'send_message' olayı da
    aynı doğrulama mantığını paylaşır."""
    uid = session['user_id']
    conn = get_db()
    result, status = _handle_send_message(conn, uid, conversation_id, request.form.get('content'),
                                           request.form.get('reply_to_id', type=int))
    conn.close()
    return jsonify(result), status


def _handle_send_message(conn, uid, conversation_id, content, reply_to_id=None):
    text = _normalize_message_text((content or '').strip())
    if not text:
        return {'error': 'Boş mesaj gönderilemez.'}, 400
    if len(text) > MAX_MESSAGE_LEN:
        return {'error': 'Mesaj çok uzun.'}, 400

    conv, err, status = _validate_send_target(conn, uid, conversation_id)
    if err:
        return err, status

    if reply_to_id:
        rmsg = conn.execute('SELECT id FROM conv_messages WHERE id=? AND conversation_id=?', (reply_to_id, conversation_id)).fetchone()
        if not rmsg:
            reply_to_id = None

    cur = conn.execute(
        'INSERT INTO conv_messages (conversation_id, sender_id, content, reply_to_id) VALUES (?, ?, ?, ?)',
        (conversation_id, uid, text, reply_to_id)
    )
    new_id = cur.lastrowid
    payload = _broadcast_new_message(conn, uid, conversation_id, conv, new_id)
    return {'ok': True, 'message': payload}, 200


MAX_IMAGE_SIZE_BYTES = 8 * 1024 * 1024  # 8 MB


@app.route('/messages/conversation/<int:conversation_id>/send-image', methods=['POST'])
@login_required
def message_send_image(conversation_id):
    """Görsel gönderme: metinle aynı yetkilendirme ve engelleme kontrollerini
    paylaşır, ardından dosyayı diske kaydedip yeni bir mesaj olarak yayınlar."""
    uid = session['user_id']
    conn = get_db()

    image_file = request.files.get('image')
    if not image_file or not image_file.filename:
        conn.close()
        return jsonify({'error': 'Görsel bulunamadı.'}), 400
    if not allowed_file(image_file.filename, IMAGE_EXT):
        conn.close()
        return jsonify({'error': 'Yalnızca resim dosyaları gönderilebilir (png, jpg, jpeg, gif, webp).'}), 400

    image_file.stream.seek(0, os.SEEK_END)
    size = image_file.stream.tell()
    image_file.stream.seek(0)
    if size > MAX_IMAGE_SIZE_BYTES:
        conn.close()
        return jsonify({'error': 'Görsel çok büyük (maksimum 8 MB).'}), 400

    conv, err, status = _validate_send_target(conn, uid, conversation_id)
    if err:
        conn.close()
        return jsonify(err), status

    caption = _normalize_message_text((request.form.get('content') or '').strip()[:MAX_MESSAGE_LEN]) or None
    reply_to_id = request.form.get('reply_to_id', type=int)
    if reply_to_id:
        rmsg = conn.execute('SELECT id FROM conv_messages WHERE id=? AND conversation_id=?', (reply_to_id, conversation_id)).fetchone()
        if not rmsg:
            reply_to_id = None

    filename = save_media(image_file, UPLOAD_MESSAGES, IMAGE_EXT)
    cur = conn.execute(
        'INSERT INTO conv_messages (conversation_id, sender_id, content, image, reply_to_id) VALUES (?, ?, ?, ?, ?)',
        (conversation_id, uid, caption, filename, reply_to_id)
    )
    new_id = cur.lastrowid
    payload = _broadcast_new_message(conn, uid, conversation_id, conv, new_id)
    conn.close()
    return jsonify({'ok': True, 'message': payload}), 200


@socketio.on('send_message')
def handle_send_message(data):
    if 'user_id' not in session:
        return
    conversation_id = data.get('conversation_id')
    if not conversation_id:
        return
    conn = get_db()
    _handle_send_message(conn, session['user_id'], conversation_id, data.get('content'), data.get('reply_to_id'))
    conn.close()


@socketio.on('typing')
def handle_typing(data):
    if 'user_id' not in session:
        return
    conversation_id = data.get('conversation_id')
    is_typing = bool(data.get('is_typing'))
    if not conversation_id:
        return
    conn = get_db()
    if not get_conversation_membership(conn, conversation_id, session['user_id']):
        conn.close()
        return
    member_ids = _conversation_member_ids(conn, conversation_id)
    sender = conn.execute('SELECT full_name FROM users WHERE id=?', (session['user_id'],)).fetchone()
    conn.close()
    _emit_to_members('user_typing', {
        'conversation_id': conversation_id, 'user_id': session['user_id'],
        'user_name': sender['full_name'].split(' ')[0] if sender else 'Biri',
        'is_typing': is_typing
    }, member_ids, exclude_user_id=session['user_id'])


@socketio.on('mark_read')
def handle_mark_read(data):
    if 'user_id' not in session:
        return
    conversation_id = data.get('conversation_id')
    if not conversation_id:
        return
    conn = get_db()
    uid = session['user_id']
    if not get_conversation_membership(conn, conversation_id, uid):
        conn.close()
        return
    last = conn.execute('SELECT MAX(id) m FROM conv_messages WHERE conversation_id=?', (conversation_id,)).fetchone()['m'] or 0
    conn.execute('UPDATE conversation_members SET last_read_message_id = ? WHERE conversation_id = ? AND user_id = ?',
                 (last, conversation_id, uid))
    conn.commit()
    member_ids = _conversation_member_ids(conn, conversation_id)
    conn.close()
    _emit_to_members('messages_read', {'conversation_id': conversation_id, 'reader_id': uid, 'up_to': last},
                      member_ids, exclude_user_id=uid)


@app.route('/messages/search-users')
@login_required
def messages_search_users():
    q = (request.args.get('q') or '').strip()
    uid = session['user_id']
    conn = get_db()
    if q:
        rows = conn.execute('''
            SELECT id, username, full_name, profile_photo FROM users
            WHERE id != ? AND (full_name LIKE ? OR username LIKE ?)
            ORDER BY full_name ASC LIMIT 30
        ''', (uid, f'%{q}%', f'%{q}%')).fetchall()
    else:
        rows = conn.execute('SELECT id, username, full_name, profile_photo FROM users WHERE id != ? ORDER BY full_name ASC LIMIT 30',
                             (uid,)).fetchall()
    conn.close()
    return jsonify({'users': [dict(r) for r in rows]})


# ==================== KAHOOT TARZI YARIŞMA (QUIZ) ====================

# ==================== YARIŞMALAR (Kahoot tarzı canlı quiz) ====================
# Editör: soru, cevaplar ve doğru şık TEK ekranda düzenlenir, otomatik kaydedilir.
# Oyun akışı sunucu tarafından yönetilir: soru -> (süre dolunca / herkes cevaplayınca)
# sonuç -> animasyonlu sıralama -> sonraki soru ... -> podyum.

QUIZ_TIME_CHOICES = [5, 10, 20, 30, 45, 60, 90, 120]
QUIZ_POINTS_MODES = OrderedDict([('standard', ('Standart', 1)), ('double', ('Çift Puan', 2)), ('none', ('Puansız', 0))])
QUIZ_TYPES = OrderedDict([('quiz', 'Çoktan Seçmeli'), ('truefalse', 'Doğru / Yanlış')])
QUIZ_INTRO_SECONDS = 4        # Soru metni önce tek başına gösterilir ("Hazırlan")
QUIZ_REVEAL_SECONDS = 5       # Doğru cevap + dağılım ekranı, ardından otomatik sıralama
QUIZ_AUTO_NEXT_SECONDS = 8    # "Otomatik ilerle" açıksa sıralamadan sonra bekleme
QUIZ_OPTION_KEYS = ('a', 'b', 'c', 'd')

QUIZ_LIVE = {}                # session_id -> canlı durum (bellekte)
QUIZ_LIVE_LOCK = threading.Lock()


def _init_quiz_extras(conn):
    _ensure_column(conn, 'quiz_questions', 'qtype', "TEXT DEFAULT 'quiz'")
    _ensure_column(conn, 'quiz_questions', 'points_mode', "TEXT DEFAULT 'standard'")
    _ensure_column(conn, 'quizzes', 'updated_at', 'TIMESTAMP')
    _ensure_column(conn, 'quizzes', 'cover_theme', "TEXT DEFAULT 'purple'")
    _ensure_column(conn, 'quiz_participants', 'streak', 'INTEGER DEFAULT 0')
    _ensure_column(conn, 'quiz_participants', 'best_streak', 'INTEGER DEFAULT 0')
    _ensure_column(conn, 'quiz_participants', 'correct_count', 'INTEGER DEFAULT 0')
    _ensure_column(conn, 'quiz_sessions', 'auto_advance', 'INTEGER DEFAULT 0')
    _ensure_column(conn, 'quiz_sessions', 'finished_at', 'TIMESTAMP')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_quiz_questions_quiz ON quiz_questions(quiz_id, order_index)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_quiz_answers_session ON quiz_answers(session_id, question_id)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_quiz_participants_session ON quiz_participants(session_id)')


def _can_edit_quiz(quiz):
    return bool(quiz) and (quiz['created_by'] == session.get('user_id') or has_perm('manage_quizzes'))


def _quiz_filled_options(q):
    return [k for k in QUIZ_OPTION_KEYS if (q['option_' + k] or '').strip() or q['option_' + k + '_image']]


def _quiz_question_issues(q):
    issues = []
    if not (q['question_text'] or '').strip():
        issues.append('Soru metni eksik')
    filled = _quiz_filled_options(q)
    if len(filled) < 2:
        issues.append('En az 2 cevap gerekli')
    if q['correct_option'] not in filled:
        issues.append('Doğru cevabı işaretle')
    return issues


def _quiz_q_dict(q, with_answer=True):
    keys = q.keys()
    d = {
        'id': q['id'],
        'question_text': q['question_text'] or '',
        'time_limit': q['time_limit'] or 20,
        'qtype': (q['qtype'] if 'qtype' in keys else None) or 'quiz',
        'points_mode': (q['points_mode'] if 'points_mode' in keys else None) or 'standard',
        'image_path': q['image_path'] if 'image_path' in keys else None,
        'audio_path': q['audio_path'] if 'audio_path' in keys else None,
        'order_index': q['order_index'],
    }
    for k in QUIZ_OPTION_KEYS:
        d['option_' + k] = q['option_' + k] or ''
        d['option_' + k + '_image'] = q['option_' + k + '_image'] if ('option_' + k + '_image') in keys else None
    if with_answer:
        d['correct_option'] = q['correct_option']
        d['issues'] = _quiz_question_issues(q)
    return d


def _quiz_questions(conn, quiz_id):
    return conn.execute('SELECT * FROM quiz_questions WHERE quiz_id = ? ORDER BY order_index ASC, id ASC', (quiz_id,)).fetchall()


def _touch_quiz(conn, quiz_id):
    conn.execute('UPDATE quizzes SET updated_at = CURRENT_TIMESTAMP WHERE id = ?', (quiz_id,))


def _remove_quiz_file(name):
    if name:
        try:
            os.remove(os.path.join(UPLOAD_QUIZ, name))
        except OSError:
            pass


@app.route('/quiz')
@login_required
def quiz_list():
    conn = get_db()
    summary_sql = '''
        SELECT q.*, u.full_name AS owner_name,
               (SELECT COUNT(*) FROM quiz_questions WHERE quiz_id = q.id) as question_count,
               (SELECT COALESCE(SUM(time_limit), 0) FROM quiz_questions WHERE quiz_id = q.id) as total_seconds,
               (SELECT image_path FROM quiz_questions WHERE quiz_id = q.id AND image_path IS NOT NULL
                    ORDER BY order_index LIMIT 1) as cover_image,
               (SELECT COUNT(*) FROM quiz_sessions WHERE quiz_id = q.id) as session_count,
               (SELECT COUNT(*) FROM quiz_sessions WHERE quiz_id = q.id AND status = 'finished') as finished_count,
               (SELECT COUNT(*) FROM quiz_participants qp JOIN quiz_sessions s2 ON qp.session_id = s2.id
                    WHERE s2.quiz_id = q.id) as total_participants
        FROM quizzes q JOIN users u ON u.id = q.created_by
    '''
    my_quizzes = conn.execute(summary_sql + ' WHERE q.created_by = ? ORDER BY COALESCE(q.updated_at, q.created_at) DESC',
                              (session['user_id'],)).fetchall()
    other_quizzes = []
    if has_perm('manage_quizzes'):
        other_quizzes = conn.execute(summary_sql + ' WHERE q.created_by != ? ORDER BY COALESCE(q.updated_at, q.created_at) DESC LIMIT 60',
                                     (session['user_id'],)).fetchall()
    active_sessions = conn.execute('''
        SELECT s.*, q.title as quiz_title,
               (SELECT COUNT(*) FROM quiz_participants WHERE session_id = s.id) as participant_count
        FROM quiz_sessions s JOIN quizzes q ON s.quiz_id = q.id
        WHERE s.host_id = ? AND s.status != 'finished'
        ORDER BY s.created_at DESC
    ''', (session['user_id'],)).fetchall()
    joined = conn.execute('''
        SELECT s.id, s.code, s.status, q.title as quiz_title, p.score
        FROM quiz_participants p JOIN quiz_sessions s ON s.id = p.session_id JOIN quizzes q ON q.id = s.quiz_id
        WHERE p.user_id = ? AND s.status != 'finished' ORDER BY s.created_at DESC LIMIT 5
    ''', (session['user_id'],)).fetchall()
    conn.close()
    live_session_by_quiz = {}
    for s in active_sessions:
        live_session_by_quiz.setdefault(s['quiz_id'], s['id'])
    return render_template('quiz_list.html', my_quizzes=my_quizzes, other_quizzes=other_quizzes,
                           active_sessions=active_sessions, joined_sessions=joined,
                           live_session_by_quiz=live_session_by_quiz)


@app.route('/quiz/create', methods=['GET', 'POST'])
@login_required
def quiz_create():
    """Kahoot gibi: tek tıkla boş bir yarışma + ilk soru oluşturulur ve editör açılır."""
    title = (request.form.get('title') or request.args.get('title') or '').strip()[:120] or 'Başlıksız Yarışma'
    description = (request.form.get('description') or '').strip()[:300]
    conn = get_db()
    cur = conn.cursor()
    cur.execute('INSERT INTO quizzes (title, description, created_by, updated_at) VALUES (?, ?, ?, CURRENT_TIMESTAMP)',
                (title, description, session['user_id']))
    quiz_id = cur.lastrowid
    conn.execute('''INSERT INTO quiz_questions (quiz_id, question_text, option_a, option_b, option_c, option_d,
                                                correct_option, time_limit, order_index, qtype, points_mode)
                    VALUES (?, '', '', '', '', '', 'a', 20, 0, 'quiz', 'standard')''', (quiz_id,))
    conn.commit()
    conn.close()
    return redirect(url_for('quiz_detail', quiz_id=quiz_id, new=1))


@app.route('/quiz/<int:quiz_id>')
@login_required
def quiz_detail(quiz_id):
    conn = get_db()
    quiz = conn.execute('SELECT * FROM quizzes WHERE id = ?', (quiz_id,)).fetchone()
    if not quiz:
        conn.close()
        flash('Yarışma bulunamadı.', 'danger')
        return redirect(url_for('quiz_list'))
    if not _can_edit_quiz(quiz):
        conn.close()
        flash('Bu yarışmayı düzenleme yetkin yok.', 'danger')
        return redirect(url_for('quiz_list'))
    questions = [_quiz_q_dict(q) for q in _quiz_questions(conn, quiz_id)]
    live_session = conn.execute('''SELECT * FROM quiz_sessions WHERE quiz_id = ? AND status != 'finished'
                                   ORDER BY created_at DESC LIMIT 1''', (quiz_id,)).fetchone()
    stats = conn.execute('''SELECT COUNT(DISTINCT s.id) sessions, COUNT(p.id) players
                            FROM quiz_sessions s LEFT JOIN quiz_participants p ON p.session_id = s.id
                            WHERE s.quiz_id = ?''', (quiz_id,)).fetchone()
    conn.close()
    return render_template('quiz_detail.html', quiz=quiz, questions=questions, live_session=live_session,
                           session_count=stats['sessions'], total_participants=stats['players'],
                           time_choices=QUIZ_TIME_CHOICES,
                           points_modes=[(k, v[0]) for k, v in QUIZ_POINTS_MODES.items()],
                           quiz_types=list(QUIZ_TYPES.items()),
                           focus_question=request.args.get('q', type=int))


def _quiz_for_edit(conn, quiz_id):
    quiz = conn.execute('SELECT * FROM quizzes WHERE id = ?', (quiz_id,)).fetchone()
    if not quiz:
        return None, (jsonify({'ok': False, 'error': 'Yarışma bulunamadı.'}), 404)
    if not _can_edit_quiz(quiz):
        return None, (jsonify({'ok': False, 'error': 'Bu yarışmayı düzenleme yetkin yok.'}), 403)
    return quiz, None


def _question_for_edit(conn, question_id):
    q = conn.execute('''SELECT qq.*, qz.created_by FROM quiz_questions qq JOIN quizzes qz ON qz.id = qq.quiz_id
                        WHERE qq.id = ?''', (question_id,)).fetchone()
    if not q:
        return None, (jsonify({'ok': False, 'error': 'Soru bulunamadı.'}), 404)
    if not _can_edit_quiz(q):
        return None, (jsonify({'ok': False, 'error': 'Bu soruyu düzenleme yetkin yok.'}), 403)
    return q, None


@app.route('/api/quiz/<int:quiz_id>/meta', methods=['POST'])
@login_required
def api_quiz_meta(quiz_id):
    data = request.get_json(silent=True) or {}
    conn = get_db()
    quiz, err = _quiz_for_edit(conn, quiz_id)
    if err:
        conn.close()
        return err
    title = (data.get('title') if data.get('title') is not None else quiz['title']).strip()[:120] or 'Başlıksız Yarışma'
    desc = (data.get('description') if data.get('description') is not None else (quiz['description'] or '')).strip()[:300]
    conn.execute('UPDATE quizzes SET title = ?, description = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
                 (title, desc, quiz_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'title': title, 'description': desc})


@app.route('/api/quiz/<int:quiz_id>/questions', methods=['POST'])
@login_required
def api_quiz_add_question(quiz_id):
    data = request.get_json(silent=True) or {}
    qtype = data.get('qtype') if data.get('qtype') in QUIZ_TYPES else 'quiz'
    conn = get_db()
    quiz, err = _quiz_for_edit(conn, quiz_id)
    if err:
        conn.close()
        return err
    rows = _quiz_questions(conn, quiz_id)
    after_id = data.get('after_id')
    # Sıralamayı 0..n olarak normalize edip yeni soruyu istenen yere yerleştir
    ids = [r['id'] for r in rows]
    insert_at = ids.index(after_id) + 1 if after_id in ids else len(ids)
    last_time = rows[-1]['time_limit'] if rows else 20
    a, b = ('Doğru', 'Yanlış') if qtype == 'truefalse' else ('', '')
    cur = conn.cursor()
    cur.execute('''INSERT INTO quiz_questions (quiz_id, question_text, option_a, option_b, option_c, option_d,
                                               correct_option, time_limit, order_index, qtype, points_mode)
                   VALUES (?, '', ?, ?, '', '', 'a', ?, 0, ?, 'standard')''', (quiz_id, a, b, last_time or 20, qtype))
    new_id = cur.lastrowid
    ids.insert(insert_at, new_id)
    for i, qid in enumerate(ids):
        conn.execute('UPDATE quiz_questions SET order_index = ? WHERE id = ?', (i, qid))
    _touch_quiz(conn, quiz_id)
    conn.commit()
    q = conn.execute('SELECT * FROM quiz_questions WHERE id = ?', (new_id,)).fetchone()
    conn.close()
    return jsonify({'ok': True, 'question': _quiz_q_dict(q), 'order': ids})


@app.route('/api/quiz/question/<int:question_id>', methods=['POST'])
@login_required
def api_quiz_update_question(question_id):
    """Tek bir soruyu (metin, şıklar, doğru cevap, süre, puan, medya) günceller.
    JSON ya da multipart (dosya yüklemeleri için) kabul eder; sadece gönderilen alanlar değişir."""
    conn = get_db()
    q, err = _question_for_edit(conn, question_id)
    if err:
        conn.close()
        return err
    data = request.get_json(silent=True) if request.is_json else request.form.to_dict()
    data = data or {}
    updates = {}

    if 'question_text' in data:
        updates['question_text'] = str(data['question_text'])[:300]
    for k in QUIZ_OPTION_KEYS:
        if 'option_' + k in data:
            updates['option_' + k] = str(data['option_' + k])[:120]
    if data.get('correct_option') in QUIZ_OPTION_KEYS:
        updates['correct_option'] = data['correct_option']
    if 'time_limit' in data:
        try:
            updates['time_limit'] = max(5, min(240, int(data['time_limit'])))
        except (TypeError, ValueError):
            pass
    if data.get('points_mode') in QUIZ_POINTS_MODES:
        updates['points_mode'] = data['points_mode']
    if data.get('qtype') in QUIZ_TYPES and data['qtype'] != (q['qtype'] or 'quiz'):
        updates['qtype'] = data['qtype']
        if data['qtype'] == 'truefalse':
            updates.update({'option_a': 'Doğru', 'option_b': 'Yanlış', 'option_c': '', 'option_d': ''})
            if (updates.get('correct_option') or q['correct_option']) not in ('a', 'b'):
                updates['correct_option'] = 'a'
            for k in ('c', 'd'):
                if q['option_' + k + '_image']:
                    _remove_quiz_file(q['option_' + k + '_image'])
                    updates['option_' + k + '_image'] = None
        else:
            if q['option_a'] == 'Doğru' and q['option_b'] == 'Yanlış':
                updates.update({'option_a': '', 'option_b': ''})

    # Dosya kaldırma bayrakları
    def truthy(v):
        return str(v).lower() in ('1', 'true', 'yes', 'on')
    if truthy(data.get('remove_image')):
        _remove_quiz_file(q['image_path'])
        updates['image_path'] = None
    if truthy(data.get('remove_audio')):
        _remove_quiz_file(q['audio_path'])
        updates['audio_path'] = None
    for k in QUIZ_OPTION_KEYS:
        if truthy(data.get(f'remove_option_{k}_image')):
            _remove_quiz_file(q[f'option_{k}_image'])
            updates[f'option_{k}_image'] = None

    # Dosya yüklemeleri
    f = request.files.get('question_image')
    if f and f.filename:
        if not allowed_file(f.filename, IMAGE_EXT):
            conn.close()
            return jsonify({'ok': False, 'error': 'Soru görseli için PNG, JPG, GIF veya WEBP yükleyin.'}), 400
        _remove_quiz_file(q['image_path'])
        updates['image_path'] = save_media(f, UPLOAD_QUIZ, IMAGE_EXT)
    f = request.files.get('question_audio')
    if f and f.filename:
        if not allowed_file(f.filename, AUDIO_EXT):
            conn.close()
            return jsonify({'ok': False, 'error': 'Ses için mp3, wav, ogg, m4a veya aac yükleyin.'}), 400
        _remove_quiz_file(q['audio_path'])
        updates['audio_path'] = save_media(f, UPLOAD_QUIZ, AUDIO_EXT)
    for k in QUIZ_OPTION_KEYS:
        f = request.files.get(f'option_{k}_image')
        if f and f.filename:
            if not allowed_file(f.filename, IMAGE_EXT):
                conn.close()
                return jsonify({'ok': False, 'error': 'Şık görselleri için resim dosyası yükleyin.'}), 400
            _remove_quiz_file(q[f'option_{k}_image'])
            updates[f'option_{k}_image'] = save_media(f, UPLOAD_QUIZ, IMAGE_EXT)

    if updates:
        cols = ', '.join(f'{c} = ?' for c in updates)
        conn.execute(f'UPDATE quiz_questions SET {cols} WHERE id = ?', list(updates.values()) + [question_id])
        _touch_quiz(conn, q['quiz_id'])
        conn.commit()
    fresh = conn.execute('SELECT * FROM quiz_questions WHERE id = ?', (question_id,)).fetchone()
    conn.close()
    return jsonify({'ok': True, 'question': _quiz_q_dict(fresh)})


@app.route('/api/quiz/question/<int:question_id>/duplicate', methods=['POST'])
@login_required
def api_quiz_duplicate_question(question_id):
    conn = get_db()
    q, err = _question_for_edit(conn, question_id)
    if err:
        conn.close()
        return err
    ids = [r['id'] for r in _quiz_questions(conn, q['quiz_id'])]

    def copy_file(name):
        if not name:
            return None
        src = os.path.join(UPLOAD_QUIZ, name)
        if not os.path.exists(src):
            return None
        new_name = f"{uuid.uuid4().hex}.{name.rsplit('.', 1)[-1]}"
        try:
            import shutil
            shutil.copyfile(src, os.path.join(UPLOAD_QUIZ, new_name))
            return new_name
        except OSError:
            return None

    cur = conn.cursor()
    cur.execute('''INSERT INTO quiz_questions (quiz_id, question_text, option_a, option_b, option_c, option_d, correct_option,
                        time_limit, order_index, image_path, audio_path, option_a_image, option_b_image, option_c_image,
                        option_d_image, qtype, points_mode)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?, ?, ?, ?, ?, ?)''',
                (q['quiz_id'], q['question_text'], q['option_a'], q['option_b'], q['option_c'], q['option_d'],
                 q['correct_option'], q['time_limit'], copy_file(q['image_path']), copy_file(q['audio_path']),
                 copy_file(q['option_a_image']), copy_file(q['option_b_image']), copy_file(q['option_c_image']),
                 copy_file(q['option_d_image']), q['qtype'] or 'quiz', q['points_mode'] or 'standard'))
    new_id = cur.lastrowid
    ids.insert(ids.index(question_id) + 1, new_id)
    for i, qid in enumerate(ids):
        conn.execute('UPDATE quiz_questions SET order_index = ? WHERE id = ?', (i, qid))
    _touch_quiz(conn, q['quiz_id'])
    conn.commit()
    fresh = conn.execute('SELECT * FROM quiz_questions WHERE id = ?', (new_id,)).fetchone()
    conn.close()
    return jsonify({'ok': True, 'question': _quiz_q_dict(fresh), 'order': ids})


@app.route('/api/quiz/question/<int:question_id>/delete', methods=['POST'])
@login_required
def api_quiz_delete_question(question_id):
    conn = get_db()
    q, err = _question_for_edit(conn, question_id)
    if err:
        conn.close()
        return err
    used = conn.execute('SELECT COUNT(*) c FROM quiz_answers WHERE question_id = ?', (question_id,)).fetchone()['c']
    if used:
        conn.execute('DELETE FROM quiz_answers WHERE question_id = ?', (question_id,))
    for col in ('image_path', 'audio_path', 'option_a_image', 'option_b_image', 'option_c_image', 'option_d_image'):
        _remove_quiz_file(q[col])
    conn.execute('DELETE FROM quiz_questions WHERE id = ?', (question_id,))
    ids = [r['id'] for r in _quiz_questions(conn, q['quiz_id'])]
    for i, qid in enumerate(ids):
        conn.execute('UPDATE quiz_questions SET order_index = ? WHERE id = ?', (i, qid))
    _touch_quiz(conn, q['quiz_id'])
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'order': ids})


@app.route('/api/quiz/<int:quiz_id>/reorder', methods=['POST'])
@login_required
def api_quiz_reorder(quiz_id):
    order = (request.get_json(silent=True) or {}).get('order') or []
    conn = get_db()
    quiz, err = _quiz_for_edit(conn, quiz_id)
    if err:
        conn.close()
        return err
    valid = [r['id'] for r in _quiz_questions(conn, quiz_id)]
    order = [int(x) for x in order if str(x).isdigit() and int(x) in valid]
    order += [x for x in valid if x not in order]
    for i, qid in enumerate(order):
        conn.execute('UPDATE quiz_questions SET order_index = ? WHERE id = ?', (i, qid))
    _touch_quiz(conn, quiz_id)
    conn.commit()
    conn.close()
    return jsonify({'ok': True, 'order': order})


@app.route('/quiz/<int:quiz_id>/delete', methods=['POST'])
@login_required
def quiz_delete(quiz_id):
    conn = get_db()
    quiz = conn.execute('SELECT * FROM quizzes WHERE id = ?', (quiz_id,)).fetchone()
    if quiz and _can_edit_quiz(quiz):
        for row in _quiz_questions(conn, quiz_id):
            for col in ('image_path', 'audio_path', 'option_a_image', 'option_b_image', 'option_c_image', 'option_d_image'):
                _remove_quiz_file(row[col])
        session_ids = [r['id'] for r in conn.execute('SELECT id FROM quiz_sessions WHERE quiz_id = ?', (quiz_id,)).fetchall()]
        for sid in session_ids:
            conn.execute('DELETE FROM quiz_answers WHERE session_id = ?', (sid,))
            conn.execute('DELETE FROM quiz_participants WHERE session_id = ?', (sid,))
            QUIZ_LIVE.pop(sid, None)
        conn.execute('DELETE FROM quiz_sessions WHERE quiz_id = ?', (quiz_id,))
        conn.execute('DELETE FROM quiz_questions WHERE quiz_id = ?', (quiz_id,))
        conn.execute('DELETE FROM quizzes WHERE id = ?', (quiz_id,))
        conn.commit()
        flash('Yarışma silindi.', 'success')
    else:
        flash('Yetkisiz işlem.', 'danger')
    conn.close()
    return redirect(url_for('quiz_list'))


@app.route('/quiz/<int:quiz_id>/host', methods=['POST'])
@login_required
def quiz_host_start(quiz_id):
    conn = get_db()
    quiz = conn.execute('SELECT * FROM quizzes WHERE id = ?', (quiz_id,)).fetchone()
    if not quiz or not _can_edit_quiz(quiz):
        conn.close()
        flash('Yetkisiz işlem.', 'danger')
        return redirect(url_for('quiz_list'))
    questions = _quiz_questions(conn, quiz_id)
    if not questions:
        conn.close()
        flash('Önce en az bir soru eklemelisin.', 'danger')
        return redirect(url_for('quiz_detail', quiz_id=quiz_id))
    for i, q in enumerate(questions):
        issues = _quiz_question_issues(q)
        if issues:
            conn.close()
            flash(f'{i + 1}. soru eksik: {", ".join(issues)}.', 'warning')
            return redirect(url_for('quiz_detail', quiz_id=quiz_id, q=q['id']))

    code = generate_quiz_code()
    cur = conn.cursor()
    cur.execute('INSERT INTO quiz_sessions (quiz_id, host_id, code, status) VALUES (?, ?, ?, ?)',
                (quiz_id, session['user_id'], code, 'waiting'))
    session_id = cur.lastrowid
    conn.commit()
    conn.close()
    return redirect(url_for('quiz_host_screen', session_id=session_id))


def _is_quiz_host(qs):
    return bool(qs) and (qs['host_id'] == session.get('user_id') or has_perm('manage_quizzes'))


@app.route('/quiz/host/<int:session_id>')
@login_required
def quiz_host_screen(session_id):
    conn = get_db()
    qs = conn.execute('SELECT * FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
    if not qs or not _is_quiz_host(qs):
        conn.close()
        flash('Bu ekrana erişimin yok.', 'danger')
        return redirect(url_for('quiz_list'))
    quiz = conn.execute('SELECT * FROM quizzes WHERE id = ?', (qs['quiz_id'],)).fetchone()
    q_count = conn.execute('SELECT COUNT(*) c FROM quiz_questions WHERE quiz_id = ?', (qs['quiz_id'],)).fetchone()['c']
    conn.close()
    join_url = request.host_url.rstrip('/') + url_for('quiz_join') + '?code=' + qs['code']
    return render_template('quiz_host.html', quiz_session=qs, quiz=quiz, question_count=q_count, join_url=join_url,
                           join_host=request.host)


@app.route('/quiz/join', methods=['GET', 'POST'])
@login_required
def quiz_join():
    if request.method == 'POST':
        code = re.sub(r'\D', '', request.form.get('code', ''))
        default_nick = (current_user_dict() or {}).get('full_name') or session.get('username', 'Oyuncu')
        nickname = (request.form.get('nickname', '').strip() or default_nick)[:24]
        conn = get_db()
        qs = conn.execute("SELECT * FROM quiz_sessions WHERE code = ? AND status != 'finished'", (code,)).fetchone()
        if not qs:
            conn.close()
            flash('Bu PIN ile açık bir yarışma bulunamadı. Kodu kontrol et.', 'danger')
            return redirect(url_for('quiz_join', code=code))

        existing = conn.execute('SELECT id FROM quiz_participants WHERE session_id = ? AND user_id = ?',
                                (qs['id'], session['user_id'])).fetchone()
        if not existing:
            taken = {r['nickname'].lower() for r in conn.execute(
                'SELECT nickname FROM quiz_participants WHERE session_id = ?', (qs['id'],)).fetchall()}
            base, n = nickname, 2
            while nickname.lower() in taken:
                nickname = f'{base[:21]} {n}'
                n += 1
            conn.execute('INSERT INTO quiz_participants (session_id, user_id, nickname) VALUES (?, ?, ?)',
                         (qs['id'], session['user_id'], nickname))
            conn.commit()
        session_id = qs['id']
        conn.close()
        return redirect(url_for('quiz_play', session_id=session_id))
    return render_template('quiz_join.html', prefill_code=re.sub(r'\D', '', request.args.get('code', ''))[:6])


@app.route('/quiz/play/<int:session_id>')
@login_required
def quiz_play(session_id):
    conn = get_db()
    qs = conn.execute('SELECT * FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
    if not qs:
        conn.close()
        flash('Yarışma bulunamadı.', 'danger')
        return redirect(url_for('quiz_join'))
    participant = conn.execute('SELECT * FROM quiz_participants WHERE session_id = ? AND user_id = ?',
                               (session_id, session['user_id'])).fetchone()
    if not participant:
        conn.close()
        flash('Bu yarışmaya katılmadın.', 'danger')
        return redirect(url_for('quiz_join', code=qs['code']))
    quiz = conn.execute('SELECT * FROM quizzes WHERE id = ?', (qs['quiz_id'],)).fetchone()
    conn.close()
    return render_template('quiz_play.html', quiz_session=qs, quiz=quiz, participant=participant)


# ---------- Canlı oyun motoru ----------

def _quiz_question_public(q):
    """Katılımcıya gönderilecek soru verisi (doğru cevap gizli)."""
    return _quiz_q_dict(q, with_answer=False)


def _quiz_players(conn, session_id):
    return conn.execute('''SELECT id, user_id, nickname, score, streak, best_streak, correct_count
                           FROM quiz_participants WHERE session_id = ? ORDER BY id''', (session_id,)).fetchall()


def _rank_players(players, score_of):
    ordered = sorted(players, key=lambda p: (-score_of(p), p['nickname'].lower(), p['id']))
    return {p['id']: i + 1 for i, p in enumerate(ordered)}, ordered


def _quiz_leaderboard(session_id):
    conn = get_db()
    rows = conn.execute('SELECT nickname, score FROM quiz_participants WHERE session_id = ? ORDER BY score DESC LIMIT 20',
                        (session_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def _quiz_emit(event, payload, session_id):
    socketio.emit(event, payload, room=quiz_room_name(session_id))


def _quiz_live(session_id):
    with QUIZ_LIVE_LOCK:
        return QUIZ_LIVE.setdefault(session_id, {'token': 0, 'phase': 'waiting', 'lock': threading.Lock()})


def _quiz_participant_payload(conn, session_id):
    rows = conn.execute('SELECT id, nickname, score FROM quiz_participants WHERE session_id = ? ORDER BY id',
                        (session_id,)).fetchall()
    return [dict(r) for r in rows]


def _quiz_begin_question(session_id, index):
    conn = get_db()
    qs = conn.execute('SELECT * FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
    questions = _quiz_questions(conn, qs['quiz_id']) if qs else []
    if not qs or index >= len(questions):
        conn.close()
        return _quiz_finish(session_id)
    q = questions[index]
    conn.execute("UPDATE quiz_sessions SET status = 'question', current_index = ?, question_started_at = datetime('now') WHERE id = ?",
                 (index, session_id))
    conn.commit()
    player_count = conn.execute('SELECT COUNT(*) c FROM quiz_participants WHERE session_id = ?', (session_id,)).fetchone()['c']
    conn.close()

    live = _quiz_live(session_id)
    with live['lock']:
        live['token'] += 1
        token = live['token']
        live.update({'phase': 'question', 'index': index, 'total': len(questions), 'question_id': q['id'],
                     'open_at': time.time() + QUIZ_INTRO_SECONDS, 'time_limit': q['time_limit'] or 20,
                     'answered': set(), 'reveal': None, 'finished': None})
        live['close_at'] = live['open_at'] + live['time_limit']
        payload = {'index': index, 'total': len(questions), 'question': _quiz_question_public(q),
                   'intro_ms': QUIZ_INTRO_SECONDS * 1000, 'time_limit': live['time_limit'],
                   'player_count': player_count, 'server_now': time.time()}
        live['question_payload'] = payload
    _quiz_emit('quiz_question', payload, session_id)

    def timer():
        socketio.sleep(QUIZ_INTRO_SECONDS + live['time_limit'] + 0.4)
        _quiz_reveal(session_id, token)
    socketio.start_background_task(timer)


def _quiz_reveal(session_id, token=None):
    live = _quiz_live(session_id)
    with live['lock']:
        if live.get('phase') != 'question' or (token is not None and token != live.get('token')):
            return
        live['phase'] = 'reveal'
        live['token'] += 1
        token = live['token']
        question_id, index, total = live['question_id'], live['index'], live['total']

    conn = get_db()
    q = conn.execute('SELECT * FROM quiz_questions WHERE id = ?', (question_id,)).fetchone()
    players = _quiz_players(conn, session_id)
    answers = {a['participant_id']: a for a in conn.execute(
        'SELECT participant_id, selected_option, is_correct, points, answer_ms FROM quiz_answers WHERE session_id = ? AND question_id = ?',
        (session_id, question_id)).fetchall()}
    # Cevap vermeyenlerin serisi bozulur
    for p in players:
        if p['id'] not in answers and p['streak']:
            conn.execute('UPDATE quiz_participants SET streak = 0 WHERE id = ?', (p['id'],))
    conn.execute("UPDATE quiz_sessions SET status = 'reveal' WHERE id = ?", (session_id,))
    conn.commit()
    players = _quiz_players(conn, session_id)
    conn.close()

    gained = {pid: (a['points'] or 0) for pid, a in answers.items()}
    new_rank, ordered = _rank_players(players, lambda p: p['score'])
    old_rank, _ = _rank_players(players, lambda p: p['score'] - gained.get(p['id'], 0))

    distribution = {k: 0 for k in QUIZ_OPTION_KEYS}
    for a in answers.values():
        if a['selected_option'] in distribution:
            distribution[a['selected_option']] += 1

    results = {}
    for i, p in enumerate(ordered):
        a = answers.get(p['id'])
        ahead = ordered[i - 1] if i > 0 else None
        results[str(p['id'])] = {
            'answered': bool(a), 'correct': bool(a and a['is_correct']),
            'selected': a['selected_option'] if a else None,
            'points': gained.get(p['id'], 0), 'score': p['score'], 'streak': p['streak'] or 0,
            'rank': new_rank[p['id']], 'prev_rank': old_rank[p['id']],
            'ahead_name': ahead['nickname'] if ahead else None,
            'ahead_gap': (ahead['score'] - p['score']) if ahead else 0,
        }
    leaderboard = [{'id': p['id'], 'nickname': p['nickname'], 'score': p['score'],
                    'prev_score': p['score'] - gained.get(p['id'], 0), 'gained': gained.get(p['id'], 0),
                    'rank': new_rank[p['id']], 'prev_rank': old_rank[p['id']], 'streak': p['streak'] or 0}
                   for p in ordered[:10]]
    answer_times = [a['answer_ms'] for a in answers.values() if a['is_correct'] and a['answer_ms'] is not None]
    payload = {
        'index': index, 'total': total, 'is_last': index + 1 >= total,
        'question': _quiz_q_dict(q), 'correct_option': q['correct_option'],
        'distribution': distribution, 'answer_count': len(answers), 'player_count': len(players),
        'correct_count': sum(1 for a in answers.values() if a['is_correct']),
        'fastest_ms': min(answer_times) if answer_times else None,
        'leaderboard': leaderboard, 'results': results, 'reveal_ms': QUIZ_REVEAL_SECONDS * 1000,
    }
    with live['lock']:
        live['reveal'] = payload
    _quiz_emit('quiz_reveal', payload, session_id)

    def to_leaderboard():
        socketio.sleep(QUIZ_REVEAL_SECONDS)
        _quiz_show_leaderboard(session_id, token)
    socketio.start_background_task(to_leaderboard)


def _quiz_show_leaderboard(session_id, token=None):
    live = _quiz_live(session_id)
    with live['lock']:
        if live.get('phase') != 'reveal' or (token is not None and token != live.get('token')) or not live.get('reveal'):
            return
        live['phase'] = 'leaderboard'
        live['token'] += 1
        token = live['token']
        payload = dict(live['reveal'])
        conn = get_db()
        auto = conn.execute('SELECT auto_advance FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
        conn.execute("UPDATE quiz_sessions SET status = 'leaderboard' WHERE id = ?", (session_id,))
        conn.commit()
        conn.close()
        payload['auto_next_ms'] = QUIZ_AUTO_NEXT_SECONDS * 1000 if (auto and auto['auto_advance'] and not payload['is_last']) else 0
        live['leaderboard_payload'] = payload
    _quiz_emit('quiz_leaderboard', payload, session_id)
    if payload['auto_next_ms']:
        def auto_next():
            socketio.sleep(QUIZ_AUTO_NEXT_SECONDS)
            with live['lock']:
                ok = live.get('phase') == 'leaderboard' and live.get('token') == token
            if ok:
                _quiz_begin_question(session_id, payload['index'] + 1)
        socketio.start_background_task(auto_next)


def _quiz_finish(session_id):
    live = _quiz_live(session_id)
    conn = get_db()
    qs = conn.execute('SELECT * FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
    already = qs and qs['status'] == 'finished'
    conn.execute("UPDATE quiz_sessions SET status = 'finished', finished_at = COALESCE(finished_at, CURRENT_TIMESTAMP) WHERE id = ?",
                 (session_id,))
    conn.commit()
    players = _quiz_players(conn, session_id)
    total = conn.execute('SELECT COUNT(*) c FROM quiz_questions WHERE quiz_id = ?', (qs['quiz_id'],)).fetchone()['c'] if qs else 0
    title_row = conn.execute('SELECT title FROM quizzes WHERE id = ?', (qs['quiz_id'],)).fetchone() if qs else None
    conn.close()
    rank, ordered = _rank_players(players, lambda p: p['score'])
    standings = [{'id': p['id'], 'nickname': p['nickname'], 'score': p['score'], 'rank': rank[p['id']],
                  'correct': p['correct_count'] or 0, 'best_streak': p['best_streak'] or 0} for p in ordered]
    payload = {'standings': standings[:50], 'total_questions': total, 'player_count': len(players),
               'results': {str(s['id']): s for s in standings},
               'leaderboard': [{'nickname': s['nickname'], 'score': s['score']} for s in standings[:20]]}
    with live['lock']:
        live['token'] += 1
        live['phase'] = 'finished'
        live['finished'] = payload
    _quiz_emit('quiz_finished', payload, session_id)

    if not already and title_row:
        # Katılım ve derece XP'si (yalnızca bir kez)
        for s in standings:
            p = next(pp for pp in ordered if pp['id'] == s['id'])
            bonus = {1: 50, 2: 30, 3: 20}.get(s['rank'], 0)
            award_xp(p['user_id'], 10 + bonus, f'"{title_row["title"]}" yarışması' + (f' — {s["rank"]}. oldun' if bonus else ''))


def _quiz_snapshot_for(session_id, qs):
    """Sayfayı yenileyen / sonradan katılan oyuncuya mevcut durumu gönderir."""
    live = QUIZ_LIVE.get(session_id)
    status = qs['status']
    if status == 'finished':
        return 'quiz_finished', (live or {}).get('finished') or None
    if not live or live.get('phase') == 'waiting':
        if status in ('question', 'reveal', 'leaderboard'):
            # Sunucu yeniden başlatılmış olabilir: host "Sonraki" ile devam eder.
            return 'quiz_state', {'phase': 'paused', 'index': qs['current_index']}
        return None, None
    phase = live['phase']
    if phase == 'question':
        p = dict(live['question_payload'])
        now = time.time()
        p['intro_ms'] = max(0, int((live['open_at'] - now) * 1000))
        p['remaining_ms'] = max(0, int((live['close_at'] - max(now, live['open_at'])) * 1000))
        return 'quiz_question', p
    if phase == 'reveal':
        return 'quiz_reveal', live.get('reveal')
    if phase == 'leaderboard':
        p = dict(live.get('leaderboard_payload') or {})
        p['auto_next_ms'] = 0
        return 'quiz_leaderboard', p
    return None, None


@socketio.on('quiz_join_room')
def handle_quiz_join_room(data):
    if 'user_id' not in session:
        return
    try:
        session_id = int((data or {}).get('session_id'))
    except (TypeError, ValueError):
        return
    conn = get_db()
    qs = conn.execute('SELECT * FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
    if not qs:
        conn.close()
        return
    is_member = conn.execute('SELECT 1 FROM quiz_participants WHERE session_id = ? AND user_id = ?',
                             (session_id, session['user_id'])).fetchone()
    if not is_member and not _is_quiz_host(qs):
        conn.close()
        return
    join_room(quiz_room_name(session_id))
    participants = _quiz_participant_payload(conn, session_id)
    conn.close()
    _quiz_emit('quiz_participants_update', {'participants': participants}, session_id)
    event, payload = _quiz_snapshot_for(session_id, qs)
    if event and payload is not None:
        payload = dict(payload)
        payload['resync'] = True
        emit(event, payload)
    if _is_quiz_host(qs):
        emit('quiz_host_settings', {'auto_advance': bool(qs['auto_advance'])})


def _host_session_or_none(data):
    if 'user_id' not in session:
        return None
    try:
        session_id = int((data or {}).get('session_id'))
    except (TypeError, ValueError):
        return None
    conn = get_db()
    qs = conn.execute('SELECT * FROM quiz_sessions WHERE id = ?', (session_id,)).fetchone()
    conn.close()
    if not qs or not _is_quiz_host(qs):
        return None
    return qs


@socketio.on('quiz_start')
def handle_quiz_start(data):
    qs = _host_session_or_none(data)
    if not qs or qs['status'] != 'waiting':
        return
    _quiz_begin_question(qs['id'], 0)


@socketio.on('quiz_next')
def handle_quiz_next(data):
    qs = _host_session_or_none(data)
    if not qs or qs['status'] == 'finished':
        return
    live = _quiz_live(qs['id'])
    phase = live.get('phase')
    if phase == 'question':
        return _quiz_reveal(qs['id'])
    if phase == 'reveal':
        return _quiz_show_leaderboard(qs['id'])
    index = live.get('index', qs['current_index'] if qs['current_index'] is not None else -1)
    _quiz_begin_question(qs['id'], (index if index is not None else -1) + 1)


@socketio.on('quiz_skip')
def handle_quiz_skip(data):
    qs = _host_session_or_none(data)
    if qs:
        _quiz_reveal(qs['id'])


@socketio.on('quiz_show_leaderboard')
def handle_quiz_show_leaderboard(data):
    qs = _host_session_or_none(data)
    if qs:
        _quiz_show_leaderboard(qs['id'])


@socketio.on('quiz_end')
def handle_quiz_end(data):
    qs = _host_session_or_none(data)
    if qs and qs['status'] != 'finished':
        _quiz_finish(qs['id'])


@socketio.on('quiz_set_auto')
def handle_quiz_set_auto(data):
    qs = _host_session_or_none(data)
    if not qs:
        return
    conn = get_db()
    conn.execute('UPDATE quiz_sessions SET auto_advance = ? WHERE id = ?', (1 if data.get('value') else 0, qs['id']))
    conn.commit()
    conn.close()
    emit('quiz_host_settings', {'auto_advance': bool(data.get('value'))})


@socketio.on('quiz_kick')
def handle_quiz_kick(data):
    qs = _host_session_or_none(data)
    if not qs or qs['status'] != 'waiting':
        return
    try:
        pid = int(data.get('participant_id'))
    except (TypeError, ValueError):
        return
    conn = get_db()
    conn.execute('DELETE FROM quiz_participants WHERE id = ? AND session_id = ?', (pid, qs['id']))
    conn.commit()
    participants = _quiz_participant_payload(conn, qs['id'])
    conn.close()
    _quiz_emit('quiz_kicked', {'participant_id': pid}, qs['id'])
    _quiz_emit('quiz_participants_update', {'participants': participants}, qs['id'])


@socketio.on('quiz_answer')
def handle_quiz_answer(data):
    if 'user_id' not in session:
        return
    try:
        session_id = int(data.get('session_id'))
        question_id = int(data.get('question_id'))
    except (TypeError, ValueError):
        return
    selected = data.get('selected_option')
    if selected not in QUIZ_OPTION_KEYS:
        return
    live = QUIZ_LIVE.get(session_id)
    now = time.time()
    if not live:
        return
    with live['lock']:
        if live.get('phase') != 'question' or live.get('question_id') != question_id:
            emit('quiz_answer_ack', {'ok': False, 'reason': 'closed'})
            return
        if now < live['open_at'] - 0.3 or now > live['close_at'] + 0.6:
            emit('quiz_answer_ack', {'ok': False, 'reason': 'closed'})
            return
        answer_ms = int(max(0, min(now, live['close_at']) - live['open_at']) * 1000)
        time_limit_ms = live['time_limit'] * 1000

    conn = get_db()
    participant = conn.execute('SELECT * FROM quiz_participants WHERE session_id = ? AND user_id = ?',
                               (session_id, session['user_id'])).fetchone()
    question = conn.execute('SELECT * FROM quiz_questions WHERE id = ?', (question_id,)).fetchone()
    if not participant or not question:
        conn.close()
        return
    already = conn.execute('SELECT id FROM quiz_answers WHERE question_id = ? AND participant_id = ?',
                           (question_id, participant['id'])).fetchone()
    if already:
        conn.close()
        emit('quiz_answer_ack', {'ok': True, 'selected': selected, 'duplicate': True})
        return

    is_correct = 1 if selected == question['correct_option'] else 0
    multiplier = QUIZ_POINTS_MODES.get(question['points_mode'] or 'standard', ('', 1))[1]
    points = 0
    streak = (participant['streak'] or 0) + 1 if is_correct else 0
    if is_correct and multiplier:
        # Kahoot formülü: hızlı cevap = 1000'e yakın, son saniye = ~500 puan
        ratio = min(1.0, answer_ms / time_limit_ms) if time_limit_ms else 1.0
        points = int(round((1 - ratio / 2) * 1000)) * multiplier
        if streak >= 2:
            points += min(streak - 1, 5) * 100  # seri bonusu
    try:
        conn.execute('''INSERT INTO quiz_answers (session_id, question_id, participant_id, selected_option, is_correct, answer_ms, points)
                        VALUES (?, ?, ?, ?, ?, ?, ?)''',
                     (session_id, question_id, participant['id'], selected, is_correct, answer_ms, points))
    except sqlite3.IntegrityError:
        conn.close()
        return
    conn.execute('''UPDATE quiz_participants SET score = score + ?, streak = ?, correct_count = correct_count + ?,
                    best_streak = MAX(COALESCE(best_streak, 0), ?) WHERE id = ?''',
                 (points, streak, is_correct, streak, participant['id']))
    conn.commit()
    answer_count = conn.execute('SELECT COUNT(*) as c FROM quiz_answers WHERE session_id = ? AND question_id = ?',
                                (session_id, question_id)).fetchone()['c']
    player_count = conn.execute('SELECT COUNT(*) c FROM quiz_participants WHERE session_id = ?', (session_id,)).fetchone()['c']
    conn.close()

    emit('quiz_answer_ack', {'ok': True, 'selected': selected, 'answer_ms': answer_ms})
    _quiz_emit('quiz_answer_count', {'answer_count': answer_count, 'player_count': player_count, 'question_id': question_id},
               session_id)
    if answer_count >= player_count:
        # Herkes cevapladı -> kısa bir nefes payından sonra sonuçları göster
        token = live['token']

        def early_reveal():
            socketio.sleep(0.8)
            _quiz_reveal(session_id, token)
        socketio.start_background_task(early_reveal)


# ==================== OYUNLAR (Mini oyunlar) ====================

def get_game_leaderboard(game_key, limit=10):
    meta = GAMES[game_key]
    agg = 'MIN' if meta['sort'] == 'asc' else 'MAX'
    order = 'ASC' if meta['sort'] == 'asc' else 'DESC'
    conn = get_db()
    rows = conn.execute(f'''
        SELECT u.id as user_id, u.username, u.full_name, u.profile_photo,
               {agg}(gs.score) as best_score
        FROM game_scores gs JOIN users u ON gs.user_id = u.id
        WHERE gs.game_key = ?
        GROUP BY gs.user_id
        ORDER BY best_score {order}
        LIMIT ?
    ''', (game_key, limit)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_user_best_score(game_key, user_id):
    meta = GAMES[game_key]
    agg = 'MIN' if meta['sort'] == 'asc' else 'MAX'
    conn = get_db()
    row = conn.execute(f'SELECT {agg}(score) as best FROM game_scores WHERE game_key = ? AND user_id = ?',
                        (game_key, user_id)).fetchone()
    conn.close()
    return row['best'] if row and row['best'] is not None else None


@app.route('/games')
@login_required
def games_hub():
    games_data = []
    for key, meta in GAMES.items():
        top = get_game_leaderboard(key, limit=1)
        games_data.append({
            'key': key, **meta,
            'top_player': top[0] if top else None,
            'my_best': get_user_best_score(key, session['user_id']),
        })
    return render_template('games_hub.html', games=games_data)


@app.route('/games/<game_key>')
@login_required
def game_play(game_key):
    if game_key not in GAMES:
        flash('Böyle bir oyun yok.', 'danger')
        return redirect(url_for('games_hub'))
    meta = GAMES[game_key]
    leaderboard = get_game_leaderboard(game_key, limit=10)
    my_best = get_user_best_score(game_key, session['user_id'])
    template_map = {
        'snake': 'game_snake.html',
        'reflex': 'game_reflex.html',
        'whackamole': 'game_whackamole.html',
        'penalty': 'game_penalty.html',
        'memory': 'game_memory.html',
        '2048': 'game_2048.html',
    }
    return render_template(template_map[game_key], game_key=game_key, meta=meta,
                            leaderboard=leaderboard, my_best=my_best)


@app.route('/games/<game_key>/score', methods=['POST'])
@login_required
def submit_game_score(game_key):
    if game_key not in GAMES:
        return jsonify({'error': 'invalid_game'}), 404

    score = request.json.get('score') if request.is_json else request.form.get('score', type=int)
    try:
        score = int(score)
    except (TypeError, ValueError):
        return jsonify({'error': 'invalid_score'}), 400

    # Basit mantık kontrolü — sıra dışı büyük değerleri reddet (hile önleme amaçlı temel kontrol).
    if score < 0 or score > 1_000_000:
        return jsonify({'error': 'invalid_score'}), 400

    conn = get_db()
    conn.execute('INSERT INTO game_scores (game_key, user_id, score) VALUES (?, ?, ?)',
                 (game_key, session['user_id'], score))
    conn.commit()
    conn.close()
    award_xp(session['user_id'], 10, f'{GAMES[game_key]["name"]} oynandı')

    leaderboard = get_game_leaderboard(game_key, limit=10)
    my_best = get_user_best_score(game_key, session['user_id'])
    return jsonify({'ok': True, 'leaderboard': leaderboard, 'my_best': my_best})



# ==================== ODA ARKADAŞI EŞLEŞTİRME ====================

ROOMMATE_TRAITS = [
    {
        'key': 'sleep_time', 'label': 'Uyku saati', 'icon': 'bi-moon-stars', 'weight': 1.25,
        'options': [('early', 'Erken uyurum (22:30–23:30)'), ('normal', 'Normal (23:30–00:30)'), ('late', 'Geç uyurum (00:30+)')]
    },
    {
        'key': 'wake_time', 'label': 'Uyanma saati', 'icon': 'bi-sunrise', 'weight': 1.0,
        'options': [('early', 'Erken kalkarım (06:00–07:00)'), ('normal', 'Normal (07:00–08:30)'), ('late', 'Geç kalkarım (08:30+)')]
    },
    {
        'key': 'noise', 'label': 'Oda içi ses seviyesi', 'icon': 'bi-volume-up', 'weight': 1.35,
        'options': [('quiet', 'Sessizim'), ('balanced', 'Normal'), ('loud', 'Hareketli / ses yapabilirim')]
    },
    {
        'key': 'cleanliness', 'label': 'Temizlik ve düzen', 'icon': 'bi-stars', 'weight': 1.35,
        'options': [('very_tidy', 'Çok düzenliyim'), ('tidy', 'Düzenli olmaya dikkat ederim'), ('relaxed', 'Dağınıklığı çok önemsemem')]
    },
    {
        'key': 'lights', 'label': 'Gece ışık kullanımı', 'icon': 'bi-lightbulb', 'weight': 1.05,
        'options': [('dark', 'Uyurken tamamen karanlık'), ('dim', 'Loş ışık sorun değil'), ('on', 'Gece ışık kullanabilirim')]
    },
    {
        'key': 'study', 'label': 'Odada ders çalışma', 'icon': 'bi-book', 'weight': 1.0,
        'options': [('often', 'Sık sık odada çalışırım'), ('sometimes', 'Bazen odada çalışırım'), ('rarely', 'Nadiren odada çalışırım')]
    },
    {
        'key': 'social', 'label': 'Sosyallik', 'icon': 'bi-people', 'weight': 0.85,
        'options': [('quiet', 'Daha sakinim'), ('balanced', 'Dengeli'), ('social', 'Sosyal ve konuşkanım')]
    },
    {
        'key': 'guests', 'label': 'Odaya misafir', 'icon': 'bi-person-plus', 'weight': 1.15,
        'options': [('never', 'Neredeyse hiç'), ('rare', 'Nadiren'), ('sometimes', 'Bazen'), ('often', 'Sık sık')]
    },
    {
        'key': 'audio', 'label': 'Telefon / müzik kullanımı', 'icon': 'bi-headphones', 'weight': 1.0,
        'options': [('headphones', 'Kulaklık kullanırım'), ('low', 'Düşük ses'), ('speaker', 'Hoparlör kullanabilirim')]
    },
    {
        'key': 'temperature', 'label': 'Oda sıcaklığı', 'icon': 'bi-thermometer-half', 'weight': 0.7,
        'options': [('cool', 'Serin severim'), ('neutral', 'Orta sıcaklık'), ('warm', 'Sıcak severim')]
    },
    {
        'key': 'sharing', 'label': 'Ortak eşya kullanımı', 'icon': 'bi-box-seam', 'weight': 0.65,
        'options': [('private', 'Kişisel eşyalar ayrı olsun'), ('ask', 'Sorarak paylaşırım'), ('share', 'Paylaşmakta sorun yok')]
    },
    {
        'key': 'privacy', 'label': 'Kişisel alan', 'icon': 'bi-shield-lock', 'weight': 0.95,
        'options': [('high', 'Kişisel alanım önemli'), ('balanced', 'Dengeli'), ('social', 'Oda içinde sürekli iletişim sorun değil')]
    },
    {
        'key': 'alarm', 'label': 'Alarm kullanımı', 'icon': 'bi-alarm', 'weight': 0.9,
        'options': [('one', 'Tek alarm'), ('few', 'Birkaç alarm'), ('many', 'Çoklu alarm / erteleme')]
    },
    {
        'key': 'weekend', 'label': 'Hafta sonu düzeni', 'icon': 'bi-calendar-week', 'weight': 0.7,
        'options': [('early', 'Hafta sonu da erken kalkarım'), ('normal', 'Normal'), ('late', 'Hafta sonu geç kalkarım')]
    },
    {
        'key': 'organization', 'label': 'Eşyaları kullanma düzeni', 'icon': 'bi-grid-3x3-gap', 'weight': 0.9,
        'options': [('organized', 'Her şeyin yeri belli olsun'), ('balanced', 'Dengeli'), ('flexible', 'Çok takılmam')]
    },
]
ROOMMATE_TRAIT_MAP = {t['key']: t for t in ROOMMATE_TRAITS}
ROOMMATE_IMPORTANCE = [(1, 'Düşük'), (2, 'Az'), (3, 'Orta'), (4, 'Yüksek'), (5, 'Çok önemli')]


def _roommate_default_profile():
    return {t['key']: {'value': t['options'][1][0], 'importance': 3} for t in ROOMMATE_TRAITS}


def _roommate_load_profile(row):
    defaults = _roommate_default_profile()
    if row:
        try:
            raw = json.loads(row['preferences_json'] or '{}')
            if isinstance(raw, dict):
                for key, item in raw.items():
                    if key in ROOMMATE_TRAIT_MAP and isinstance(item, dict):
                        allowed = {v for v, _ in ROOMMATE_TRAIT_MAP[key]['options']}
                        value = item.get('value')
                        importance = int(item.get('importance', 3)) if str(item.get('importance', '')).isdigit() else 3
                        if value in allowed:
                            defaults[key]['value'] = value
                        defaults[key]['importance'] = max(1, min(5, importance))
        except (ValueError, TypeError, json.JSONDecodeError):
            pass
    return defaults


def _roommate_option_index(trait, value):
    values = [v for v, _ in trait['options']]
    try:
        return values.index(value)
    except ValueError:
        return 0


def _roommate_pair_score(my_profile, other_profile):
    weighted = 0.0
    total = 0.0
    mismatches = []
    for trait in ROOMMATE_TRAITS:
        key = trait['key']
        mine = my_profile[key]
        theirs = other_profile[key]
        max_idx = max(1, len(trait['options']) - 1)
        distance = abs(_roommate_option_index(trait, mine['value']) - _roommate_option_index(trait, theirs['value']))
        similarity = 1.0 - (distance / max_idx)
        importance = (mine['importance'] + theirs['importance']) / 2
        weight = trait['weight'] * importance
        weighted += similarity * weight
        total += weight
        if distance == max_idx and mine['importance'] >= 4 and theirs['importance'] >= 4:
            mismatches.append(trait['label'])
    score = round((weighted / total) * 100) if total else 0
    # Kritik uyuşmazlıklar skoru aşırı iyimser göstermesin.
    if len(mismatches) >= 2:
        score = min(score, 64)
    elif len(mismatches) == 1:
        score = min(score, 78)
    return max(0, min(100, score)), mismatches


def _roommate_match_reasons(my_profile, other_profile):
    similarities = []
    for trait in ROOMMATE_TRAITS:
        key = trait['key']
        mine = my_profile[key]
        theirs = other_profile[key]
        max_idx = max(1, len(trait['options']) - 1)
        distance = abs(_roommate_option_index(trait, mine['value']) - _roommate_option_index(trait, theirs['value']))
        similarity = 1.0 - (distance / max_idx)
        importance = (mine['importance'] + theirs['importance']) / 2
        if similarity >= 0.999 and importance >= 3:
            similarities.append(trait['label'])
        elif similarity >= 0.66 and importance >= 4:
            similarities.append(trait['label'])
    return similarities[:4]


@app.route('/roommates', methods=['GET', 'POST'])
@login_required
def roommate_matching():
    conn = get_db()
    current = conn.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    profile_row = conn.execute('SELECT * FROM roommate_profiles WHERE user_id = ?', (session['user_id'],)).fetchone()

    if request.method == 'POST':
        submitted = {}
        for trait in ROOMMATE_TRAITS:
            key = trait['key']
            value = request.form.get(f'{key}_value', '').strip()
            allowed = {v for v, _ in trait['options']}
            if value not in allowed:
                flash(f'{trait["label"]} için geçerli bir seçim yapmalısın.', 'danger')
                conn.close()
                return redirect(url_for('roommate_matching'))
            try:
                importance = int(request.form.get(f'{key}_importance', '3'))
            except ValueError:
                importance = 3
            importance = max(1, min(5, importance))
            submitted[key] = {'value': value, 'importance': importance}

        conn.execute('''
            INSERT INTO roommate_profiles(user_id, preferences_json, completed, updated_at)
            VALUES (?, ?, 1, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id) DO UPDATE SET preferences_json=excluded.preferences_json,
                completed=1, updated_at=CURRENT_TIMESTAMP
        ''', (session['user_id'], json.dumps(submitted, ensure_ascii=False)))
        conn.commit()
        profile_row = conn.execute('SELECT * FROM roommate_profiles WHERE user_id = ?', (session['user_id'],)).fetchone()
        flash('Oda arkadaşı tercihlerin kaydedildi. Sana en uyumlu kişiler yeniden hesaplandı.', 'success')

    my_profile = _roommate_load_profile(profile_row)
    completed = bool(profile_row and profile_row['completed'])
    matches = []
    if completed:
        rows = conn.execute('''
            SELECT u.id, u.username, u.full_name, u.profile_photo, u.room_number, rp.preferences_json, rp.completed
            FROM roommate_profiles rp
            JOIN users u ON u.id = rp.user_id
            WHERE rp.completed = 1 AND rp.user_id != ? AND COALESCE(u.is_admin, 0) = 0
        ''', (session['user_id'],)).fetchall()
        for row in rows:
            other = _roommate_load_profile(row)
            score, mismatches = _roommate_pair_score(my_profile, other)
            reasons = _roommate_match_reasons(my_profile, other)
            matches.append({
                'id': row['id'], 'username': row['username'], 'full_name': row['full_name'],
                'profile_photo': row['profile_photo'], 'room_number': row['room_number'],
                'score': score, 'reasons': reasons, 'mismatches': mismatches
            })
        matches.sort(key=lambda item: (-item['score'], item['full_name'].lower()))
        matches = matches[:12]
    conn.close()

    return render_template('roommates.html', traits=ROOMMATE_TRAITS, importance_options=ROOMMATE_IMPORTANCE,
                           importance_labels=dict(ROOMMATE_IMPORTANCE),
                           roommate_profile=my_profile, completed=completed, matches=matches)


# ==================== PROFİL ====================

@app.route('/profile/<username>')
@login_required
def profile(username):
    conn = get_db()

    user = conn.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
    if not user:
        conn.close()
        flash('Kullanıcı bulunamadı.', 'danger')
        return redirect(url_for('dashboard'))

    posts = conn.execute('''
        SELECT p.*,
               (SELECT COUNT(*) FROM post_likes WHERE post_id = p.id) as like_count,
               (SELECT COUNT(*) FROM post_comments WHERE post_id = p.id) as comment_count
        FROM posts p WHERE p.user_id = ?
        ORDER BY p.created_at DESC
    ''', (user['id'],)).fetchall()

    post_count = len(posts)

    followers_count = conn.execute('SELECT COUNT(*) as c FROM follows WHERE followed_id = ?',
                                    (user['id'],)).fetchone()['c']
    following_count = conn.execute('SELECT COUNT(*) as c FROM follows WHERE follower_id = ?',
                                    (user['id'],)).fetchone()['c']
    am_following = False
    if user['id'] != session['user_id']:
        am_following = conn.execute(
            'SELECT id FROM follows WHERE follower_id = ? AND followed_id = ?',
            (session['user_id'], user['id'])
        ).fetchone() is not None
    conn.close()

    return render_template('profile.html', profile_user=user, posts=posts, post_count=post_count,
                            followers_count=followers_count, following_count=following_count, am_following=am_following)


@app.route('/profile/<username>/follow', methods=['POST'])
@login_required
def toggle_follow(username):
    conn = get_db()
    target = conn.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
    if not target:
        conn.close()
        flash('Kullanıcı bulunamadı.', 'danger')
        return redirect(url_for('dashboard'))

    if target['id'] == session['user_id']:
        conn.close()
        flash('Kendini takip edemezsin.', 'warning')
        return redirect(url_for('profile', username=username))

    existing = conn.execute('SELECT id FROM follows WHERE follower_id = ? AND followed_id = ?',
                             (session['user_id'], target['id'])).fetchone()
    if existing:
        conn.execute('DELETE FROM follows WHERE id = ?', (existing['id'],))
        conn.commit()
        following = False
    else:
        conn.execute('INSERT INTO follows (follower_id, followed_id) VALUES (?, ?)',
                      (session['user_id'], target['id']))
        conn.commit()
        following = True
        me = conn.execute('SELECT full_name FROM users WHERE id = ?', (session['user_id'],)).fetchone()
        create_notification(
            target['id'], 'Yeni takipçi',
            f"{me['full_name']} seni takip etmeye başladı.",
            link=url_for('profile', username=session['username'])
        )

    followers_count = conn.execute('SELECT COUNT(*) as c FROM follows WHERE followed_id = ?',
                                    (target['id'],)).fetchone()['c']
    conn.close()

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({'following': following, 'followers_count': followers_count})

    return redirect(url_for('profile', username=username))


@app.route('/profile/<username>/followers')
@login_required
def followers_list(username):
    conn = get_db()
    user = conn.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
    if not user:
        conn.close()
        flash('Kullanıcı bulunamadı.', 'danger')
        return redirect(url_for('dashboard'))
    people = conn.execute('''
        SELECT u.id, u.username, u.full_name, u.profile_photo,
               EXISTS(SELECT 1 FROM follows f2 WHERE f2.follower_id = ? AND f2.followed_id = u.id) as am_following
        FROM follows f JOIN users u ON f.follower_id = u.id
        WHERE f.followed_id = ?
        ORDER BY f.created_at DESC
    ''', (session['user_id'], user['id'])).fetchall()
    conn.close()
    return render_template('follow_list.html', profile_user=user, people=people,
                            list_title='Takipçiler')


@app.route('/profile/<username>/following')
@login_required
def following_list(username):
    conn = get_db()
    user = conn.execute('SELECT * FROM users WHERE username = ?', (username,)).fetchone()
    if not user:
        conn.close()
        flash('Kullanıcı bulunamadı.', 'danger')
        return redirect(url_for('dashboard'))
    people = conn.execute('''
        SELECT u.id, u.username, u.full_name, u.profile_photo,
               EXISTS(SELECT 1 FROM follows f2 WHERE f2.follower_id = ? AND f2.followed_id = u.id) as am_following
        FROM follows f JOIN users u ON f.followed_id = u.id
        WHERE f.follower_id = ?
        ORDER BY f.created_at DESC
    ''', (session['user_id'], user['id'])).fetchall()
    conn.close()
    return render_template('follow_list.html', profile_user=user, people=people,
                            list_title='Takip Edilenler')


@app.route('/profile/edit', methods=['GET', 'POST'])
@login_required
def edit_profile():
    conn = get_db()
    user = conn.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()

    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        room_number = request.form.get('room_number', '').strip()
        if not full_name:
            flash('Ad Soyad alanı zorunludur.', 'danger')
            conn.close()
            return redirect(url_for('edit_profile', required=1))
        bio = request.form.get('bio', '').strip()
        in_dorm = 1 if request.form.get('in_dorm') else 0
        photo_filename = user['profile_photo']

        file = request.files.get('profile_photo')
        if file and file.filename:
            if allowed_file(file.filename, IMAGE_EXT):
                photo_filename = save_media(file, UPLOAD_PROFILE, IMAGE_EXT)
            else:
                flash('Profil fotoğrafı sadece resim dosyası olabilir (png, jpg, jpeg, gif, webp).', 'danger')
                conn.close()
                return redirect(url_for('edit_profile'))

        new_password = request.form.get('new_password', '')
        password_sql = user['password']
        if new_password:
            if len(new_password) < 4:
                flash('Yeni şifre en az 4 karakter olmalı.', 'danger')
                conn.close()
                return redirect(url_for('edit_profile'))
            password_sql = generate_password_hash(new_password)

        conn.execute('''
            UPDATE users SET full_name = ?, room_number = ?, bio = ?, in_dorm = ?,
                              profile_photo = ?, password = ?
            WHERE id = ?
        ''', (full_name, room_number, bio, in_dorm, photo_filename, password_sql, session['user_id']))
        conn.commit()
        conn.close()

        flash('Profil güncellendi!', 'success')
        return redirect(url_for('profile', username=session['username']))

    conn.close()
    return render_template('edit_profile.html', user=user)


# ==================== SOSYAL AKIŞ (FEED) ====================

@app.route('/feed')
@login_required
def feed():
    page = request.args.get('page', 1, type=int)
    if page < 1:
        page = 1
    offset = (page - 1) * POSTS_PER_PAGE

    conn = get_db()
    # PERFORMANS: bir sayfa fazladan çekip "sonraki sayfa var mı" bilgisini
    # ekstra bir COUNT(*) sorgusu olmadan anlıyoruz.
    posts = conn.execute('''
        SELECT p.*, u.full_name, u.username, u.profile_photo, u.role, u.is_admin,
               (SELECT COUNT(*) FROM post_likes WHERE post_id = p.id) as like_count,
               (SELECT COUNT(*) FROM post_comments WHERE post_id = p.id) as comment_count,
               (SELECT COUNT(*) FROM post_likes WHERE post_id = p.id AND user_id = ?) as liked_by_me
        FROM posts p
        JOIN users u ON p.user_id = u.id
        ORDER BY p.created_at DESC
        LIMIT ? OFFSET ?
    ''', (session['user_id'], POSTS_PER_PAGE + 1, offset)).fetchall()

    has_more = len(posts) > POSTS_PER_PAGE
    posts = posts[:POSTS_PER_PAGE]

    posts_with_comments = []
    for p in posts:
        comments = conn.execute('''
            SELECT pc.*, u.full_name, u.username, u.role, u.is_admin
            FROM post_comments pc
            JOIN users u ON pc.user_id = u.id
            WHERE pc.post_id = ?
            ORDER BY pc.created_at ASC
        ''', (p['id'],)).fetchall()
        row = dict(p)
        row['comments'] = nest_comments(comments)
        posts_with_comments.append(row)

    conn.close()

    # Aktif hikayeler: yurttaki tüm kullanıcıların (kendim dahil) süresi dolmamış hikayeleri,
    # kullanıcı bazında gruplu ve "görüldü mü" bilgisiyle.
    conn = get_db()
    story_rows = conn.execute('''
        SELECT s.*, u.username, u.full_name, u.profile_photo,
               EXISTS(SELECT 1 FROM story_views v WHERE v.story_id = s.id AND v.user_id = ?) as seen
        FROM stories s
        JOIN users u ON s.user_id = u.id
        WHERE s.expires_at > datetime('now')
        ORDER BY s.user_id = ? DESC, s.created_at ASC
    ''', (session['user_id'], session['user_id'])).fetchall()
    conn.close()

    stories_by_user = OrderedDict()
    for s in story_rows:
        uid = s['user_id']
        if uid not in stories_by_user:
            stories_by_user[uid] = {
                'user_id': uid, 'username': s['username'], 'full_name': s['full_name'],
                'profile_photo': s['profile_photo'], 'items': [], 'all_seen': True
            }
        stories_by_user[uid]['items'].append(dict(s))
        if not s['seen']:
            stories_by_user[uid]['all_seen'] = False

    return render_template('feed.html', posts=posts_with_comments, page=page, has_more=has_more,
                            story_groups=list(stories_by_user.values()))


@app.route('/post/create', methods=['GET', 'POST'])
@login_required
def create_post():
    if request.method == 'POST':
        caption = request.form.get('caption', '').strip()
        file = request.files.get('media')
        has_file = file and file.filename

        if not has_file and not caption:
            flash('Bir fotoğraf/video seç ya da en azından bir şeyler yaz.', 'danger')
            return redirect(url_for('create_post'))

        if has_file:
            if allowed_file(file.filename, IMAGE_EXT):
                media_type = 'image'
                filename = save_media(file, UPLOAD_POSTS, IMAGE_EXT)
            elif allowed_file(file.filename, VIDEO_EXT):
                media_type = 'video'
                filename = save_media(file, UPLOAD_POSTS, VIDEO_EXT)
            else:
                flash('Desteklenmeyen dosya türü. Resim ya da video yükleyin.', 'danger')
                return redirect(url_for('create_post'))
        else:
            # Sadece metin paylaşımı — medya dosyası yok.
            media_type = 'text'
            filename = None

        conn = get_db()
        conn.execute('''
            INSERT INTO posts (user_id, caption, media_path, media_type)
            VALUES (?, ?, ?, ?)
        ''', (session['user_id'], caption, filename, media_type))
        conn.commit()
        conn.close()

        flash('Gönderi paylaşıldı!', 'success')
        return redirect(url_for('feed'))

    return render_template('create_post.html')


@app.route('/story/create', methods=['POST'])
@login_required
def create_story():
    file = request.files.get('media')
    caption = request.form.get('caption', '').strip()
    text_overlays = request.form.get('text_overlays', '').strip()

    if not file or not file.filename:
        flash('Hikaye için bir fotoğraf ya da video seç.', 'danger')
        return redirect(url_for('feed'))

    # Metin katmanlarının geçerli JSON olduğundan emin ol; değilse yok say.
    if text_overlays:
        try:
            json.loads(text_overlays)
        except (ValueError, TypeError):
            text_overlays = None
    else:
        text_overlays = None

    if allowed_file(file.filename, IMAGE_EXT):
        media_type = 'image'
        filename = save_media(file, UPLOAD_STORIES, IMAGE_EXT)
    elif allowed_file(file.filename, VIDEO_EXT):
        media_type = 'video'
        filename = save_media(file, UPLOAD_STORIES, VIDEO_EXT)
    else:
        flash('Desteklenmeyen dosya türü. Resim ya da video yükleyin.', 'danger')
        return redirect(url_for('feed'))

    conn = get_db()
    conn.execute('''
        INSERT INTO stories (user_id, media_path, media_type, caption, text_overlays, expires_at)
        VALUES (?, ?, ?, ?, ?, datetime('now', '+24 hours'))
    ''', (session['user_id'], filename, media_type, caption, text_overlays))
    conn.commit()
    conn.close()

    flash('Hikayen paylaşıldı! 24 saat sonra otomatik kalkacak.', 'success')
    return redirect(url_for('feed'))


@app.route('/story/user/<int:user_id>')
@login_required
def get_user_stories(user_id):
    """Bir kullanıcının aktif hikayelerini JSON olarak döner (görüntüleyici için)."""
    conn = get_db()
    user = conn.execute('SELECT id, username, full_name, profile_photo FROM users WHERE id = ?',
                         (user_id,)).fetchone()
    if not user:
        conn.close()
        return jsonify({'error': 'not_found'}), 404

    items = conn.execute('''
        SELECT * FROM stories WHERE user_id = ? AND expires_at > datetime('now')
        ORDER BY created_at ASC
    ''', (user_id,)).fetchall()
    conn.close()

    return jsonify({
        'user': dict(user),
        'items': [dict(i) for i in items]
    })


@app.route('/story/<int:story_id>/view', methods=['POST'])
@login_required
def view_story(story_id):
    conn = get_db()
    story = conn.execute('SELECT * FROM stories WHERE id = ?', (story_id,)).fetchone()
    if story:
        conn.execute('''
            INSERT OR IGNORE INTO story_views (story_id, user_id) VALUES (?, ?)
        ''', (story_id, session['user_id']))
        conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/story/<int:story_id>/delete', methods=['POST'])
@login_required
def delete_story(story_id):
    conn = get_db()
    story = conn.execute('SELECT * FROM stories WHERE id = ?', (story_id,)).fetchone()
    if not story or story['user_id'] != session['user_id']:
        conn.close()
        flash('Bu hikayeyi silemezsin.', 'danger')
        return redirect(url_for('feed'))
    try:
        if story['media_path']:
            os.remove(os.path.join(UPLOAD_STORIES, story['media_path']))
    except OSError:
        pass
    conn.execute('DELETE FROM story_views WHERE story_id = ?', (story_id,))
    conn.execute('DELETE FROM stories WHERE id = ?', (story_id,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/post/<int:post_id>/like', methods=['POST'])
@login_required
def like_post(post_id):
    conn = get_db()
    existing = conn.execute('SELECT id FROM post_likes WHERE post_id = ? AND user_id = ?',
                             (post_id, session['user_id'])).fetchone()
    post = conn.execute('SELECT id, user_id FROM posts WHERE id = ?', (post_id,)).fetchone()
    if not post:
        conn.close()
        return jsonify({'error': 'Gönderi bulunamadı.'}), 404
    if existing:
        conn.execute('DELETE FROM post_likes WHERE id = ?', (existing['id'],))
        liked = False
    else:
        conn.execute('INSERT INTO post_likes (post_id, user_id) VALUES (?, ?)', (post_id, session['user_id']))
        liked = True
    conn.commit()
    like_count = conn.execute('SELECT COUNT(*) as c FROM post_likes WHERE post_id = ?', (post_id,)).fetchone()['c']
    liker = conn.execute('SELECT full_name FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    conn.close()
    if liked and post['user_id'] != session['user_id']:
        create_notification(post['user_id'], 'Gönderin beğenildi ❤️', f"{liker['full_name']} gönderini beğendi.", link=url_for('feed') + f'#post-{post_id}')
    return jsonify({'liked': liked, 'like_count': like_count})


@app.route('/post/<int:post_id>/comment', methods=['POST'], endpoint='comment_post')
@app.route('/post/<int:post_id>/add_comment', methods=['POST'], endpoint='add_comment')
@login_required
def comment_post(post_id):
    text = request.form.get('comment', '').strip()
    parent_id = request.form.get('parent_id', type=int)

    if text:
        conn = get_db()
        post = conn.execute('SELECT * FROM posts WHERE id = ?', (post_id,)).fetchone()
        if not post:
            conn.close()
            flash('Gönderi bulunamadı.', 'danger')
            return redirect(url_for('feed'))

        parent_comment = None
        if parent_id:
            parent_comment = conn.execute('SELECT * FROM post_comments WHERE id = ? AND post_id = ?',
                                           (parent_id, post_id)).fetchone()
            if not parent_comment:
                parent_id = None

        cur = conn.cursor()
        cur.execute('''
            INSERT INTO post_comments (post_id, user_id, comment, parent_id) VALUES (?, ?, ?, ?)
        ''', (post_id, session['user_id'], text, parent_id))
        conn.commit()

        commenter = conn.execute('SELECT full_name FROM users WHERE id = ?',
                                  (session['user_id'],)).fetchone()
        conn.close()

        if parent_comment and parent_comment['user_id'] != session['user_id']:
            # Yanıt bildirimi
            create_notification(
                parent_comment['user_id'],
                'Yorumuna yanıt geldi',
                f"{commenter['full_name']} yorumuna yanıt verdi: \"{text[:80]}\"",
                link=url_for('feed') + f'#post-{post_id}'
            )
        elif not parent_comment and post['user_id'] != session['user_id']:
            # Yeni yorum bildirimi
            create_notification(
                post['user_id'],
                'Gönderine yorum yapıldı',
                f"{commenter['full_name']} gönderine yorum yaptı: \"{text[:80]}\"",
                link=url_for('feed') + f'#post-{post_id}'
            )
    return redirect(url_for('feed') + f'#post-{post_id}')


@app.route('/post/<int:post_id>/delete')
@login_required
def delete_post(post_id):
    conn = get_db()
    post = conn.execute('SELECT * FROM posts WHERE id = ?', (post_id,)).fetchone()
    if not post:
        conn.close()
        flash('Gönderi bulunamadı.', 'danger')
        return redirect(url_for('feed'))

    if post['user_id'] != session['user_id'] and not has_perm('moderate_content'):
        conn.close()
        flash('Bu gönderiyi silme yetkiniz yok.', 'danger')
        return redirect(url_for('feed'))

    conn.execute('DELETE FROM post_likes WHERE post_id = ?', (post_id,))
    conn.execute('DELETE FROM post_comments WHERE post_id = ?', (post_id,))
    conn.execute('DELETE FROM posts WHERE id = ?', (post_id,))
    conn.commit()
    conn.close()

    try:
        if post['media_path']:
            os.remove(os.path.join(UPLOAD_POSTS, post['media_path']))
    except OSError:
        pass

    flash('Gönderi silindi.', 'success')
    return redirect(url_for('feed'))


# ==================== ETKİNLİKLER ====================

@app.route('/events')
@login_required
def events():
    conn = get_db()
    all_events = conn.execute('''
        SELECT e.*, u.full_name as creator_name,
               (SELECT COUNT(*) FROM event_participants WHERE event_id = e.id) as participant_count
        FROM events e
        JOIN users u ON e.created_by = u.id
        ORDER BY e.event_date ASC
    ''').fetchall()

    my_events = conn.execute('''
        SELECT event_id FROM event_participants WHERE user_id = ?
    ''', (session['user_id'],)).fetchall()
    my_event_ids = [e['event_id'] for e in my_events]
    conn.close()

    return render_template('events.html', events=all_events, my_event_ids=my_event_ids)


@app.route('/event/create', methods=['GET', 'POST'])
@login_required
def create_event():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        description = request.form.get('description', '').strip()
        event_type = request.form.get('event_type', 'Diğer').strip()
        location = request.form.get('location', '').strip()
        event_date = request.form.get('event_date', '').strip()
        max_participants = request.form.get('max_participants') or None

        if not title or not location or not event_date:
            flash('Başlık, konum ve tarih/saat zorunludur.', 'danger')
            return render_template('create_event.html')

        if max_participants is not None:
            try:
                max_participants = int(max_participants)
                if max_participants < 1:
                    max_participants = None
            except (ValueError, TypeError):
                max_participants = None

        image_path = None
        image_file = request.files.get('image')
        if image_file and image_file.filename:
            if not allowed_file(image_file.filename, IMAGE_EXT):
                flash('Etkinlik görseli PNG, JPG, JPEG, GIF veya WEBP olmalıdır.', 'danger')
                return render_template('create_event.html')
            try:
                image_path = save_media(image_file, UPLOAD_EVENTS, IMAGE_EXT)
            except Exception as exc:
                print(f'[EVENT IMAGE] görsel kaydedilemedi: {exc}')
                flash('Etkinlik görseli yüklenemedi. Görseli kaldırıp tekrar deneyebilirsin.', 'danger')
                return render_template('create_event.html')

        conn = get_db()
        try:
            # Eski database.db dosyalarında image_path olmayabilir. Her etkinlik
            # oluşturma isteğinde güvenli migration yaparak uygulamanın çökmesini önle.
            _ensure_column(conn, 'events', 'image_path', 'TEXT')
            cur = conn.cursor()
            cur.execute('''
                INSERT INTO events (title, description, event_type, location, event_date, max_participants, image_path, created_by)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            ''', (title, description, event_type, location, event_date, max_participants, image_path, session['user_id']))
            event_id = cur.lastrowid
            conn.commit()
        except Exception as exc:
            conn.rollback()
            print(f'[EVENT CREATE] etkinlik oluşturulamadı: {exc}')
            if image_path:
                try:
                    os.remove(os.path.join(UPLOAD_EVENTS, image_path))
                except OSError:
                    pass
            conn.close()
            flash('Etkinlik oluşturulurken bir hata oluştu. Lütfen tekrar deneyin.', 'danger')
            return render_template('create_event.html')
        finally:
            try:
                conn.close()
            except Exception:
                pass

        # Bildirim gönderimi etkinlik kaydından bağımsızdır; push tarafındaki
        # olası bir hata etkinliğin oluşturulmasını asla geri almamalı.
        try:
            create_notification_for_all(
                title="Yeni Etkinlik!",
                message=f'"{title}" etkinliği oluşturuldu. {location} - {event_date}',
                link=url_for('event_detail', event_id=event_id)
            )
        except Exception as exc:
            print(f'[EVENT NOTIFICATION] bildirim gönderilemedi: {exc}')

        flash('Etkinlik başarıyla oluşturuldu!', 'success')
        return redirect(url_for('events'))

    return render_template('create_event.html')


@app.route('/event/<int:event_id>')
@login_required
def event_detail(event_id):
    conn = get_db()
    event = conn.execute('''
        SELECT e.*, u.full_name as creator_name
        FROM events e
        JOIN users u ON e.created_by = u.id
        WHERE e.id = ?
    ''', (event_id,)).fetchone()

    if not event:
        conn.close()
        flash('Etkinlik bulunamadı!', 'danger')
        return redirect(url_for('events'))

    participants = conn.execute('''
        SELECT u.username, u.full_name, u.room_number, ep.joined_at
        FROM event_participants ep
        JOIN users u ON ep.user_id = u.id
        WHERE ep.event_id = ?
        ORDER BY ep.joined_at ASC
    ''', (event_id,)).fetchall()

    is_joined = conn.execute('''
        SELECT id FROM event_participants WHERE event_id = ? AND user_id = ?
    ''', (event_id, session['user_id'])).fetchone()

    conn.close()

    return render_template('event_detail.html', event=event,
                            participants=participants, is_joined=is_joined)


@app.route('/event/<int:event_id>/join')
@login_required
def join_event(event_id):
    conn = get_db()
    existing = conn.execute('''
        SELECT id FROM event_participants WHERE event_id = ? AND user_id = ?
    ''', (event_id, session['user_id'])).fetchone()

    if existing:
        conn.close()
        flash('Zaten bu etkinliğe katıldınız!', 'info')
        return redirect(url_for('event_detail', event_id=event_id))

    event = conn.execute('SELECT max_participants FROM events WHERE id = ?',
                          (event_id,)).fetchone()
    count = conn.execute('SELECT COUNT(*) as c FROM event_participants WHERE event_id = ?',
                          (event_id,)).fetchone()['c']

    if event['max_participants'] and count >= int(event['max_participants']):
        conn.close()
        flash('Etkinlik kontenjanı dolu!', 'danger')
        return redirect(url_for('event_detail', event_id=event_id))

    conn.execute('''
        INSERT INTO event_participants (event_id, user_id)
        VALUES (?, ?)
    ''', (event_id, session['user_id']))
    conn.commit()
    conn.close()

    flash('Etkinliğe başarıyla katıldınız!', 'success')
    return redirect(url_for('event_detail', event_id=event_id))


@app.route('/event/<int:event_id>/leave')
@login_required
def leave_event(event_id):
    conn = get_db()
    conn.execute('''
        DELETE FROM event_participants WHERE event_id = ? AND user_id = ?
    ''', (event_id, session['user_id']))
    conn.commit()
    conn.close()

    flash('Etkinlikten ayrıldınız.', 'info')
    return redirect(url_for('event_detail', event_id=event_id))


# ==================== YOKLAMA SİSTEMİ ====================

@app.route('/attendance')
@login_required
def attendance_list():
    conn = get_db()
    current_u = conn.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    current_u_dict = dict(current_u) if current_u else {}
    can_manage = can_user_manage_attendance(current_u_dict)

    if can_manage:
        attendances = conn.execute('''
            SELECT a.*, u.full_name as creator_name,
                   (SELECT COUNT(*) FROM attendance_records WHERE attendance_id = a.id AND status = 'present') as present_count,
                   (SELECT COUNT(*) FROM attendance_records WHERE attendance_id = a.id AND status = 'absent') as absent_count
            FROM attendances a
            JOIN users u ON a.created_by = u.id
            ORDER BY a.created_at DESC
        ''').fetchall()
    else:
        attendances = conn.execute('''
            SELECT a.*,
                   (SELECT status FROM attendance_records
                    WHERE attendance_id = a.id AND user_id = ?) as my_status
            FROM attendances a
            WHERE a.is_active = 1
            ORDER BY a.created_at DESC
        ''', (session['user_id'],)).fetchall()

    conn.close()
    return render_template('attendance.html', attendances=attendances, can_manage=can_manage)


@app.route('/attendance/create', methods=['POST'])
@attendance_manager_required
def create_attendance():
    title = request.form['title']
    description = request.form.get('description', '')
    attendance_date = request.form['attendance_date']
    is_mandatory = 1 if request.form.get('is_mandatory') else 0

    conn = get_db()
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO attendances (title, description, attendance_date, created_by, is_mandatory)
        VALUES (?, ?, ?, ?, ?)
    ''', (title, description, attendance_date, session['user_id'], is_mandatory))
    attendance_id = cur.lastrowid
    conn.commit()
    conn.close()

    create_notification_for_all(
        title="Yeni Yoklama!",
        message=f'"{title}" yoklaması açıldı. Lütfen durumunuzu işaretleyin!',
        link=url_for('attendance_detail', attendance_id=attendance_id),
        only_in_dorm=False
    )

    flash('Yoklama oluşturuldu ve bildirim gönderildi!', 'success')
    return redirect(url_for('attendance_list'))


@app.route('/attendance/<int:attendance_id>')
@login_required
def attendance_detail(attendance_id):
    conn = get_db()
    current_u = conn.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    current_u_dict = dict(current_u) if current_u else {}
    can_manage = can_user_manage_attendance(current_u_dict)

    attendance = conn.execute('''
        SELECT a.*, u.full_name as creator_name
        FROM attendances a
        JOIN users u ON a.created_by = u.id
        WHERE a.id = ?
    ''', (attendance_id,)).fetchone()

    if not attendance:
        conn.close()
        flash('Yoklama bulunamadı!', 'danger')
        return redirect(url_for('attendance_list'))

    if can_manage:
        records = conn.execute('''
            SELECT ar.*, u.full_name, u.username, u.room_number, u.role
            FROM attendance_records ar
            JOIN users u ON ar.user_id = u.id
            WHERE ar.attendance_id = ?
            ORDER BY ar.marked_at DESC
        ''', (attendance_id,)).fetchall()

        all_users = conn.execute('''
            SELECT id, full_name, username, room_number, role FROM users
            WHERE is_admin = 0
            ORDER BY full_name COLLATE NOCASE ASC
        ''').fetchall()

        marked_user_ids = [r['user_id'] for r in records]
        unmarked_users = [u for u in all_users if u['id'] not in marked_user_ids]
        my_status = None
    else:
        records = None
        unmarked_users = None
        my_status = conn.execute('''
            SELECT status FROM attendance_records
            WHERE attendance_id = ? AND user_id = ?
        ''', (attendance_id, session['user_id'])).fetchone()

    conn.close()

    return render_template('attendance_detail.html', attendance=attendance,
                            records=records, my_status=my_status, unmarked_users=unmarked_users, can_manage=can_manage)


@app.route('/attendance/<int:attendance_id>/mark', methods=['POST'])
@login_required
def mark_attendance(attendance_id):
    status = request.form.get('status', '').strip().lower()
    if status not in ('present', 'absent'):
        flash('Geçersiz yoklama durumu.', 'danger')
        return redirect(url_for('attendance_detail', attendance_id=attendance_id))

    conn = get_db()
    attendance = conn.execute('SELECT is_active FROM attendances WHERE id = ?',
                               (attendance_id,)).fetchone()

    if not attendance or not attendance['is_active']:
        conn.close()
        flash('Bu yoklama artık aktif değil!', 'danger')
        return redirect(url_for('attendance_list'))

    conn.execute('''
        DELETE FROM attendance_records WHERE attendance_id = ? AND user_id = ?
    ''', (attendance_id, session['user_id']))

    conn.execute('''
        INSERT INTO attendance_records (attendance_id, user_id, status)
        VALUES (?, ?, ?)
    ''', (attendance_id, session['user_id'], status))
    conn.commit()
    conn.close()

    award_xp(session['user_id'], 5, 'Yoklama bildirimi yapıldı')
    flash('Yoklama durumunuz kaydedildi.', 'success')
    return redirect(url_for('attendance_detail', attendance_id=attendance_id))


@app.route('/attendance/<int:attendance_id>/close')
@attendance_manager_required
def close_attendance(attendance_id):
    conn = get_db()
    conn.execute('UPDATE attendances SET is_active = 0 WHERE id = ?', (attendance_id,))
    conn.commit()
    conn.close()
    flash('Yoklama sonlandırıldı.', 'info')
    return redirect(url_for('attendance_detail', attendance_id=attendance_id))


@app.route('/attendance/<int:attendance_id>/delete')
@attendance_manager_required
def delete_attendance(attendance_id):
    conn = get_db()
    conn.execute('DELETE FROM attendance_records WHERE attendance_id = ?', (attendance_id,))
    conn.execute('DELETE FROM attendances WHERE id = ?', (attendance_id,))
    conn.commit()
    conn.close()
    flash('Yoklama silindi.', 'success')
    return redirect(url_for('admin_panel'))


def send_nightly_attendance():
    """Yönetici tarafından belirlenen saatte günlük otomatik zorunlu yoklama açar."""
    try:
        if get_app_setting('attendance_auto_enabled', '1') != '1':
            return

        now = now_tr()
        configured = get_app_setting('attendance_auto_time', '22:00') or '22:00'
        try:
            hour, minute = [int(x) for x in configured.split(':', 1)]
        except (ValueError, TypeError):
            hour, minute = 22, 0

        # Scheduler her dakika çalışır; yalnızca seçilen dakika geldiğinde işlem yapar.
        if now.hour != hour or now.minute != minute:
            return

        today_str = now.strftime('%Y-%m-%d')
        title = f"Gece Yoklaması - {today_str}"
        conn = get_db()
        existing = conn.execute(
            'SELECT id FROM attendances WHERE title = ? AND attendance_date LIKE ? LIMIT 1',
            (title, today_str + '%')
        ).fetchone()
        if existing:
            conn.close()
            return

        cur = conn.cursor()
        cur.execute('''
            INSERT INTO attendances (title, description, attendance_date, created_by, is_mandatory, is_active)
            VALUES (?, ?, ?, 1, 1, 1)
        ''', (title, 'Zorunlu gece yoklaması. Lütfen "Yurttayım / Yurtta değilim" durumunuzu işaretleyin.', now.strftime('%Y-%m-%d %H:%M:%S')))
        attendance_id = cur.lastrowid
        conn.commit()
        conn.close()

        try:
            create_notification_for_all(
                title="Gece Yoklaması Açıldı!",
                message='Bu akşamın yoklaması açıldı. Lütfen durumunuzu işaretleyin, doldurmak zorunludur.',
                link=f'/attendance/{attendance_id}',
                only_in_dorm=False
            )
        except Exception as exc:
            print(f'[ATTENDANCE NOTIFICATION] bildirim gönderilemedi: {exc}')
    except Exception as exc:
        print(f'[AUTO ATTENDANCE] hata: {exc}')


def check_scheduled_attendance():
    # Her dakika tetiklenir; Europe/Istanbul saatine göre ayarı kontrol eder.
    send_nightly_attendance()

def cleanup_expired_stories():
    """Süresi dolan (24 saatten eski) hikayeleri ve dosyalarını temizler."""
    conn = get_db()
    expired = conn.execute("SELECT * FROM stories WHERE expires_at <= datetime('now')").fetchall()
    for s in expired:
        try:
            if s['media_path']:
                os.remove(os.path.join(UPLOAD_STORIES, s['media_path']))
        except OSError:
            pass
    conn.execute("DELETE FROM story_views WHERE story_id IN (SELECT id FROM stories WHERE expires_at <= datetime('now'))")
    conn.execute("DELETE FROM stories WHERE expires_at <= datetime('now')")
    conn.commit()
    conn.close()


scheduler = BackgroundScheduler(daemon=True, timezone='Europe/Istanbul')


def start_scheduler():
    if not scheduler.running:
        scheduler.add_job(check_scheduled_attendance, 'interval', minutes=1,
                           id='nightly_attendance', replace_existing=True,
                           max_instances=1, coalesce=True)
        scheduler.add_job(cleanup_expired_stories, 'interval', hours=1,
                           id='cleanup_stories', replace_existing=True)
        scheduler.add_job(process_laundry_notifications, 'interval', minutes=1,
                           id='laundry_notifications', replace_existing=True, max_instances=1, coalesce=True)
        scheduler.start()


# ==================== YEMEKHANE ====================
@app.route('/menu')
@login_required
def menu_page():
    conn = get_db()
    today = now_tr().date().isoformat()
    menus = conn.execute('''
        SELECT m.*,
               ROUND(AVG(r.overall_rating),1) AS avg_rating,
               COUNT(r.id) AS rating_count,
               (SELECT overall_rating FROM menu_ratings WHERE menu_date=m.menu_date AND user_id=?) AS my_rating,
               (SELECT breakfast_rating FROM menu_ratings WHERE menu_date=m.menu_date AND user_id=?) AS my_breakfast_rating,
               (SELECT dinner_rating FROM menu_ratings WHERE menu_date=m.menu_date AND user_id=?) AS my_dinner_rating,
               (SELECT comment FROM menu_ratings WHERE menu_date=m.menu_date AND user_id=?) AS my_comment
        FROM meal_menus m
        LEFT JOIN menu_ratings r ON r.menu_date=m.menu_date
        GROUP BY m.id
        ORDER BY m.menu_date ASC
        LIMIT 100
    ''', (session['user_id'], session['user_id'], session['user_id'], session['user_id'])).fetchall()
    
    today_menu = next((m for m in menus if m['menu_date'] == today), None)
    
    # Mevcut ayların listesini çıkar (örn: '2026-06', '2026-08')
    months_available = []
    for m in menus:
        d = m['menu_date']
        if len(d) >= 7:
            ym = d[:7]
            if ym not in months_available:
                months_available.append(ym)

    conn.close()

    # Hafta içi: 07:00–09:00 kahvaltı. Hafta sonu: 09:00–11:00. Akşam: 18:00–20:00.
    is_weekend = now_tr().weekday() >= 5
    breakfast_start = get_app_setting('meal_breakfast_weekend_start' if is_weekend else 'meal_breakfast_weekday_start', '09:00' if is_weekend else '07:00')
    breakfast_end = get_app_setting('meal_breakfast_weekend_end' if is_weekend else 'meal_breakfast_weekday_end', '11:00' if is_weekend else '09:00')
    dinner_start = get_app_setting('meal_dinner_start', '18:00')
    dinner_end = get_app_setting('meal_dinner_end', '20:00')
    meal_schedule = {
        'breakfast': f'{breakfast_start} – {breakfast_end}',
        'dinner': f'{dinner_start} – {dinner_end}',
        'is_weekend': is_weekend
    }
    return render_template('menu.html', menus=menus, today=today, today_menu=today_menu, months_available=months_available, meal_schedule=meal_schedule)

@app.route('/menu/rate', methods=['POST'])
@login_required
def rate_menu():
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    try:
        if request.is_json:
            data = request.get_json() or {}
            menu_date = str(data.get('menu_date', '')).strip()
            overall = int(data.get('overall_rating', '0'))
            breakfast = int(data.get('breakfast_rating', '0') or 0) or None
            dinner = int(data.get('dinner_rating', '0') or 0) or None
            comment = str(data.get('comment', '')).strip()[:300]
        else:
            menu_date = request.form.get('menu_date', '').strip()
            overall = int(request.form.get('overall_rating', '0'))
            breakfast = int(request.form.get('breakfast_rating', '0') or 0) or None
            dinner = int(request.form.get('dinner_rating', '0') or 0) or None
            comment = request.form.get('comment', '').strip()[:300]
    except (ValueError, TypeError):
        if is_ajax:
            return jsonify({'ok': False, 'error': 'Geçersiz puan değeri.'}), 400
        flash('Geçersiz puan.', 'danger')
        return redirect(url_for('menu_page'))

    if menu_date != now_tr().date().isoformat():
        msg = 'Yalnızca bugünün menüsünü puanlayabilirsin. Her gün yeni bir puanlama açılır. ⭐'
        if is_ajax:
            return jsonify({'ok': False, 'error': msg}), 400
        flash(msg, 'warning')
        return redirect(url_for('menu_page'))

    if overall not in range(1, 6) or (breakfast is not None and breakfast not in range(1,6)) or (dinner is not None and dinner not in range(1,6)):
        msg = 'Puan 1 ile 5 arasında olmalıdır.'
        if is_ajax:
            return jsonify({'ok': False, 'error': msg}), 400
        flash(msg, 'danger')
        return redirect(url_for('menu_page'))

    conn = get_db()
    exists = conn.execute('SELECT id FROM meal_menus WHERE menu_date=?', (menu_date,)).fetchone()
    if not exists:
        conn.close()
        msg = 'Bu tarihe ait menü bulunamadı.'
        if is_ajax:
            return jsonify({'ok': False, 'error': msg}), 404
        flash(msg, 'danger')
        return redirect(url_for('menu_page'))

    conn.execute('''
        INSERT INTO menu_ratings(menu_date,user_id,breakfast_rating,dinner_rating,overall_rating,comment)
        VALUES(?,?,?,?,?,?)
        ON CONFLICT(menu_date,user_id) DO UPDATE SET
            breakfast_rating=excluded.breakfast_rating,
            dinner_rating=excluded.dinner_rating,
            overall_rating=excluded.overall_rating,
            comment=excluded.comment,
            updated_at=CURRENT_TIMESTAMP
    ''', (menu_date, session['user_id'], breakfast, dinner, overall, comment))
    conn.commit()

    # Güncel ortalama ve oy sayısını çek
    stat = conn.execute('SELECT ROUND(AVG(overall_rating),1) as avg_r, COUNT(id) as count_r FROM menu_ratings WHERE menu_date=?', (menu_date,)).fetchone()
    conn.close()

    award_xp(session['user_id'], 10, 'Günün menüsü değerlendirildi')

    if is_ajax:
        return jsonify({
            'ok': True,
            'message': 'Günün menüsü puanın kaydedildi! (+10 XP) ⭐',
            'my_rating': overall,
            'avg_rating': stat['avg_r'] if stat and stat['avg_r'] is not None else overall,
            'rating_count': stat['count_r'] if stat else 1
        })

    flash('Günün menüsü puanın kaydedildi! (+10 XP) Teşekkürler! ⭐', 'success')
    return redirect(url_for('menu_page'))

@app.route('/admin/menu/schedule', methods=['POST'])
@permission_required('manage_menu')
def update_meal_schedule():
    keys = [
        'meal_breakfast_weekday_start', 'meal_breakfast_weekday_end',
        'meal_breakfast_weekend_start', 'meal_breakfast_weekend_end',
        'meal_dinner_start', 'meal_dinner_end'
    ]
    values = {k: request.form.get(k, '').strip() for k in keys}
    if any(not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d', values[k]) for k in keys):
        flash('Geçersiz yemek saati. Lütfen tüm saatleri HH:MM biçiminde girin.', 'danger')
        return redirect(url_for('admin_panel'))

    ranges = [
        ('meal_breakfast_weekday_start', 'meal_breakfast_weekday_end'),
        ('meal_breakfast_weekend_start', 'meal_breakfast_weekend_end'),
        ('meal_dinner_start', 'meal_dinner_end')
    ]
    for start_key, end_key in ranges:
        if values[start_key] >= values[end_key]:
            flash('Yemek başlangıç saati, bitiş saatinden önce olmalıdır.', 'danger')
            return redirect(url_for('admin_panel'))

    for key, value in values.items():
        set_app_setting(key, value)
    flash('Yemekhane saatleri güncellendi.', 'success')
    return redirect(url_for('admin_panel'))


@app.route('/admin/menu/upload',methods=['POST'])
@permission_required('manage_menu')
def upload_menu():
    f=request.files.get('menu_file')
    meal_type = request.form.get('meal_type', 'auto').strip().lower()
    if meal_type not in ('auto','breakfast','dinner','lunch'):
        meal_type = 'auto'
    if not f or not f.filename.lower().endswith(('.xlsx','.xlsm','.xls')):
        flash('Lütfen .xls, .xlsx veya .xlsm Excel dosyası yükleyin.','danger')
        return redirect(url_for('admin_panel'))
    name=f'{uuid.uuid4().hex}_{secure_filename(f.filename)}'
    path=os.path.join(UPLOAD_MENU,name)
    f.save(path)
    try:
        rows=parse_menu_excel(path, meal_type=meal_type)
        if not rows:
            raise ValueError('Tarih ve menü satırı bulunamadı. Excel başlıklarını kontrol edin.')
        conn=get_db()
        for row in rows:
            ds, breakfast, lunch, dinner, snack = row
            if meal_type == 'breakfast':
                lunch = ''; dinner = ''; snack = ''
            elif meal_type == 'dinner':
                breakfast = ''; lunch = ''; snack = ''
            elif meal_type == 'lunch':
                breakfast = ''; dinner = ''; snack = ''
            
            conn.execute('''
                INSERT INTO meal_menus(menu_date,breakfast,lunch,dinner,snack,source_file)
                VALUES (?,?,?,?,?,?)
                ON CONFLICT(menu_date) DO UPDATE SET
                    breakfast = CASE 
                        WHEN excluded.breakfast != '' THEN excluded.breakfast 
                        ELSE meal_menus.breakfast 
                    END,
                    lunch = CASE 
                        WHEN excluded.lunch != '' THEN excluded.lunch 
                        ELSE meal_menus.lunch 
                    END,
                    dinner = CASE 
                        WHEN excluded.dinner != '' THEN excluded.dinner 
                        ELSE meal_menus.dinner 
                    END,
                    snack = CASE 
                        WHEN excluded.snack != '' THEN excluded.snack 
                        ELSE meal_menus.snack 
                    END,
                    source_file = excluded.source_file
            ''', (ds, breakfast, lunch, dinner, snack, name))
        conn.commit()
        conn.close()
        flash(f'{len(rows)} günlük menü başarıyla içe aktarıldı. 🍽️','success')
    except Exception as e:
        flash(f'Excel okunamadı: {e}','danger')
    return redirect(url_for('admin_panel'))

@app.route('/admin/menu/save', methods=['POST'])
@permission_required('manage_menu')
def save_menu_manual():
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    if request.is_json:
        data = request.get_json() or {}
        menu_date = str(data.get('menu_date', '')).strip()
        breakfast = str(data.get('breakfast', '')).strip()
        dinner = str(data.get('dinner', '')).strip()
        snack = str(data.get('snack', '')).strip()
    else:
        menu_date = request.form.get('menu_date', '').strip()
        breakfast = request.form.get('breakfast', '').strip()
        dinner = request.form.get('dinner', '').strip()
        snack = request.form.get('snack', '').strip()

    if not menu_date:
        if is_ajax:
            return jsonify({'ok': False, 'error': 'Tarih seçimi zorunludur.'}), 400
        flash('Tarih seçimi zorunludur.', 'danger')
        return redirect(url_for('admin_panel'))

    conn = get_db()
    conn.execute('''
        INSERT INTO meal_menus(menu_date, breakfast, dinner, snack, source_file)
        VALUES (?, ?, ?, ?, 'manual_edit')
        ON CONFLICT(menu_date) DO UPDATE SET
            breakfast = excluded.breakfast,
            dinner = excluded.dinner,
            snack = excluded.snack,
            source_file = 'manual_edit'
    ''', (menu_date, breakfast, dinner, snack))
    conn.commit()
    conn.close()

    msg = f'{format_date_tr(menu_date)} tarihli menü başarıyla kaydedildi/güncellendi. 🍽️'
    if is_ajax:
        return jsonify({
            'ok': True,
            'message': msg,
            'menu_date': menu_date,
            'breakfast': breakfast,
            'dinner': dinner,
            'snack': snack
        })
    flash(msg, 'success')
    return redirect(request.referrer or url_for('admin_panel'))

@app.route('/admin/menu/delete-day/<int:menu_id>', methods=['POST'])
@permission_required('manage_menu')
def delete_single_day_menu(menu_id):
    conn = get_db()
    conn.execute('DELETE FROM meal_menus WHERE id = ?', (menu_id,))
    conn.commit()
    conn.close()
    flash('Günlük menü kaydı silindi.', 'success')
    return redirect(request.referrer or url_for('admin_panel'))

@app.route('/admin/menu/delete',methods=['POST'])
@permission_required('manage_menu')
def delete_menu():
    conn=get_db(); conn.execute('DELETE FROM meal_menus'); conn.commit(); conn.close(); flash('Menü kayıtları temizlendi.','success'); return redirect(url_for('admin_panel'))

# ==================== DUYURULAR ====================
@app.route('/announcements')
@login_required
def announcements_page():
    conn=get_db(); items=conn.execute('SELECT * FROM announcements ORDER BY CASE priority WHEN "urgent" THEN 0 WHEN "high" THEN 1 ELSE 2 END,created_at DESC').fetchall(); conn.close(); return render_template('announcements.html',announcements=items)
@app.route('/admin/announcement/create',methods=['POST'])
@permission_required('manage_announcements')
def create_announcement():
    title=request.form.get('title','').strip(); content=request.form.get('content','').strip(); priority=request.form.get('priority','normal'); expires=request.form.get('expires_at') or None
    if title and content:
        conn=get_db(); conn.execute('INSERT INTO announcements(title,content,priority,expires_at,created_by) VALUES (?,?,?,?,?)',(title,content,priority,expires,session['user_id'])); conn.commit(); conn.close(); flash('Duyuru yayınlandı.','success')
    return redirect(url_for('admin_panel'))
@app.route('/admin/announcement/<int:announcement_id>/delete',methods=['POST'])
@permission_required('manage_announcements')
def delete_announcement(announcement_id):
    conn=get_db(); conn.execute('DELETE FROM announcements WHERE id=?',(announcement_id,)); conn.commit(); conn.close(); return redirect(url_for('admin_panel'))

# ==================== TALEPLER / ARIZA ====================
@app.route('/maintenance')
@login_required
def maintenance_page():
    conn = get_db()
    current_u = conn.execute('SELECT * FROM users WHERE id = ?', (session['user_id'],)).fetchone()
    current_u_dict = dict(current_u) if current_u else {}
    is_staff = can_user_manage_maintenance(current_u_dict)
    
    if is_staff:
        all_requests = conn.execute('''
            SELECT m.*, u.full_name, u.username, u.room_number as user_room
            FROM maintenance_requests m
            JOIN users u ON u.id = m.user_id
            ORDER BY CASE m.status WHEN 'pending' THEN 0 WHEN 'in_progress' THEN 1 ELSE 2 END, m.created_at DESC
        ''').fetchall()
    else:
        all_requests = []
        
    mine = conn.execute('''
        SELECT * FROM maintenance_requests
        WHERE user_id = ?
        ORDER BY created_at DESC
    ''', (session['user_id'],)).fetchall()
    
    conn.close()
    return render_template('maintenance.html', requests=mine, all_requests=all_requests, is_staff=is_staff)

@app.route('/maintenance/create', methods=['POST'])
@login_required
def create_maintenance():
    desc = request.form.get('description', '').strip()
    if not desc:
        flash('Sorun açıklaması gerekli.', 'danger')
        return redirect(url_for('maintenance_page'))
    photo = None
    f = request.files.get('photo')
    if f and allowed_file(f.filename, IMAGE_EXT):
        photo = save_media(f, UPLOAD_POSTS, IMAGE_EXT)
    
    conn = get_db()
    conn.execute('''
        INSERT INTO maintenance_requests(user_id, category, room_number, description, photo_path, priority)
        VALUES (?, ?, ?, ?, ?, ?)
    ''', (session['user_id'], request.form.get('category', 'Diğer'), request.form.get('room_number', '').strip(), desc, photo, request.form.get('priority', 'normal')))
    conn.commit()
    conn.close()
    award_xp(session['user_id'], 5, 'Arıza/talep bildirimi')
    flash('Arıza talebiniz teknik ekibe ve yönetime iletildi. 🔧', 'success')
    return redirect(url_for('maintenance_page'))

@app.route('/admin/maintenance/<int:req_id>/status', methods=['POST'])
@maintenance_manager_required
def update_maintenance(req_id):
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    if request.is_json:
        data = request.get_json() or {}
        status = str(data.get('status', 'pending')).strip()
        note = str(data.get('admin_note', '')).strip()
    else:
        status = request.form.get('status', 'pending').strip()
        note = request.form.get('admin_note', '').strip()

    status_map = {
        'Bekliyor': 'pending',
        'İnceleniyor': 'in_progress',
        'Çözüldü': 'resolved',
        'İptal': 'cancelled',
        'pending': 'pending',
        'in_progress': 'in_progress',
        'resolved': 'resolved',
        'cancelled': 'cancelled'
    }
    status = status_map.get(status, status)
    if status not in ['pending', 'in_progress', 'resolved', 'cancelled']:
        status = 'pending'

    conn = get_db()
    row = conn.execute('SELECT user_id FROM maintenance_requests WHERE id = ?', (req_id,)).fetchone()
    conn.execute('''
        UPDATE maintenance_requests
        SET status = ?, admin_note = ?, updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
    ''', (status, note, req_id))
    conn.commit()
    conn.close()
    
    status_tr = {
        'pending': 'Beklemede',
        'in_progress': 'İnceleniyor / İşlemde',
        'resolved': 'Tamamlandı / Çözüldü',
        'cancelled': 'İptal Edildi'
    }.get(status, status)

    if row:
        create_notification(
            row['user_id'],
            'Arıza Talebiniz Güncellendi 🔧',
            f'Talebinizin durumu "{status_tr}" olarak güncellendi. Not: {note if note else "İşlem yapıldı."}',
            url_for('maintenance_page')
        )
        if status == 'resolved':
            award_xp(row['user_id'], 10, 'Arıza talebi çözüldü')
            
    msg = f'Arıza talebi durumu "{status_tr}" olarak güncellendi.'
    if is_ajax:
        return jsonify({'ok': True, 'status': status, 'status_tr': status_tr, 'admin_note': note, 'message': msg})

    flash(msg, 'success')
    return redirect(request.referrer or url_for('maintenance_page'))

# ==================== İLANLAR ====================
@app.route('/market')
@login_required
def market_page():
    conn=get_db(); items=conn.execute('SELECT m.*,u.full_name,u.username FROM marketplace_items m JOIN users u ON u.id=m.user_id WHERE m.status="active" ORDER BY m.created_at DESC').fetchall(); conn.close(); return render_template('market.html',items=items)
@app.route('/market/create',methods=['POST'])
@login_required
def market_create():
    title=request.form.get('title','').strip()
    if not title: flash('İlan başlığı gerekli.','danger'); return redirect(url_for('market_page'))
    conn=get_db(); conn.execute('INSERT INTO marketplace_items(user_id,title,description,price,category) VALUES (?,?,?,?,?)',(session['user_id'],title,request.form.get('description','').strip(),request.form.get('price') or None,request.form.get('category','Diğer'))); conn.commit(); conn.close(); award_xp(session['user_id'],5,'İlan yayınlandı'); flash('İlan yayınlandı.','success'); return redirect(url_for('market_page'))
@app.route('/market/<int:item_id>/delete',methods=['POST'])
@login_required
def market_delete(item_id):
    conn=get_db(); conn.execute('UPDATE marketplace_items SET status="sold" WHERE id=? AND (user_id=? OR ?=1)',(item_id,session['user_id'],int(has_perm('moderate_content')))); conn.commit(); conn.close(); return redirect(url_for('market_page'))

# ==================== XP / SIRALAMA ====================
@app.route('/leaderboard')
@login_required
def leaderboard_page():
    conn=get_db(); leaders=conn.execute('SELECT u.id,u.full_name,u.username,u.profile_photo,u.role,u.is_admin,COALESCE(SUM(x.amount),0) xp FROM users u LEFT JOIN xp_transactions x ON x.user_id=u.id GROUP BY u.id ORDER BY xp DESC LIMIT 50').fetchall(); badges=conn.execute('SELECT * FROM badges ORDER BY id').fetchall(); conn.close(); return render_template('leaderboard.html',leaders=leaders,badges=badges)

# ==================== ÇAMAŞIRHANE ====================
def _laundry_parse_dt(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value)).replace(tzinfo=ISTANBUL_TZ)
    except Exception:
        return None

LAUNDRY_SLOT_MINUTES = 10  # Rezervasyon başlangıçları bu dakikanın katlarında seçilebilir.


def _laundry_slot_start(dt):
    """Bir sonraki 10 dakikalık rezervasyon başlangıcını verir."""
    dt = dt.astimezone(ISTANBUL_TZ).replace(second=0, microsecond=0)
    remainder = dt.minute % LAUNDRY_SLOT_MINUTES
    if remainder == 0:
        return dt
    return dt + timedelta(minutes=LAUNDRY_SLOT_MINUTES - remainder)


def _laundry_cleanup_overlaps(conn):
    """Eski sürümlerde oluşmuş çakışan aktif rezervasyonları tek seferlik temizler."""
    now_s = now_tr().strftime('%Y-%m-%d %H:%M:%S')
    rows = conn.execute("""
        SELECT id, machine_id, start_at, end_at, created_at
        FROM laundry_reservations
        WHERE status='reserved' AND end_at > ?
        ORDER BY machine_id, start_at, created_at, id
    """, (now_s,)).fetchall()
    kept_by_machine = {}
    for row in rows:
        previous = kept_by_machine.setdefault(row['machine_id'], [])
        if any(row['start_at'] < p['end_at'] and row['end_at'] > p['start_at'] for p in previous):
            conn.execute("UPDATE laundry_reservations SET status='cancelled' WHERE id=?", (row['id'],))
        else:
            previous.append(row)


def _laundry_conflict(conn, machine_id, start_dt, end_dt, exclude_id=None):
    params=[machine_id, start_dt.strftime('%Y-%m-%d %H:%M:%S'), end_dt.strftime('%Y-%m-%d %H:%M:%S')]
    sql='''SELECT id FROM laundry_reservations
           WHERE machine_id=? AND status='reserved'
             AND start_at < ? AND end_at > ?'''
    if exclude_id:
        sql += ' AND id != ?'; params.append(exclude_id)
    return conn.execute(sql, params).fetchone()

@app.route('/laundry')
@login_required
def laundry_page():
    conn=get_db()
    machines=conn.execute('SELECT * FROM laundry_machines WHERE enabled=1 ORDER BY id').fetchall()
    now=now_tr()
    reservations=conn.execute('''
        SELECT r.*, m.name machine_name, u.full_name user_name
        FROM laundry_reservations r JOIN laundry_machines m ON m.id=r.machine_id
        JOIN users u ON u.id=r.user_id
        WHERE r.status='reserved' AND r.end_at >= ?
        ORDER BY r.start_at ASC
    ''',(now.strftime('%Y-%m-%d %H:%M:%S'),)).fetchall()
    active=[]; upcoming=[]
    for r in reservations:
        rd=dict(r); rd['start_iso']=r['start_at']; rd['end_iso']=r['end_at']; rd['is_mine']=r['user_id']==session['user_id']
        start=_laundry_parse_dt(r['start_at']); end=_laundry_parse_dt(r['end_at'])
        if start and start <= now < end: active.append(rd)
        elif start and start > now: upcoming.append(rd)
    wait=conn.execute('SELECT machine_id FROM laundry_waitlist WHERE user_id=? AND notified_at IS NULL',(session['user_id'],)).fetchall()
    wait_ids={r['machine_id'] for r in wait}
    # Önümüzdeki 24 saat için LAUNDRY_SLOT_MINUTES dakikalık başlangıç seçenekleri.
    first_slot=_laundry_slot_start(now)
    slot_count=(24*60)//LAUNDRY_SLOT_MINUTES
    slots=[first_slot+timedelta(minutes=LAUNDRY_SLOT_MINUTES*i) for i in range(slot_count)]
    bookings_by_machine={}
    for r in reservations:
        bookings_by_machine.setdefault(r['machine_id'], []).append({
            'start': r['start_at'], 'end': r['end_at'], 'mine': bool(r['user_id']==session['user_id'])
        })
    conn.close()
    return render_template('laundry.html', machines=machines, active_reservations=active,
                           upcoming_reservations=upcoming, wait_ids=wait_ids, slots=slots,
                           bookings_by_machine=bookings_by_machine, now_iso=now.isoformat())

@app.route('/laundry/reserve', methods=['POST'])
@login_required
def laundry_reserve():
    machine_id=request.form.get('machine_id', type=int)
    start_raw=request.form.get('start_at','').strip()
    if not machine_id or not start_raw:
        flash('Makine ve saat seçmelisin.','danger'); return redirect(url_for('laundry_page'))
    try:
        start=datetime.fromisoformat(start_raw).replace(tzinfo=ISTANBUL_TZ)
    except ValueError:
        flash('Geçersiz rezervasyon saati.','danger'); return redirect(url_for('laundry_page'))
    now=now_tr()
    if start < _laundry_slot_start(now):
        flash('Geçmiş bir saat için rezervasyon yapılamaz.','danger'); return redirect(url_for('laundry_page'))
    if start.minute % LAUNDRY_SLOT_MINUTES != 0 or start.second != 0:
        flash(f'Rezervasyon başlangıcı {LAUNDRY_SLOT_MINUTES} dakikalık aralıklarda seçilebilir.','danger'); return redirect(url_for('laundry_page'))
    duration_raw=request.form.get('duration_minutes','60').strip()
    try:
        duration=int(duration_raw)
    except ValueError:
        duration=60
    if duration not in (30,60,90,120,150,180):
        flash('Geçersiz rezervasyon süresi.','danger'); return redirect(url_for('laundry_page'))
    conn=get_db()
    try:
        # Aynı anda gelen iki rezervasyon isteğini seri hale getir.
        conn.execute('BEGIN IMMEDIATE')
        machine=conn.execute('SELECT * FROM laundry_machines WHERE id=? AND enabled=1',(machine_id,)).fetchone()
        if not machine:
            conn.rollback(); conn.close(); flash('Makine bulunamadı.','danger'); return redirect(url_for('laundry_page'))
        end=start+timedelta(minutes=duration)
        if _laundry_conflict(conn,machine_id,start,end):
            conn.rollback(); conn.close(); flash('Bu saat artık dolu. Başka bir saat seçebilirsin.','warning'); return redirect(url_for('laundry_page'))
        if conn.execute("SELECT id FROM laundry_reservations WHERE user_id=? AND status='reserved' AND start_at < ? AND end_at > ?",(session['user_id'],end.strftime('%Y-%m-%d %H:%M:%S'),start.strftime('%Y-%m-%d %H:%M:%S'))).fetchone():
            conn.rollback(); conn.close(); flash('Bu saat için zaten başka bir makine rezervasyonun var.','warning'); return redirect(url_for('laundry_page'))
        conn.execute('INSERT INTO laundry_reservations(machine_id,user_id,start_at,end_at) VALUES(?,?,?,?)',(machine_id,session['user_id'],start.strftime('%Y-%m-%d %H:%M:%S'),end.strftime('%Y-%m-%d %H:%M:%S')))
        conn.execute('DELETE FROM laundry_waitlist WHERE machine_id=? AND user_id=?',(machine_id,session['user_id']))
        conn.commit(); machine_name=machine['name']
    except sqlite3.OperationalError as exc:
        try: conn.rollback()
        except Exception: pass
        conn.close()
        flash('Çamaşırhane şu anda yoğun. Birkaç saniye sonra tekrar dene.' if 'locked' in str(exc).lower() else 'Rezervasyon oluşturulurken bir hata oluştu.','warning' if 'locked' in str(exc).lower() else 'danger')
        return redirect(url_for('laundry_page'))
    except Exception:
        try: conn.rollback()
        except Exception: pass
        conn.close(); flash('Rezervasyon oluşturulurken bir hata oluştu.','danger'); return redirect(url_for('laundry_page'))
    conn.close()
    flash(f'{machine_name} için {start.strftime("%H:%M")}–{end.strftime("%H:%M")} rezervasyonun oluşturuldu.','success')
    return redirect(url_for('laundry_page'))

@app.route('/laundry/reservation/<int:reservation_id>/finish', methods=['POST'])
@login_required
def laundry_finish(reservation_id):
    conn=get_db()
    row=conn.execute("""SELECT r.*, m.name machine_name FROM laundry_reservations r
                        JOIN laundry_machines m ON m.id=r.machine_id
                        WHERE r.id=? AND r.user_id=? AND r.status='reserved'""", (reservation_id, session['user_id'])).fetchone()
    if not row:
        conn.close(); flash('Aktif rezervasyon bulunamadı.','danger'); return redirect(url_for('laundry_page'))
    now_s=now_tr().strftime('%Y-%m-%d %H:%M:%S')
    conn.execute('UPDATE laundry_reservations SET status="completed", end_at=? WHERE id=?', (now_s, reservation_id))
    waiters=conn.execute('SELECT * FROM laundry_waitlist WHERE machine_id=? AND notified_at IS NULL ORDER BY created_at ASC', (row['machine_id'],)).fetchall()
    conn.commit(); conn.close()
    for w in waiters:
        create_notification(w['user_id'], 'Makine boşaldı 🧺', f'{row["machine_name"]} artık boş. Hemen rezerve edebilirsin.', url_for('laundry_page'))
    if waiters:
        conn=get_db()
        conn.execute('DELETE FROM laundry_waitlist WHERE machine_id=? AND notified_at IS NULL', (row['machine_id'],))
        conn.commit(); conn.close()
    flash(f'{row["machine_name"]} boş olarak işaretlendi. Bekleyenlere bildirim gönderildi.', 'success')
    return redirect(url_for('laundry_page'))

@app.route('/laundry/reservation/<int:reservation_id>/cancel', methods=['POST'])
@login_required
def laundry_cancel(reservation_id):
    conn=get_db(); row=conn.execute('SELECT * FROM laundry_reservations WHERE id=? AND user_id=? AND status="reserved"',(reservation_id,session['user_id'])).fetchone()
    if not row:
        conn.close(); flash('Rezervasyon bulunamadı.','danger'); return redirect(url_for('laundry_page'))
    if _laundry_parse_dt(row['start_at']) <= now_tr():
        conn.close(); flash('Başlamış rezervasyon iptal edilemez.','warning'); return redirect(url_for('laundry_page'))
    conn.execute('UPDATE laundry_reservations SET status="cancelled" WHERE id=?',(reservation_id,)); conn.commit(); conn.close()
    flash('Rezervasyon iptal edildi.','success'); return redirect(url_for('laundry_page'))

@app.route('/laundry/wait/<int:machine_id>', methods=['POST'])
@login_required
def laundry_wait(machine_id):
    conn=get_db(); machine=conn.execute('SELECT * FROM laundry_machines WHERE id=? AND enabled=1',(machine_id,)).fetchone()
    if not machine:
        conn.close(); return jsonify({'ok':False,'error':'Makine bulunamadı.'}),404
    conn.execute('INSERT OR IGNORE INTO laundry_waitlist(machine_id,user_id) VALUES(?,?)',(machine_id,session['user_id']))
    conn.commit(); conn.close()
    return jsonify({'ok':True,'message':f'{machine["name"]} boşaldığında sana bildirim göndereceğiz.'})

@app.route('/laundry/wait/<int:machine_id>/cancel', methods=['POST'])
@login_required
def laundry_wait_cancel(machine_id):
    conn=get_db(); conn.execute('DELETE FROM laundry_waitlist WHERE machine_id=? AND user_id=?',(machine_id,session['user_id'])); conn.commit(); conn.close()
    return jsonify({'ok':True})

def process_laundry_notifications():
    """Dakikada bir çalışır: süresi dolan makineleri tamamlar ve bekleyen herkese haber verir.

    10 dakika önceden hatırlatma yoktur. Bildirim yalnızca makine gerçekten boşaldığında
    gönderilir; kullanıcı daha erken bitirirse /finish rotası aynı bildirimi anında tetikler.
    """
    now=now_tr(); now_s=now.strftime('%Y-%m-%d %H:%M:%S')
    conn=get_db()
    ended=conn.execute("""SELECT r.*,m.name machine_name FROM laundry_reservations r JOIN laundry_machines m ON m.id=r.machine_id
                          WHERE r.status='reserved' AND r.end_at <= ?""",(now_s,)).fetchall()
    notifications=[]
    for r in ended:
        conn.execute('UPDATE laundry_reservations SET status="completed" WHERE id=?',(r['id'],))
        # Makinenin hemen ardından başka bir rezervasyon varsa gerçekten boşalmış sayılmaz.
        next_booking = conn.execute("""
            SELECT id FROM laundry_reservations
            WHERE machine_id=? AND status='reserved' AND start_at < ?
            ORDER BY start_at ASC LIMIT 1
        """, (r['machine_id'], r['end_at'])).fetchone()
        if not next_booking:
            waiters=conn.execute('SELECT * FROM laundry_waitlist WHERE machine_id=? AND notified_at IS NULL ORDER BY created_at ASC',(r['machine_id'],)).fetchall()
            for w in waiters:
                notifications.append((w['user_id'], r['machine_name']))
            conn.execute('DELETE FROM laundry_waitlist WHERE machine_id=? AND notified_at IS NULL',(r['machine_id'],))
    conn.commit(); conn.close()
    for user_id, machine_name in notifications:
        create_notification(user_id,'Makine boşaldı 🧺',f'{machine_name} artık boş. Hemen rezerve edebilirsin.',url_for('laundry_page'))


# ==================== ADMIN PANELİ ====================

@app.route('/admin/attendance-schedule', methods=['POST'])
@permission_required('manage_attendance')
def update_attendance_schedule():
    enabled = '1' if request.form.get('auto_attendance_enabled') == '1' else '0'
    schedule_time = request.form.get('auto_attendance_time', '22:00').strip()
    if not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d', schedule_time):
        flash('Geçersiz otomatik yoklama saati.', 'danger')
        return redirect(url_for('admin_panel'))
    set_app_setting('attendance_auto_enabled', enabled)
    set_app_setting('attendance_auto_time', schedule_time)
    flash(f'Otomatik yoklama {"açıldı" if enabled == "1" else "kapatıldı"}. Saat: {schedule_time}', 'success')
    return redirect(url_for('admin_panel'))

@app.route('/admin')
@admin_required
def admin_panel():
    conn = get_db()
    users = conn.execute('SELECT * FROM users ORDER BY created_at DESC').fetchall()
    events = conn.execute('''
        SELECT e.*, u.full_name as creator_name,
               (SELECT COUNT(*) FROM event_participants WHERE event_id = e.id) as participant_count
        FROM events e
        JOIN users u ON e.created_by = u.id
        ORDER BY e.created_at DESC
    ''').fetchall()
    messages = conn.execute('SELECT * FROM messages ORDER BY created_at DESC LIMIT 50').fetchall()

    attendances = conn.execute('''
        SELECT a.*, u.full_name as creator_name,
               (SELECT COUNT(*) FROM attendance_records WHERE attendance_id = a.id AND status = 'present') as present_count,
               (SELECT COUNT(*) FROM attendance_records WHERE attendance_id = a.id AND status = 'absent') as absent_count
        FROM attendances a
        JOIN users u ON a.created_by = u.id
        ORDER BY a.created_at DESC
    ''').fetchall()

    posts_count = conn.execute('SELECT COUNT(*) as c FROM posts').fetchone()['c']
    menu_count = conn.execute('SELECT COUNT(*) as c FROM meal_menus').fetchone()['c']
    announcement_list = conn.execute('SELECT * FROM announcements ORDER BY created_at DESC LIMIT 20').fetchall()
    maintenance = conn.execute('SELECT m.*,u.full_name,u.username FROM maintenance_requests m JOIN users u ON u.id=m.user_id ORDER BY m.created_at DESC LIMIT 50').fetchall()
    marketplace = conn.execute('SELECT m.*,u.full_name FROM marketplace_items m JOIN users u ON u.id=m.user_id ORDER BY m.created_at DESC LIMIT 50').fetchall()
    top_players = conn.execute('SELECT u.full_name,COALESCE(SUM(x.amount),0) xp FROM users u LEFT JOIN xp_transactions x ON x.user_id=u.id GROUP BY u.id ORDER BY xp DESC LIMIT 10').fetchall()
    recent_xp = conn.execute('SELECT x.*,u.full_name FROM xp_transactions x JOIN users u ON u.id=x.user_id ORDER BY x.created_at DESC LIMIT 20').fetchall()
    menu_rating_summary = conn.execute('''
        SELECT m.menu_date, m.breakfast, m.dinner,
               ROUND(AVG(r.overall_rating),1) avg_rating,
               ROUND(AVG(r.breakfast_rating),1) avg_breakfast,
               ROUND(AVG(r.dinner_rating),1) avg_dinner,
               COUNT(r.id) rating_count
        FROM meal_menus m
        LEFT JOIN menu_ratings r ON r.menu_date=m.menu_date
        GROUP BY m.menu_date
        ORDER BY m.menu_date DESC LIMIT 31
    ''').fetchall()

    poll_rows = conn.execute('SELECT p.*, (SELECT COUNT(*) FROM poll_votes v WHERE v.poll_id=p.id) vote_count FROM polls p ORDER BY p.created_at DESC LIMIT 10').fetchall()
    all_menus = conn.execute('SELECT * FROM meal_menus ORDER BY menu_date DESC LIMIT 60').fetchall()
    laundry_machines = conn.execute('SELECT * FROM laundry_machines ORDER BY id').fetchall()
    laundry_reservations = conn.execute('''SELECT r.*,m.name machine_name,u.full_name user_name FROM laundry_reservations r JOIN laundry_machines m ON m.id=r.machine_id JOIN users u ON u.id=r.user_id WHERE r.status='reserved' AND r.end_at >= ? ORDER BY r.start_at''',(now_tr().strftime('%Y-%m-%d %H:%M:%S'),)).fetchall()

    conn.close()

    stats = {
        'total_users': len(users),
        'total_events': len(events),
        'total_messages': len(messages),
        'total_attendances': len(attendances),
        'total_posts': posts_count,
        'total_menu': menu_count,
    }

    attendance_auto_enabled = get_app_setting('attendance_auto_enabled', '1') == '1'
    attendance_auto_time = get_app_setting('attendance_auto_time', '22:00') or '22:00'
    meal_schedule_admin = {
        'weekday_breakfast_start': get_app_setting('meal_breakfast_weekday_start', '07:00'),
        'weekday_breakfast_end': get_app_setting('meal_breakfast_weekday_end', '09:00'),
        'weekend_breakfast_start': get_app_setting('meal_breakfast_weekend_start', '09:00'),
        'weekend_breakfast_end': get_app_setting('meal_breakfast_weekend_end', '11:00'),
        'dinner_start': get_app_setting('meal_dinner_start', '18:00'),
        'dinner_end': get_app_setting('meal_dinner_end', '20:00')
    }
    conn = get_db()
    role_counts = _role_user_counts(conn)
    conn.close()
    roles_payload = [_role_public(r, role_counts.get(k, 0)) for k, r in USER_ROLES.items()]
    perm_catalog = {
        'pages': [{'key': k, 'label': v['label'], 'icon': v['icon'], 'desc': v['desc']} for k, v in PAGE_PERMISSIONS.items()],
        'admin': [{'key': k, 'label': v['label'], 'icon': v['icon'], 'desc': v['desc']} for k, v in ADMIN_PERMISSIONS.items()],
    }
    return render_template('admin.html', users=users, events=events, messages=messages, attendances=attendances, stats=stats, announcement_list=announcement_list, maintenance=maintenance, marketplace=marketplace, top_players=top_players, recent_xp=recent_xp, menu_rating_summary=menu_rating_summary, all_menus=all_menus, attendance_auto_enabled=attendance_auto_enabled, attendance_auto_time=attendance_auto_time, meal_schedule_admin=meal_schedule_admin, laundry_machines=laundry_machines, laundry_reservations=laundry_reservations, poll_rows=poll_rows, roles_payload=roles_payload, perm_catalog=perm_catalog, role_counts=role_counts, role_icon_choices=ROLE_ICON_CHOICES, role_color_palette=ROLE_COLOR_PALETTE, is_full_admin=has_perm('full_admin'))


@app.route('/admin/laundry/machine/add', methods=['POST'])
@permission_required('manage_laundry')
def admin_laundry_machine_add():
    name=request.form.get('name','').strip()
    if not name:
        flash('Makine adı gerekli.','danger'); return redirect(url_for('admin_panel'))
    conn=get_db()
    try:
        conn.execute('INSERT INTO laundry_machines(name,slot_minutes,enabled) VALUES(?,60,1)',(name,)); conn.commit(); flash('Makine eklendi.','success')
    except sqlite3.IntegrityError:
        flash('Bu makine adı zaten mevcut.','warning')
    conn.close(); return redirect(url_for('admin_panel'))

@app.route('/admin/laundry/machine/<int:machine_id>/toggle', methods=['POST'])
@permission_required('manage_laundry')
def admin_laundry_machine_toggle(machine_id):
    conn=get_db(); conn.execute('UPDATE laundry_machines SET enabled=CASE WHEN enabled=1 THEN 0 ELSE 1 END WHERE id=?',(machine_id,)); conn.commit(); conn.close(); return redirect(url_for('admin_panel'))

@app.route('/admin/user/<int:user_id>/delete')
@permission_required('manage_users')
def delete_user(user_id):
    if user_id == session['user_id']:
        flash('Kendinizi silemezsiniz!', 'danger')
        return redirect(url_for('admin_panel'))

    # NOT: PRAGMA foreign_keys, açık bir transaction sırasında değiştirilemez
    # (SQLite bunu no-op sayar). Bu yüzden bu işlem için ayrı, taze bir bağlantı
    # açıp FK kontrolünü herhangi bir DELETE çalışmadan ÖNCE kapatıyoruz.
    conn = sqlite3.connect(DATABASE, timeout=10)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA foreign_keys = OFF')
    try:
        # Bu uygulamada 'users' tablosuna referans veren onlarca tablo var
        # (posts, follows, quiz, laundry, vb.) ve bunların tamamını burada
        # tek tek temizlemek bu değişikliğin kapsamı dışında. Mesajlaşmayla
        # doğrudan ilgili tabloları güvenli şekilde temizliyoruz; FK kontrolü
        # bu bağlantı için kapalı olduğundan diğer (mesajlaşma dışı) tablolardaki
        # geçmiş kayıtlar silme işlemini bloke etmiyor.
        conn.execute('DELETE FROM event_participants WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM attendance_records WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM notifications WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM posts WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM post_likes WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM post_comments WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM direct_messages WHERE sender_id = ? OR receiver_id = ?', (user_id, user_id))
        conn.execute('DELETE FROM conv_messages WHERE sender_id = ?', (user_id,))
        conn.execute('DELETE FROM conversation_members WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM message_hides WHERE user_id = ?', (user_id,))
        conn.execute('DELETE FROM blocked_users WHERE blocker_id = ? OR blocked_id = ?', (user_id, user_id))
        conn.execute('DELETE FROM message_reports WHERE reported_by = ?', (user_id,))
        conn.execute('DELETE FROM users WHERE id = ?', (user_id,))
        conn.commit()
        flash('Kullanıcı silindi.', 'success')
    except Exception:
        conn.rollback()
        flash('Kullanıcı silinemedi: bu kullanıcıya bağlı kayıtlar mevcut. Lütfen tekrar deneyin.', 'danger')
    finally:
        conn.close()

    return redirect(url_for('admin_panel'))


@app.route('/admin/user/<int:user_id>/toggle-dorm')
@permission_required('manage_users')
def toggle_in_dorm(user_id):
    conn = get_db()
    user = conn.execute('SELECT in_dorm FROM users WHERE id = ?', (user_id,)).fetchone()
    if user:
        conn.execute('UPDATE users SET in_dorm = ? WHERE id = ?',
                      (0 if user['in_dorm'] else 1, user_id))
        conn.commit()
    conn.close()
    flash('Kullanıcının yurt durumu güncellendi.', 'success')
    return redirect(url_for('admin_panel'))


@app.route('/admin/user/<int:user_id>/role', methods=['POST'])
@permission_required('manage_users')
def set_user_role(user_id):
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    if request.is_json:
        data = request.get_json() or {}
        role = data.get('role', 'ogrenci')
    else:
        role = request.form.get('role', 'ogrenci')

    if role not in USER_ROLES:
        role = 'ogrenci'

    def fail(msg, code=400):
        if is_ajax:
            return jsonify({'ok': False, 'error': msg}), code
        flash(msg, 'danger')
        return redirect(request.referrer or url_for('admin_panel'))

    if user_id == 1 and role != 'yurt_muduru':
        return fail('Kurucu yönetici hesabının rolü değiştirilemez.')
    me = current_user_dict()
    target_perms = USER_ROLES[role].get('permissions') or {}
    if not has_perm('full_admin'):
        if target_perms.get('full_admin') or target_perms.get('manage_roles'):
            return fail('Tam yönetici ya da rol yönetimi içeren bir rolü yalnızca Tam Yönetici atayabilir.', 403)
        if user_id == me.get('id'):
            return fail('Kendi rolünüzü değiştiremezsiniz.', 403)

    is_admin = 1 if is_admin_role(role) else 0

    conn = get_db()
    user = conn.execute('SELECT id, username, full_name FROM users WHERE id = ?', (user_id,)).fetchone()
    if user:
        conn.execute('UPDATE users SET role = ?, is_admin = ? WHERE id = ?', (role, is_admin, user_id))
        conn.commit()
    conn.close()

    role_title = USER_ROLES[role]['title']
    if user:
        create_notification(
            user_id,
            'Yetki & Rol Tanımlandı! 🎖️',
            f'Yönetim tarafından hesabınıza "{role_title}" rolü atandı.',
            url_for('profile', username=user['username'])
        )

    msg = f'Kullanıcı rolü "{role_title}" olarak güncellendi.'
    if is_ajax:
        return jsonify({
            'ok': True,
            'role': role,
            'role_title': role_title,
            'role_badge': USER_ROLES[role]['badge'],
            'role_class': USER_ROLES[role]['class'],
            'role_icon': USER_ROLES[role]['icon'],
            'role_color': USER_ROLES[role]['color'],
            'message': msg
        })

    flash(msg, 'success')
    return redirect(request.referrer or url_for('admin_panel'))


# ---------- ROL & YETKİ YÖNETİMİ (JSON API) ----------

def _role_public(role, user_count=None):
    d = {k: role.get(k) for k in ('key', 'title', 'badge', 'icon', 'color', 'desc', 'is_system', 'sort_order')}
    d['permissions'] = dict(role.get('permissions') or {})
    d['locked'] = role['key'] in LOCKED_ROLE_KEYS
    d['is_admin'] = role_has_admin_access(role)
    if user_count is not None:
        d['user_count'] = user_count
    return d


def _role_user_counts(conn):
    rows = conn.execute("SELECT COALESCE(role,'ogrenci') role, COUNT(*) c FROM users GROUP BY COALESCE(role,'ogrenci')").fetchall()
    return {r['role']: r['c'] for r in rows}


@app.route('/admin/roles', methods=['GET'])
@permission_required('manage_roles')
def admin_roles_list():
    conn = get_db()
    counts = _role_user_counts(conn)
    conn.close()
    return jsonify({'ok': True, 'roles': [_role_public(r, counts.get(k, 0)) for k, r in USER_ROLES.items()]})


@app.route('/admin/roles/save', methods=['POST'])
@permission_required('manage_roles')
def admin_role_save():
    data = request.get_json(silent=True) or {}
    key = (data.get('key') or '').strip()
    title = (data.get('title') or '').strip()[:60]
    badge = (data.get('badge') or '').strip()[:30] or title[:30]
    icon = (data.get('icon') or 'bi-person').strip()
    color = (data.get('color') or '#6b7280').strip()
    desc = (data.get('desc') or '').strip()[:200]
    raw_perms = data.get('permissions') or {}

    if not title:
        return jsonify({'ok': False, 'error': 'Rol adı boş olamaz.'}), 400
    if not re.fullmatch(r'bi-[a-z0-9-]{1,40}', icon):
        icon = 'bi-person'
    if not re.fullmatch(r'#[0-9a-fA-F]{6}', color):
        color = '#6b7280'

    # Sayfalar için yalnızca kapalı olanları (False), yönetim yetkileri için
    # yalnızca açık olanları (True) sakla.
    perms = {}
    for k in PAGE_PERMISSIONS:
        if raw_perms.get(k) is False:
            perms[k] = False
    for k in ADMIN_PERMISSIONS:
        if raw_perms.get(k):
            perms[k] = True
    if perms.get('full_admin'):
        perms = {'full_admin': True}

    is_full = has_perm('full_admin')
    if not is_full and (perms.get('full_admin') or perms.get('manage_roles')):
        return jsonify({'ok': False, 'error': 'Tam Yönetici ve Rol Yönetimi yetkilerini yalnızca Tam Yönetici verebilir.'}), 403

    conn = get_db()
    if key:
        existing = USER_ROLES.get(key)
        if not existing:
            conn.close()
            return jsonify({'ok': False, 'error': 'Rol bulunamadı.'}), 404
        if key == 'yurt_muduru':
            perms = {'full_admin': True}  # kurucu rol her zaman tam yetkili
        if not is_full and (existing.get('permissions') or {}).get('full_admin'):
            conn.close()
            return jsonify({'ok': False, 'error': 'Tam yetkili bir rolü yalnızca Tam Yönetici düzenleyebilir.'}), 403
        conn.execute('UPDATE roles SET title=?, badge=?, icon=?, color=?, desc=?, permissions=?, '
                     'updated_at=CURRENT_TIMESTAMP WHERE key=?',
                     (title, badge, icon, color, desc, json.dumps(perms), key))
        msg = f'"{title}" rolü güncellendi.'
    else:
        base_key = slugify_role_key(title)
        key, i = base_key, 2
        while key in USER_ROLES:
            key = f'{base_key}_{i}'
            i += 1
        max_order = conn.execute("SELECT COALESCE(MAX(sort_order), 0) m FROM roles WHERE key != 'ogrenci'").fetchone()['m']
        conn.execute('INSERT INTO roles (key, title, badge, icon, color, desc, permissions, is_system, sort_order) '
                     'VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?)',
                     (key, title, badge, icon, color, desc, json.dumps(perms), min(max_order + 1, 98)))
        msg = f'"{title}" rolü oluşturuldu.'
    flag = 1 if any(perms.get(k) for k in ADMIN_PERMISSIONS) else 0
    conn.execute('UPDATE users SET is_admin = ? WHERE role = ? AND id != 1', (flag, key))
    conn.commit()
    counts = _role_user_counts(conn)
    conn.close()
    reload_user_roles()
    return jsonify({'ok': True, 'message': msg, 'role': _role_public(USER_ROLES[key], counts.get(key, 0))})


@app.route('/admin/roles/<role_key>/delete', methods=['POST'])
@permission_required('manage_roles')
def admin_role_delete(role_key):
    role = USER_ROLES.get(role_key)
    if not role:
        return jsonify({'ok': False, 'error': 'Rol bulunamadı.'}), 404
    if role_key in LOCKED_ROLE_KEYS:
        return jsonify({'ok': False, 'error': 'Bu temel rol silinemez.'}), 400
    if (role.get('permissions') or {}).get('full_admin') and not has_perm('full_admin'):
        return jsonify({'ok': False, 'error': 'Tam yetkili bir rolü yalnızca Tam Yönetici silebilir.'}), 403
    conn = get_db()
    moved = conn.execute('SELECT COUNT(*) c FROM users WHERE role = ?', (role_key,)).fetchone()['c']
    conn.execute("UPDATE users SET role = 'ogrenci', is_admin = 0 WHERE role = ?", (role_key,))
    conn.execute('DELETE FROM roles WHERE key = ?', (role_key,))
    conn.commit()
    conn.close()
    reload_user_roles()
    extra = f' {moved} kullanıcı "Öğrenci" rolüne taşındı.' if moved else ''
    return jsonify({'ok': True, 'message': f'"{role["title"]}" rolü silindi.{extra}', 'moved': moved})


@app.route('/admin/roles/reorder', methods=['POST'])
@permission_required('manage_roles')
def admin_role_reorder():
    order = (request.get_json(silent=True) or {}).get('order') or []
    conn = get_db()
    for i, k in enumerate(order):
        if k in USER_ROLES and k != 'ogrenci':
            conn.execute('UPDATE roles SET sort_order = ? WHERE key = ?', (i + 1, k))
    conn.commit()
    conn.close()
    reload_user_roles()
    return jsonify({'ok': True})


@app.route('/admin/user/<int:user_id>/theme', methods=['POST'])
@permission_required('manage_users')
def set_user_theme(user_id):
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json
    if request.is_json:
        data = request.get_json() or {}
        theme = data.get('theme', 'default')
    else:
        theme = request.form.get('theme', 'default')

    if theme not in PROFILE_THEMES:
        theme = 'default'

    conn = get_db()
    user = conn.execute('SELECT id, username, full_name FROM users WHERE id = ?', (user_id,)).fetchone()
    if user:
        conn.execute('UPDATE users SET profile_theme = ? WHERE id = ?', (theme, user_id))
        conn.commit()
    conn.close()

    msg = f'Profil teması "{PROFILE_THEMES[theme]}" olarak güncellendi.'
    if is_ajax:
        return jsonify({'ok': True, 'theme': theme, 'theme_label': PROFILE_THEMES[theme], 'message': msg})

    flash(msg, 'success')
    return redirect(request.referrer or url_for('admin_panel'))


@app.route('/admin/event/<int:event_id>/delete')
@permission_required('manage_events')
def delete_event(event_id):
    conn = get_db()
    event = conn.execute('SELECT image_path FROM events WHERE id = ?', (event_id,)).fetchone()
    conn.execute('DELETE FROM event_participants WHERE event_id = ?', (event_id,))
    conn.execute('DELETE FROM events WHERE id = ?', (event_id,))
    conn.commit()
    conn.close()
    if event and event['image_path']:
        try: os.remove(os.path.join(UPLOAD_EVENTS, event['image_path']))
        except OSError: pass

    flash('Etkinlik silindi.', 'success')
    return redirect(url_for('admin_panel'))


@app.route('/admin/message/<int:msg_id>/delete')
@permission_required('moderate_content')
def delete_message(msg_id):
    conn = get_db()
    conn.execute('DELETE FROM messages WHERE id = ?', (msg_id,))
    conn.commit()
    conn.close()
    flash('Mesaj silindi.', 'success')
    return redirect(url_for('admin_panel'))


# ==================== BAŞLATMA ====================

init_db()
normalize_existing_usernames()
start_scheduler()

if __name__ == '__main__':
    # DEBUG MODU: Varsayılan olarak KAPALI. Werkzeug'un interaktif hata ayıklayıcısı
    # (debug=True) tünel üzerinden dışarıya açık bir uygulamada güvenlik açığı
    # oluşturabilir (uzaktan kod çalıştırmaya izin verebilir) ve her istekte ekstra
    # yük bindirerek "tünelde yavaşlık" hissini artırır. Yerel geliştirme yaparken
    # açmak isterseniz: YURT_DEBUG=true python app.py
    debug_mode = os.environ.get('YURT_DEBUG', 'false').lower() == 'true'

    # NOT: Flask-SocketIO, async_mode='threading' iken app.run()'ı zaten kendi
    # içinde threaded=True ile çağırıyor — bunu burada TEKRAR vermeye gerek yok
    # (verirsek "got multiple values for keyword argument 'threaded'" hatası
    # oluşur). İstekler zaten thread'lenmiş şekilde işleniyor, ekstra ayar
    # gerekmiyor.
    #
    # NOT 2: use_reloader=False -> arka planda çalışan zamanlayıcının (scheduler)
    # Werkzeug'un otomatik yeniden başlatma özelliği yüzünden iki kez başlamasını
    # engeller. Kod değişikliklerinde sunucuyu elle yeniden başlatın.
    socketio.run(app, host='0.0.0.0', port=int(os.environ.get('PORT', 5000)),
             debug=debug_mode, use_reloader=False, allow_unsafe_werkzeug=True)
